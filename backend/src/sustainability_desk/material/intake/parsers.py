# ABOUTME: DOCX 与 XLSX 的确定性 typed fragment 解析器。
# ABOUTME: 解析输出保留格式专属 SourceLocator，不执行公式，处理步骤记录输入输出指纹和质量标志。
# ABOUTME(en): Deterministic typed fragment parsers for DOCX and XLSX, plus PDF and PPTX text extraction.
# ABOUTME(en): Output keeps format-specific SourceLocators, evaluates no formulas, and records step fingerprints.
from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from sustainability_desk.material.intake.models import (
    DocxParagraphLocator,
    DocxTableLocator,
    NormalizedMaterial,
    PdfPageLocator,
    PptxSlideLocator,
    ProcessingStep,
    TableFragment,
    TextFragment,
    XlsxRangeLocator,
)
from sustainability_desk.material.intake.scanned_pdf_ocr import ScannedPdfOcrError, ocr_pdf_pages

MAX_FRAGMENT_CHARS = 20_000
MAX_XLSX_ROWS = 20_000
MAX_XLSX_CELLS = 100_000
MIN_PDF_PAGE_CHARS = 30
MIN_PDF_TOTAL_CHARS = 100
MIN_PDF_TEXT_PAGE_RATIO = 0.8


class MaterialParseError(ValueError):
    """已通过文件边界，但无法确定性解析。"""


def _clean(value: object) -> str:
    text = str(value).replace("\x00", " ")
    # pypdf 可能从可读文件中暴露异常 UTF-16 代理码点；必须在解析边界替换，
    # 否则持久化或计算指纹会把可恢复的提取结果变成失败。
    utf8_safe = text.encode("utf-8", errors="replace").decode("utf-8")
    return " ".join(utf8_safe.split()).strip()


def _fingerprint(data: bytes | str) -> str:
    payload = data if isinstance(data, bytes) else data.encode("utf-8", errors="replace")
    return hashlib.sha256(payload).hexdigest()


def pdf_has_usable_text(data: bytes) -> bool:
    """以保守、可测试的阈值判断 PDF 是否适合本地文本解析。"""
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise MaterialParseError("PDF 已加密")
        texts = [_clean(page.extract_text() or "") for page in reader.pages]
    except (PdfReadError, ValueError, TypeError) as error:
        raise MaterialParseError("PDF 无法解析") from error
    if not texts:
        raise MaterialParseError("PDF 不包含可处理页面")
    usable_pages = sum(len(text) >= MIN_PDF_PAGE_CHARS for text in texts)
    return (
        sum(len(text) for text in texts) >= MIN_PDF_TOTAL_CHARS
        and usable_pages / len(texts) >= MIN_PDF_TEXT_PAGE_RATIO
    )


def _page_has_image_resource(page: object) -> bool:
    """判断页面资源字典是否声明了图片 XObject；无法判断时保守视为无图片。"""
    try:
        resources = page.get_inherited("/Resources", {})
        resources = resources.get_object() if resources is not None else {}
        xobjects = resources.get("/XObject", {}) if resources else {}
        xobjects = xobjects.get_object() if xobjects else {}
    except Exception:
        return False
    if not xobjects:
        return False
    for reference in xobjects.values():
        try:
            if str(reference.get_object().get("/Subtype", "")) == "/Image":
                return True
        except Exception:
            continue
    return False


def parse_pdf_text(data: bytes, *, source_id: UUID, source_label: str) -> NormalizedMaterial:
    """按页提取有文本层的 PDF；无文本层但含图片资源的页触发本地 OCR 补全。"""
    started = datetime.now(timezone.utc)
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        pages = list(reader.pages)
        page_texts = [_clean(page.extract_text() or "") for page in pages]
    except (PdfReadError, ValueError, TypeError) as error:
        raise MaterialParseError("PDF 无法解析") from error

    ocr_page_indexes = [
        index
        for index, (page, text) in enumerate(zip(pages, page_texts))
        if not text and _page_has_image_resource(page)
    ]
    quality_flags: list[str] = []
    if ocr_page_indexes:
        ocr_results = ocr_pdf_pages(data, ocr_page_indexes)
        for index in ocr_page_indexes:
            result = ocr_results.get(index)
            if isinstance(result, ScannedPdfOcrError) or result is None:
                quality_flags.append("pdf_scanned_page_ocr_failed")
                continue
            if result.text.strip():
                page_texts[index] = result.text.strip()
            else:
                quality_flags.append("pdf_scanned_page_ocr_empty")

    fragments = [
        TextFragment(
            fragment_id=f"page-{page_number}",
            text=text[:MAX_FRAGMENT_CHARS],
            locator=PdfPageLocator(page=page_number),
        )
        for page_number, text in enumerate(page_texts, start=1)
        if text
    ]
    if not fragments:
        raise MaterialParseError("PDF 未提取到可读文本")
    rendered = "\n".join(fragment.text for fragment in fragments)
    completed = datetime.now(timezone.utc)
    message = "部分疑似扫描页的本地 OCR 未能识别文字" if quality_flags else ""
    return NormalizedMaterial(
        source_id=source_id,
        source_label=source_label,
        summary=f"已从 PDF 提取 {len(fragments)} 个逐页文本片段",
        fragments=fragments,
        processing_steps=[ProcessingStep(
            step="parsed",
            processor="deterministic",
            status="warning" if quality_flags else "succeeded",
            parser="pypdf",
            input_fingerprint=_fingerprint(data),
            output_fingerprint=_fingerprint(rendered),
            quality_flags=quality_flags,
            started_at=started,
            completed_at=completed,
            message=message,
        )],
        warnings=[message] if message else [],
    )


def _archive_quality_flags(data: bytes, kind: str) -> list[str]:
    """只在容器中实际出现当前解析器不覆盖的对象时标记人工关注。"""
    with ZipFile(BytesIO(data)) as archive:
        names = archive.namelist()
        if kind == "docx":
            document_xml = archive.read("word/document.xml")
            complex_markup = (b"<w:drawing", b"<w:txbxContent", b"<w:ins", b"<w:del")
            return ["docx_complex_objects_unparsed"] if (
                any(name.startswith("word/media/") for name in names)
                or any(marker in document_xml for marker in complex_markup)
            ) else []
        if kind == "pptx":
            return ["pptx_visual_objects_unparsed"] if any(
                name.startswith(("ppt/media/", "ppt/charts/", "ppt/embeddings/"))
                for name in names
            ) else []
        flags: list[str] = []
        if any(name.startswith(("xl/charts/", "xl/drawings/", "xl/media/")) for name in names):
            flags.append("xlsx_visual_objects_unparsed")
        if any(
            b"<f" in archive.read(name)
            for name in names
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        ):
            flags.append("xlsx_formula_cached_values_only")
        return flags


def parse_docx(data: bytes, *, source_id: UUID, source_label: str) -> NormalizedMaterial:
    """提取 DOCX 段落和表格；复杂浮动对象由 quality flag 明确保留。"""
    started = datetime.now(timezone.utc)
    try:
        document = Document(BytesIO(data))
    except Exception as error:
        raise MaterialParseError("DOCX 无法解析") from error
    fragments: list[TextFragment | TableFragment] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean(paragraph.text)
        if text:
            fragments.append(TextFragment(
                fragment_id=f"paragraph-{index}",
                text=text[:MAX_FRAGMENT_CHARS],
                locator=DocxParagraphLocator(paragraph_index=index),
            ))
    for index, table in enumerate(document.tables, start=1):
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            fragments.append(TableFragment(
                fragment_id=f"table-{index}",
                rows=rows,
                locator=DocxTableLocator(
                    table_index=index, row_start=1, row_end=len(rows)
                ),
            ))
    if not fragments:
        raise MaterialParseError("DOCX 未提取到可读文本")
    rendered = "\n".join(fragment.text for fragment in fragments)
    completed = datetime.now(timezone.utc)
    quality_flags = _archive_quality_flags(data, "docx")
    message = "浮动文本框、嵌入图片或修订内容未被提取" if quality_flags else ""
    step = ProcessingStep(
        step="parsed", processor="deterministic",
        status="warning" if quality_flags else "succeeded", parser="python-docx",
        input_fingerprint=_fingerprint(data), output_fingerprint=_fingerprint(rendered),
        quality_flags=quality_flags, started_at=started, completed_at=completed,
        message=message,
    )
    return NormalizedMaterial(
        source_id=source_id,
        source_label=source_label,
        summary=f"已从 DOCX 提取 {len(fragments)} 个可定位片段",
        fragments=fragments,
        processing_steps=[step],
        warnings=[message] if message else [],
    )


def parse_xlsx(data: bytes, *, source_id: UUID, source_label: str) -> NormalizedMaterial:
    """以 data_only/read_only 模式按工作表提取 typed table，避免执行公式和加载完整工作簿。"""
    started = datetime.now(timezone.utc)
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise MaterialParseError("XLSX 无法解析") from error
    fragments: list[TableFragment] = []
    cell_count = 0
    row_count = 0
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            first_row: int | None = None
            last_row = 0
            max_columns = 1
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                row_count += 1
                cell_count += len(row)
                if row_count > MAX_XLSX_ROWS or cell_count > MAX_XLSX_CELLS:
                    raise MaterialParseError("XLSX 内容超过解析上限")
                values = [_clean(value) if value is not None else "" for value in row]
                while values and not values[-1]:
                    values.pop()
                if not any(values):
                    continue
                first_row = row_index if first_row is None else first_row
                last_row = row_index
                max_columns = max(max_columns, len(values))
                rows.append(values)
            if rows and first_row is not None:
                fragments.append(TableFragment(
                    fragment_id=f"sheet-{len(fragments) + 1}",
                    rows=rows,
                    locator=XlsxRangeLocator(
                        sheet_name=sheet.title,
                        cell_range=f"A{first_row}:{get_column_letter(max_columns)}{last_row}",
                    ),
                ))
    finally:
        workbook.close()
    if not fragments:
        raise MaterialParseError("XLSX 未提取到可读单元格")
    rendered = "\n".join(fragment.text for fragment in fragments)
    completed = datetime.now(timezone.utc)
    quality_flags = _archive_quality_flags(data, "xlsx")
    messages = []
    if "xlsx_visual_objects_unparsed" in quality_flags:
        messages.append("图表或嵌入图片未被提取")
    if "xlsx_formula_cached_values_only" in quality_flags:
        messages.append("公式仅读取工作簿中的缓存值")
    message = "；".join(messages)
    step = ProcessingStep(
        step="parsed", processor="deterministic",
        status="warning" if quality_flags else "succeeded", parser="openpyxl",
        input_fingerprint=_fingerprint(data), output_fingerprint=_fingerprint(rendered),
        quality_flags=quality_flags, started_at=started, completed_at=completed,
        message=message,
    )
    return NormalizedMaterial(
        source_id=source_id,
        source_label=source_label,
        summary=f"已从 XLSX 提取 {len(fragments)} 个工作表片段",
        fragments=fragments,
        processing_steps=[step],
        warnings=[message] if message else [],
    )


def _pptx_shape_reading_order(shape: object) -> tuple[int, int]:
    """按视觉阅读顺序排序：先纵坐标（top）后横坐标（left）。"""
    top = getattr(shape, "top", None)
    left = getattr(shape, "left", None)
    return (int(top) if top is not None else 0, int(left) if left is not None else 0)


def parse_pptx(data: bytes, *, source_id: UUID, source_label: str) -> NormalizedMaterial:
    """按幻灯片提取 PPTX 形状文本、表格与演讲者备注；不做 OCR。"""
    started = datetime.now(timezone.utc)
    try:
        presentation = Presentation(BytesIO(data))
        slides = list(presentation.slides)
    except (PackageNotFoundError, KeyError, ValueError, OSError) as error:
        raise MaterialParseError("PPTX 无法解析") from error
    fragments: list[TextFragment | TableFragment] = []
    for slide_number, slide in enumerate(slides, start=1):
        ordered_shapes = sorted(slide.shapes, key=_pptx_shape_reading_order)
        slide_has_content = False
        pending_text: list[str] = []

        def _flush_pending_text() -> None:
            nonlocal slide_has_content
            if not pending_text:
                return
            fragments.append(TextFragment(
                fragment_id=f"slide-{slide_number}-text-{len(fragments) + 1}",
                text="\n".join(pending_text)[:MAX_FRAGMENT_CHARS],
                locator=PptxSlideLocator(slide=slide_number),
            ))
            pending_text.clear()
            slide_has_content = True

        for shape in ordered_shapes:
            if getattr(shape, "has_table", False):
                rows = [
                    [_clean(cell.text) for cell in row.cells]
                    for row in shape.table.rows
                ]
                if any(any(cell for cell in row) for row in rows):
                    _flush_pending_text()
                    fragments.append(TableFragment(
                        fragment_id=f"slide-{slide_number}-table-{len(fragments) + 1}",
                        rows=rows,
                        locator=PptxSlideLocator(slide=slide_number),
                    ))
                    slide_has_content = True
                continue
            if getattr(shape, "has_text_frame", False):
                shape_text = _clean(shape.text_frame.text)
                if shape_text:
                    pending_text.append(shape_text)
        _flush_pending_text()

        if getattr(slide, "has_notes_slide", False):
            notes_text = _clean(slide.notes_slide.notes_text_frame.text)
            if notes_text:
                fragments.append(TextFragment(
                    fragment_id=f"slide-{slide_number}-notes",
                    text=f"备注：{notes_text}"[:MAX_FRAGMENT_CHARS],
                    locator=PptxSlideLocator(slide=slide_number),
                ))
                slide_has_content = True

        if not slide_has_content:
            fragments.append(TextFragment(
                fragment_id=f"slide-{slide_number}-placeholder",
                text="（本页为图片内容，未提取文字）",
                locator=PptxSlideLocator(slide=slide_number),
            ))
    if not fragments:
        raise MaterialParseError("PPTX 未提取到可读文本")
    rendered = "\n".join(fragment.text for fragment in fragments)
    completed = datetime.now(timezone.utc)
    quality_flags = _archive_quality_flags(data, "pptx")
    message = "图表、嵌入图片或 OLE 对象未被提取" if quality_flags else ""
    step = ProcessingStep(
        step="parsed", processor="deterministic",
        status="warning" if quality_flags else "succeeded", parser="python-pptx",
        input_fingerprint=_fingerprint(data), output_fingerprint=_fingerprint(rendered),
        quality_flags=quality_flags, started_at=started, completed_at=completed,
        message=message,
    )
    return NormalizedMaterial(
        source_id=source_id,
        source_label=source_label,
        summary=f"已从 PPTX 提取 {len(fragments)} 个逐页片段",
        fragments=fragments,
        processing_steps=[step],
        warnings=[message] if message else [],
    )
