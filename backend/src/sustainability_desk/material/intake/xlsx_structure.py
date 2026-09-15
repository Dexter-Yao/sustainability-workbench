# ABOUTME: XLSX 的无截断结构解析器，按工作表与绝对行保留原始单元格事实和对象覆盖率。
# ABOUTME: 解析器不执行公式、不聚合整张工作表，也不把公式缓存值伪装为用户常量。
# ABOUTME(en): Untruncated XLSX structural parser, keeping raw cell facts per worksheet and absolute row.
# ABOUTME(en): It evaluates no formulas, aggregates no whole sheet, and never passes cached values off as constants.
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
from numbers import Real
from typing import Any
from uuid import UUID
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.worksheet.worksheet import Worksheet

from sustainability_desk.material.intake.parse_coverage import (
    CapabilityCoverage,
    ParseCoverage,
    ParseGap,
)
from sustainability_desk.material.intake.parsed_material import (
    ParsedCell,
    ParsedChartNode,
    ParsedDocumentNode,
    ParsedImageNode,
    ParsedMaterial,
    ParsedNamedRangeNode,
    ParsedNode,
    ParsedSheetNode,
    ParsedSheetRowNode,
    ParsedTableNode,
    build_parsed_material,
    build_parsed_node,
    stable_parsed_node_id,
)
from sustainability_desk.material.intake.parsed_material_locators import (
    DocumentRootLocator,
    EmbeddedImageLocator,
    XlsxCellLocator,
    XlsxChartLocator,
    XlsxNamedRangeLocator,
    XlsxRowLocator,
    XlsxSheetLocator,
    XlsxTableLocator,
)

MAX_XLSX_ROWS = 20_000
MAX_XLSX_CELLS = 100_000
PARSER_PROFILE_ID = "xlsx-structural@1"
NORMALIZATION_PROFILE_ID = "source-faithful@1"
PARSER_FINGERPRINT = sha256(
    (
        f"{PARSER_PROFILE_ID}|openpyxl:{openpyxl.__version__}|"
        "parsed-material-contract@1"
    ).encode("utf-8")
).hexdigest()


class XlsxStructureParseError(ValueError):
    """XLSX 已通过文件准入，但结构无法在资源边界内完整解析。"""


@dataclass(frozen=True)
class _MergedCellFacts:
    anchor: str
    cell_range: str
    row_span: int
    column_span: int


@dataclass(frozen=True)
class _SheetExtent:
    max_row: int
    max_column: int

    @property
    def logical_cell_count(self) -> int:
        return self.max_row * self.max_column


@dataclass
class _ParseStats:
    row_count: int = 0
    cell_count: int = 0
    nonblank_cell_count: int = 0
    text_character_count: int = 0
    formula_count: int = 0
    merge_count: int = 0
    hidden_row_count: int = 0
    hidden_column_count: int = 0
    defined_table_count: int = 0
    named_range_count: int = 0
    chart_count: int = 0
    image_node_count: int = 0
    image_extracted_count: int = 0
    comment_count: int = 0


@dataclass(frozen=True)
class _ArchiveObjectCounts:
    defined_tables: int
    charts: int
    images: int


def _archive_object_counts(data: bytes) -> _ArchiveObjectCounts:
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = tuple(archive.namelist())
    except BadZipFile as error:
        raise XlsxStructureParseError("XLSX 容器已损坏") from error
    return _ArchiveObjectCounts(
        defined_tables=sum(
            name.startswith("xl/tables/") and name.endswith(".xml")
            for name in names
        ),
        charts=sum(
            name.startswith("xl/charts/") and name.endswith(".xml")
            for name in names
        ),
        images=sum(
            name.startswith("xl/media/") and not name.endswith("/")
            for name in names
        ),
    )


def _sheet_extent(worksheet: Worksheet) -> _SheetExtent:
    physical_positions = tuple(worksheet._cells)
    max_row = max((row for row, _ in physical_positions), default=0)
    max_column = max((column for _, column in physical_positions), default=0)

    max_row = max(
        max_row,
        max(
            (
                int(index)
                for index in worksheet.row_dimensions
                if isinstance(index, int) or str(index).isdigit()
            ),
            default=0,
        ),
    )
    for merged_range in worksheet.merged_cells.ranges:
        max_row = max(max_row, merged_range.max_row)
        max_column = max(max_column, merged_range.max_col)
    for table in worksheet.tables.values():
        _, _, table_max_column, table_max_row = range_boundaries(table.ref)
        max_row = max(max_row, table_max_row)
        max_column = max(max_column, table_max_column)

    if max_column > 0 and max_row == 0:
        max_row = 1
    return _SheetExtent(max_row=max_row, max_column=max_column)


def _validate_resource_bounds(
    extents: tuple[_SheetExtent, ...],
) -> None:
    row_count = sum(extent.max_row for extent in extents)
    cell_count = sum(extent.logical_cell_count for extent in extents)
    if row_count > MAX_XLSX_ROWS:
        raise XlsxStructureParseError(
            f"XLSX 原始行数超过 {MAX_XLSX_ROWS} 行解析上限"
        )
    if cell_count > MAX_XLSX_CELLS:
        raise XlsxStructureParseError(
            f"XLSX 原始单元格位置超过 {MAX_XLSX_CELLS} 个解析上限"
        )


def _hidden_columns(
    worksheet: Worksheet,
    *,
    effective_max_column: int,
) -> set[int]:
    hidden: set[int] = set()
    for key, dimension in worksheet.column_dimensions.items():
        if not dimension.hidden:
            continue
        start = dimension.min
        end = dimension.max
        if start is None:
            start = column_index_from_string(key)
        if end is None:
            end = start
        if start > effective_max_column:
            continue
        hidden.update(range(start, min(end, effective_max_column) + 1))
    return hidden


def _merged_cell_facts(
    worksheet: Worksheet,
) -> dict[tuple[int, int], _MergedCellFacts]:
    facts: dict[tuple[int, int], _MergedCellFacts] = {}
    for merged_range in worksheet.merged_cells.ranges:
        anchor = f"{get_column_letter(merged_range.min_col)}{merged_range.min_row}"
        cell_range = str(merged_range)
        for row_index in range(merged_range.min_row, merged_range.max_row + 1):
            for column_index in range(
                merged_range.min_col, merged_range.max_col + 1
            ):
                is_anchor = (
                    row_index == merged_range.min_row
                    and column_index == merged_range.min_col
                )
                facts[(row_index, column_index)] = _MergedCellFacts(
                    anchor=anchor,
                    cell_range=cell_range,
                    row_span=(
                        merged_range.max_row - merged_range.min_row + 1
                        if is_anchor
                        else 1
                    ),
                    column_span=(
                        merged_range.max_col - merged_range.min_col + 1
                        if is_anchor
                        else 1
                    ),
                )
    return facts


def _formula_text(value: object) -> str:
    native_text = getattr(value, "text", None)
    if isinstance(native_text, str) and native_text:
        return native_text
    return str(value)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _cached_scalar(value: object) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _formula_cached_values(
    data: bytes,
    workbook: Any,
    extents: tuple[_SheetExtent, ...],
) -> tuple[dict[str, str | int | float | bool | None], ...]:
    formula_references = tuple(
        {
            cell.coordinate
            for cell in worksheet._cells.values()
            if not isinstance(cell, MergedCell) and cell.data_type == "f"
        }
        for worksheet in workbook.worksheets
    )
    if not any(formula_references):
        return tuple({} for _ in workbook.worksheets)
    try:
        cached_workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as error:
        raise XlsxStructureParseError("XLSX 公式缓存读取失败") from error
    cached_values: list[dict[str, str | int | float | bool | None]] = []
    try:
        for cached_sheet, extent, references in zip(
            cached_workbook.worksheets,
            extents,
            formula_references,
            strict=True,
        ):
            values: dict[str, str | int | float | bool | None] = {}
            if references:
                for row_index, row in enumerate(
                    cached_sheet.iter_rows(
                        min_row=1,
                        max_row=extent.max_row,
                        min_col=1,
                        max_col=extent.max_column,
                    ),
                    start=1,
                ):
                    for column_index, cell in enumerate(row, start=1):
                        reference = f"{get_column_letter(column_index)}{row_index}"
                        if reference in references:
                            values[reference] = _cached_scalar(cell.value)
            cached_values.append(values)
    finally:
        cached_workbook.close()
    return tuple(cached_values)


def _parsed_cell(
    cell: Any,
    *,
    column_index: int,
    physical: bool,
    cached_formula_value: str | int | float | bool | None,
    hidden_row: bool,
    hidden_column: bool,
    merged: _MergedCellFacts | None,
) -> ParsedCell:
    value = cell.value
    native_data_type = str(cell.data_type) if physical else None
    formula: str | None = None
    if native_data_type == "f":
        value_kind = "formula"
        formula = _formula_text(value)
        text = formula
    elif value is None:
        value_kind = "blank"
        text = ""
    elif native_data_type == "e":
        value_kind = "error"
        text = str(value)
    elif isinstance(value, bool):
        value_kind = "boolean"
        text = str(value)
    elif isinstance(value, (datetime, date, time)) or (
        physical and not isinstance(cell, MergedCell) and cell.is_date
    ):
        value_kind = "date"
        text = _cell_text(value)
    elif isinstance(value, Real):
        value_kind = "number"
        text = str(value)
    else:
        value_kind = "text"
        text = str(value)
    return ParsedCell(
        column_index=column_index,
        text=text,
        value_kind=value_kind,
        cell_reference=cell.coordinate,
        formula=formula,
        cached_value=(cached_formula_value if value_kind == "formula" else None),
        native_data_type=native_data_type,
        number_format=(str(cell.number_format) if physical else None),
        merged_anchor=merged.anchor if merged is not None else None,
        merged_range=merged.cell_range if merged is not None else None,
        row_span=merged.row_span if merged is not None else 1,
        column_span=merged.column_span if merged is not None else 1,
        hidden_row=hidden_row,
        hidden_column=hidden_column,
    )


def _chart_title(chart: Any) -> str | None:
    if isinstance(chart.title, str):
        return chart.title
    title = getattr(chart, "title", None)
    rich = getattr(getattr(title, "tx", None), "rich", None)
    paragraphs = getattr(rich, "p", ()) if rich is not None else ()
    pieces: list[str] = []
    for paragraph in paragraphs:
        for run in getattr(paragraph, "r", ()) or ():
            text = getattr(run, "t", None)
            if text:
                pieces.append(str(text))
        for field in getattr(paragraph, "fld", ()) or ():
            text = getattr(field, "t", None)
            if text:
                pieces.append(str(text))
    joined = "".join(pieces)
    return joined or None


def _reference_formula(candidate: object) -> str | None:
    for reference_name in ("numRef", "strRef"):
        reference = getattr(candidate, reference_name, None)
        formula = getattr(reference, "f", None)
        if formula:
            return str(formula)
    return None


def _chart_source_ranges(chart: Any) -> tuple[str, ...]:
    ranges: list[str] = []
    for series in getattr(chart, "ser", ()):
        for attribute in ("tx", "cat", "val", "xVal", "yVal", "bubbleSize"):
            candidate = getattr(series, attribute, None)
            if candidate is None:
                continue
            formula = _reference_formula(candidate)
            if formula and formula not in ranges:
                ranges.append(formula)
    return tuple(ranges)


def _image_anchor_reference(image: Any) -> str:
    anchor = image.anchor
    if isinstance(anchor, str):
        return anchor
    start = getattr(anchor, "_from", None)
    if start is None:
        return "unknown"
    first = f"{get_column_letter(start.col + 1)}{start.row + 1}"
    end = getattr(anchor, "to", None)
    if end is None:
        return first
    last = f"{get_column_letter(end.col + 1)}{end.row + 1}"
    return f"{first}:{last}"


def _image_sha256(image: Any) -> str | None:
    try:
        return sha256(image._data()).hexdigest()
    except (AttributeError, OSError, ValueError):
        return None


def _named_range_definitions(
    workbook: Any,
) -> tuple[tuple[str, str | None, str, int], ...]:
    definitions: list[tuple[str, str | None, str, int]] = []
    occurrences: dict[tuple[str, str | None], int] = {}

    def append(name: str, scope: str | None, definition: str) -> None:
        identity = (name, scope)
        occurrence = occurrences.get(identity, 0) + 1
        occurrences[identity] = occurrence
        definitions.append((name, scope, definition, occurrence))

    for name, definition in workbook.defined_names.items():
        native = getattr(definition, "attr_text", None)
        if native:
            append(str(name), None, str(native))
    for worksheet in workbook.worksheets:
        for name, definition in worksheet.defined_names.items():
            native = getattr(definition, "attr_text", None)
            if native:
                append(str(name), worksheet.title, str(native))
    return tuple(definitions)


def _comment_gaps(
    worksheet: Worksheet,
    *,
    physical_positions: set[tuple[int, int]],
) -> tuple[ParseGap, ...]:
    gaps: list[ParseGap] = []
    for row_index, column_index in sorted(physical_positions):
        cell = worksheet._cells[(row_index, column_index)]
        if isinstance(cell, MergedCell) or cell.comment is None:
            continue
        gaps.append(
            ParseGap(
                code="xlsx_comment_not_parsed",
                message="单元格批注未进入 ParsedNode，仅可使用已解析单元格正文。",
                effect="negative_evidence_blocked",
                locator=XlsxCellLocator(
                    sheet_name=worksheet.title,
                    cell_reference=cell.coordinate,
                ),
            )
        )
    return tuple(gaps)


def _coverage(
    *,
    sheet_count: int,
    stats: _ParseStats,
    archive_counts: _ArchiveObjectCounts,
    gaps: tuple[ParseGap, ...],
) -> ParseCoverage:
    defined_tables_parsed = stats.defined_table_count
    charts_parsed = stats.chart_count
    structural_units = (
        sheet_count
        + stats.row_count
        + stats.cell_count
        + archive_counts.defined_tables
        + stats.named_range_count
        + archive_counts.charts
        + archive_counts.images
        + stats.comment_count
    )
    parsed_units = (
        sheet_count
        + stats.row_count
        + stats.cell_count
        + min(archive_counts.defined_tables, defined_tables_parsed)
        + stats.named_range_count
        + min(archive_counts.charts, charts_parsed)
        + min(archive_counts.images, stats.image_extracted_count)
    )
    capabilities = (
        CapabilityCoverage(
            capability="text",
            source_unit_count=stats.nonblank_cell_count,
            parsed_unit_count=stats.nonblank_cell_count,
        ),
        CapabilityCoverage(
            capability="tables",
            source_unit_count=stats.row_count,
            parsed_unit_count=stats.row_count,
        ),
        CapabilityCoverage(
            capability="defined_tables",
            source_unit_count=archive_counts.defined_tables,
            parsed_unit_count=min(
                archive_counts.defined_tables, defined_tables_parsed
            ),
        ),
        CapabilityCoverage(
            capability="named_ranges",
            source_unit_count=stats.named_range_count,
            parsed_unit_count=stats.named_range_count,
        ),
        CapabilityCoverage(
            capability="charts",
            source_unit_count=archive_counts.charts,
            parsed_unit_count=min(archive_counts.charts, charts_parsed),
        ),
        CapabilityCoverage(
            capability="images",
            source_unit_count=archive_counts.images,
            parsed_unit_count=min(
                archive_counts.images,
                stats.image_extracted_count,
            ),
        ),
        CapabilityCoverage(
            capability="formulas",
            source_unit_count=stats.formula_count,
            parsed_unit_count=stats.formula_count,
        ),
        CapabilityCoverage(
            capability="comments",
            source_unit_count=stats.comment_count,
            parsed_unit_count=0,
        ),
        CapabilityCoverage(
            capability="layout",
            source_unit_count=(
                sheet_count
                + stats.row_count
                + stats.cell_count
                + stats.merge_count
                + stats.hidden_row_count
                + stats.hidden_column_count
            ),
            parsed_unit_count=(
                sheet_count
                + stats.row_count
                + stats.cell_count
                + stats.merge_count
                + stats.hidden_row_count
                + stats.hidden_column_count
            ),
        ),
    )
    incomplete = bool(gaps) or any(not capability.complete for capability in capabilities)
    return ParseCoverage(
        disposition=(
            "incomplete_usable"
            if incomplete
            else "complete_for_declared_capabilities"
        ),
        source_units_total=structural_units,
        source_units_parsed=parsed_units,
        text_characters_parsed=stats.text_character_count,
        table_cells_parsed=stats.cell_count,
        capabilities=capabilities,
        gaps=gaps,
    )


def parse_xlsx_structure(
    data: bytes,
    *,
    source_id: UUID,
    source_label: str,
) -> ParsedMaterial:
    """把 XLSX 原子解析为无截断 ParsedMaterial，不执行公式。"""

    source_sha256 = sha256(data).hexdigest()
    archive_counts = _archive_object_counts(data)
    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise XlsxStructureParseError("XLSX 无法解析") from error

    try:
        extents = tuple(_sheet_extent(sheet) for sheet in workbook.worksheets)
        _validate_resource_bounds(extents)
        cached_formula_values = _formula_cached_values(data, workbook, extents)
        root = build_parsed_node(
            ParsedDocumentNode,
            source_sha256=source_sha256,
            parser_profile_id=PARSER_PROFILE_ID,
            locator=DocumentRootLocator(),
            parent_node_id=None,
            ordinal_path=(1,),
            label=source_label,
        )
        nodes: list[ParsedNode] = [root]
        gaps: list[ParseGap] = []
        stats = _ParseStats()

        for sheet_index, (worksheet, extent, cached_values) in enumerate(
            zip(
                workbook.worksheets,
                extents,
                cached_formula_values,
                strict=True,
            ),
            start=1,
        ):
            row_locators = tuple(
                XlsxRowLocator(
                    sheet_name=worksheet.title,
                    row_index=row_index,
                )
                for row_index in range(1, extent.max_row + 1)
            )
            row_node_ids = tuple(
                stable_parsed_node_id(
                    source_sha256=source_sha256,
                    parser_profile_id=PARSER_PROFILE_ID,
                    locator=locator,
                )
                for locator in row_locators
            )
            sheet = build_parsed_node(
                ParsedSheetNode,
                source_sha256=source_sha256,
                parser_profile_id=PARSER_PROFILE_ID,
                locator=XlsxSheetLocator(
                    sheet_name=worksheet.title,
                    sheet_index=sheet_index,
                ),
                parent_node_id=root.node_id,
                ordinal_path=(1, sheet_index),
                label=worksheet.title,
                sheet_state=worksheet.sheet_state,
                cell_range=(
                    None
                    if extent.max_column == 0
                    else (
                        f"A1:{get_column_letter(extent.max_column)}"
                        f"{extent.max_row}"
                    )
                ),
                row_count=extent.max_row,
                column_count=extent.max_column,
                cell_count=extent.logical_cell_count,
                row_node_ids=row_node_ids,
            )
            nodes.append(sheet)
            physical_positions = set(worksheet._cells)
            hidden_columns = _hidden_columns(
                worksheet,
                effective_max_column=extent.max_column,
            )
            merged_facts = _merged_cell_facts(worksheet)
            stats.merge_count += len(worksheet.merged_cells.ranges)
            stats.hidden_column_count += len(hidden_columns)

            for row_index, row_locator in enumerate(row_locators, start=1):
                row_dimension = worksheet.row_dimensions.get(row_index)
                hidden_row = bool(row_dimension and row_dimension.hidden)
                if hidden_row:
                    stats.hidden_row_count += 1
                cells: list[ParsedCell] = []
                for column_index in range(1, extent.max_column + 1):
                    cell = worksheet.cell(row=row_index, column=column_index)
                    parsed_cell = _parsed_cell(
                        cell,
                        column_index=column_index,
                        physical=(row_index, column_index) in physical_positions,
                        cached_formula_value=cached_values.get(cell.coordinate),
                        hidden_row=hidden_row,
                        hidden_column=column_index in hidden_columns,
                        merged=merged_facts.get((row_index, column_index)),
                    )
                    cells.append(parsed_cell)
                    stats.cell_count += 1
                    stats.text_character_count += len(parsed_cell.text)
                    if parsed_cell.value_kind != "blank":
                        stats.nonblank_cell_count += 1
                    if parsed_cell.value_kind == "formula":
                        stats.formula_count += 1
                row = build_parsed_node(
                    ParsedSheetRowNode,
                    source_sha256=source_sha256,
                    parser_profile_id=PARSER_PROFILE_ID,
                    locator=row_locator,
                    parent_node_id=sheet.node_id,
                    ordinal_path=(1, sheet_index, row_index),
                    row_index=row_index,
                    hidden=hidden_row,
                    height_points=(
                        row_dimension.height if row_dimension is not None else None
                    ),
                    cells=tuple(cells),
                )
                nodes.append(row)
                stats.row_count += 1

            child_ordinal = extent.max_row
            for table in worksheet.tables.values():
                child_ordinal += 1
                min_column, _, max_column, _ = range_boundaries(table.ref)
                nodes.append(
                    build_parsed_node(
                        ParsedTableNode,
                        source_sha256=source_sha256,
                        parser_profile_id=PARSER_PROFILE_ID,
                        locator=XlsxTableLocator(
                            sheet_name=worksheet.title,
                            table_name=table.name,
                            cell_range=table.ref,
                        ),
                        parent_node_id=sheet.node_id,
                        ordinal_path=(1, sheet_index, child_ordinal),
                        column_count=max_column - min_column + 1,
                        header_row_node_ids=(),
                        title=table.displayName,
                    )
                )
                stats.defined_table_count += 1

            for chart_index, chart in enumerate(worksheet._charts, start=1):
                child_ordinal += 1
                title = _chart_title(chart)
                nodes.append(
                    build_parsed_node(
                        ParsedChartNode,
                        source_sha256=source_sha256,
                        parser_profile_id=PARSER_PROFILE_ID,
                        locator=XlsxChartLocator(
                            sheet_name=worksheet.title,
                            chart_index=chart_index,
                            chart_name=title,
                        ),
                        parent_node_id=sheet.node_id,
                        ordinal_path=(1, sheet_index, child_ordinal),
                        chart_type=type(chart).__name__,
                        title=title,
                        source_ranges=_chart_source_ranges(chart),
                    )
                )
                stats.chart_count += 1

            for image_index, image in enumerate(worksheet._images, start=1):
                child_ordinal += 1
                anchor_reference = _image_anchor_reference(image)
                media_sha256 = _image_sha256(image)
                if media_sha256 is None:
                    media_extraction_status = "failed"
                    gap_code = "xlsx_image_resource_extraction_failed"
                    gap_message = (
                        "已定位嵌入图片，但无法提取原始媒体字节。"
                    )
                else:
                    media_extraction_status = "complete"
                    gap_code = "xlsx_image_semantics_unparsed"
                    gap_message = (
                        "已提取嵌入图片原始媒体，但未解析图片中的文字或语义。"
                    )
                locator = EmbeddedImageLocator(
                    container_kind="xlsx",
                    container_ref=f"{worksheet.title}!{anchor_reference}",
                    image_index=image_index,
                )
                nodes.append(
                    build_parsed_node(
                        ParsedImageNode,
                        source_sha256=source_sha256,
                        parser_profile_id=PARSER_PROFILE_ID,
                        locator=locator,
                        parent_node_id=sheet.node_id,
                        ordinal_path=(1, sheet_index, child_ordinal),
                        media_sha256=media_sha256,
                        media_extraction_status=media_extraction_status,
                    )
                )
                gaps.append(
                    ParseGap(
                        code=gap_code,
                        message=gap_message,
                        effect="negative_evidence_blocked",
                        locator=locator,
                    )
                )
                stats.image_node_count += 1
                if media_extraction_status == "complete":
                    stats.image_extracted_count += 1

            comment_gaps = _comment_gaps(
                worksheet,
                physical_positions=physical_positions,
            )
            gaps.extend(comment_gaps)
            stats.comment_count += len(comment_gaps)

        for ordinal, (
            name,
            scope_sheet_name,
            definition,
            definition_index,
        ) in enumerate(_named_range_definitions(workbook), start=1):
            nodes.append(
                build_parsed_node(
                    ParsedNamedRangeNode,
                    source_sha256=source_sha256,
                    parser_profile_id=PARSER_PROFILE_ID,
                    locator=XlsxNamedRangeLocator(
                        name=name,
                        scope_sheet_name=scope_sheet_name,
                        definition_index=definition_index,
                    ),
                    parent_node_id=root.node_id,
                    ordinal_path=(1, len(workbook.worksheets) + ordinal),
                    name=name,
                    scope_sheet_name=scope_sheet_name,
                    definition=definition,
                )
            )
            stats.named_range_count += 1

        if stats.defined_table_count < archive_counts.defined_tables:
            gaps.append(
                ParseGap(
                    code="xlsx_defined_table_structure_unparsed",
                    message="部分 XLSX 定义表未能进入结构解析树。",
                    effect="negative_evidence_blocked",
                    locator=DocumentRootLocator(),
                )
            )
        if stats.chart_count < archive_counts.charts:
            gaps.append(
                ParseGap(
                    code="xlsx_chart_structure_unparsed",
                    message="部分 XLSX 图表未能进入结构解析树。",
                    effect="negative_evidence_blocked",
                    locator=DocumentRootLocator(),
                )
            )
        if stats.image_node_count < archive_counts.images:
            gaps.append(
                ParseGap(
                    code="xlsx_image_structure_unparsed",
                    message="部分 XLSX 媒体对象未能获得可定位的图片节点。",
                    effect="negative_evidence_blocked",
                    locator=DocumentRootLocator(),
                )
            )

        coverage = _coverage(
            sheet_count=len(workbook.worksheets),
            stats=stats,
            archive_counts=archive_counts,
            gaps=tuple(gaps),
        )
        return build_parsed_material(
            source_id=source_id,
            source_sha256=source_sha256,
            document_kind="xlsx",
            parser_profile_id=PARSER_PROFILE_ID,
            parser_fingerprint=PARSER_FINGERPRINT,
            normalization_profile_id=NORMALIZATION_PROFILE_ID,
            root_node_id=root.node_id,
            nodes=tuple(sorted(nodes, key=lambda node: node.ordinal_path)),
            coverage=coverage,
        )
    except XlsxStructureParseError:
        raise
    except Exception as error:
        raise XlsxStructureParseError("XLSX 结构解析失败") from error
    finally:
        workbook.close()
