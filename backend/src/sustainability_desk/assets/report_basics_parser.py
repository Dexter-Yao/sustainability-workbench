# ABOUTME: 基础资料工作簿严格解析器：按隐藏键列（kind:key）逐行解析字段、报告级题、披露准则与附录联系。
# ABOUTME: 值按合同字段类型 parse-first（年份/日期/数值/枚举/邮箱），整簿有效才返回 typed 结果，禁止部分导入。
# ABOUTME(en): Strict parser for the report basics workbook: fields, report-level questions, disclosure requirements and
# ABOUTME(en): appendix contacts by hidden key column. Values parse-first by contract field type; whole-workbook only.
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field as dataclass_field
from datetime import date, datetime

import openpyxl
from pydantic import ValidationError

from sustainability_desk.assets.report_basics_template import (
    BOOLEAN_NO,
    BOOLEAN_YES,
    REPORT_BASICS_ROWS,
    REPORT_BASICS_SHEET,
    ReportBasicsRow,
    report_basics_context_fingerprint,
    report_basics_row_key,
)
from sustainability_desk.assets.workbook_safety import (
    validate_workbook_dimensions,
    validate_xlsx_container,
)
from sustainability_desk.contract.models import (
    AppendixPackage,
    DisclosureProfile,
    ExternalAssuranceReport,
    ReaderFeedbackContactInformation,
    Report,
)
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.contract.report_values import is_valid_contact_email
from sustainability_desk.contract.stored_report_state import StoredIntakeAnswer
from sustainability_desk.contract.structured_inputs import (
    StructuredInputCellError,
    StructuredInputContext,
    StructuredInputWorkbookError,
    context_metadata_values,
    validate_metadata_sheet,
)

_DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")


@dataclass
class ReportBasicsImport:
    """基础资料工作簿的整批解析结果；键集合即模板范围，导入为整批替换。"""

    fields: dict[str, str | int | float | None] = dataclass_field(default_factory=dict)
    intake_answers: dict[str, StoredIntakeAnswer] = dataclass_field(default_factory=dict)
    intake_keys: frozenset[str] = frozenset()
    disclosure_profile: DisclosureProfile | None = None
    appendix_package: AppendixPackage | None = None


def _text(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return "" if value is None else str(value).strip()


def _cell_error(*, code: str, row: int, message: str) -> StructuredInputCellError:
    return StructuredInputCellError(
        code=code,
        sheet=REPORT_BASICS_SHEET,
        row=row,
        column="F",
        message=message,
    )


def _parse_field_value(
    *,
    report: Report,
    key: str,
    raw: object,
    row_number: int,
    errors: list[StructuredInputCellError],
) -> str | int | float | None:
    field = report.fields[key]
    text = _text(raw)
    if not text:
        return None
    label = field.label
    if field.type == "enum":
        options = [str(option) for option in (field.options or [])]
        if text not in options:
            errors.append(
                _cell_error(
                    code="report_basics_enum_invalid",
                    row=row_number,
                    message=f"「{label}」必须是模板提供的选项之一",
                )
            )
            return None
        return text
    if field.type == "month":
        # 合同口径是 YYYY-MM（与前端 <input type="month"> 同源）；Excel 可能把年月
        # 自动转成日期单元格，此处按原始值归一化，纯数字月份视为格式错误。
        if isinstance(raw, (datetime, date)):
            return raw.strftime("%Y-%m")
        if re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", text):
            return text
        if re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])-\d{2}", text):
            return text[:7]
        errors.append(
            _cell_error(
                code="report_basics_month_invalid",
                row=row_number,
                message=f"「{label}」必须是 YYYY-MM 格式的年月",
            )
        )
        return None
    if field.type in {"year", "number", "percent"}:
        try:
            numeric = float(text)
        except ValueError:
            errors.append(
                _cell_error(
                    code="report_basics_number_invalid",
                    row=row_number,
                    message=f"「{label}」必须是数字",
                )
            )
            return None
        if field.type == "year":
            year = int(numeric)
            if year != numeric or not 1900 <= year <= 9999:
                errors.append(
                    _cell_error(
                        code="report_basics_year_invalid",
                        row=row_number,
                        message=f"「{label}」必须是四位年份",
                    )
                )
                return None
            return year
        return int(numeric) if numeric == int(numeric) else numeric
    if field.type == "date":
        if not _DATE_PATTERN.fullmatch(text):
            errors.append(
                _cell_error(
                    code="report_basics_date_invalid",
                    row=row_number,
                    message=f"「{label}」必须是 YYYY-MM-DD 格式的日期",
                )
            )
            return None
        return text
    if field.type == "email":
        if not is_valid_contact_email(text):
            errors.append(
                _cell_error(
                    code="report_basics_email_invalid",
                    row=row_number,
                    message=f"「{label}」不是有效邮箱地址",
                )
            )
            return None
        return text
    return text


def _validate_against_field_contract(
    *,
    report: Report,
    parsed_fields: dict[str, str | int | float | None],
    locations: dict[str, int],
    errors: list[StructuredInputCellError],
) -> None:
    """写入前用合同 Field 自身的类型校验兜底（parse-first 的最后一道闸）。

    解析分支与合同校验若有口径漂移（如 month=YYYY-MM），必须在导入边界 fail-loud，
    不允许宽 state 落库后由下游装配（/api/plan、生成门禁）以 422 暴露。
    """

    for key, value in parsed_fields.items():
        if value is None:
            continue
        contract_field = report.fields[key]
        try:
            type(contract_field).model_validate(
                {**contract_field.model_dump(mode="json"), "value": value}
            )
        except ValidationError as error:
            errors.append(
                _cell_error(
                    code="report_basics_field_contract_invalid",
                    row=locations.get(f"field:{key}", 1),
                    message=(
                        f"「{contract_field.label}」不符合字段合同："
                        f"{error.errors()[0]['msg']}"
                    ),
                )
            )


def parse_report_basics_workbook(
    content: bytes,
    *,
    report: Report,
    context: StructuredInputContext,
) -> ReportBasicsImport:
    """严格解析基础资料工作簿；成功即返回可整批替换的 typed 结果。"""

    validate_xlsx_container(content)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as error:  # noqa: BLE001 — 统一投影为 typed 工作簿错误
        raise StructuredInputWorkbookError(
            [
                StructuredInputCellError(
                    code="workbook_invalid",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"无法读取基础资料工作簿：{error}",
                )
            ]
        ) from error
    validate_workbook_dimensions(workbook)

    fingerprint = report_basics_context_fingerprint(report, context)
    errors = validate_metadata_sheet(
        workbook,
        context_metadata_values(
            input_kind="report_basics",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    rows_by_key: dict[str, ReportBasicsRow] = {
        report_basics_row_key(row): row for row in REPORT_BASICS_ROWS
    }

    if REPORT_BASICS_SHEET not in workbook.sheetnames:
        errors.append(
            StructuredInputCellError(
                code="report_basics_sheet_missing",
                sheet=REPORT_BASICS_SHEET,
                row=1,
                column="A",
                message=f"工作簿缺少工作表：{REPORT_BASICS_SHEET}",
            )
        )
        raise StructuredInputWorkbookError(errors)

    raw_values: dict[str, object] = {}
    locations: dict[str, int] = {}
    seen: set[str] = set()
    worksheet = workbook[REPORT_BASICS_SHEET]
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = tuple(row) + (None,) * max(0, 7 - len(row))
        raw_key = _text(values[0])
        if not raw_key or raw_key == "键":
            continue
        if raw_key not in rows_by_key:
            errors.append(
                StructuredInputCellError(
                    code="report_basics_key_unknown",
                    sheet=REPORT_BASICS_SHEET,
                    row=row_number,
                    column="A",
                    message=f"未知的填写项目键：{raw_key}",
                )
            )
            continue
        if raw_key in seen:
            errors.append(
                StructuredInputCellError(
                    code="report_basics_key_duplicate",
                    sheet=REPORT_BASICS_SHEET,
                    row=row_number,
                    column="A",
                    message=f"填写项目重复出现：{raw_key}",
                )
            )
            continue
        seen.add(raw_key)
        raw_values[raw_key] = values[5]
        locations[raw_key] = row_number

    for row_key in rows_by_key:
        if row_key not in seen:
            errors.append(
                StructuredInputCellError(
                    code="report_basics_key_missing",
                    sheet=REPORT_BASICS_SHEET,
                    row=1,
                    column="A",
                    message=f"工作簿缺少填写项目行：{row_key}（请重新下载模板）",
                )
            )

    result = ReportBasicsImport(
        intake_keys=frozenset(
            row.key for row in REPORT_BASICS_ROWS if row.kind == "intake"
        )
    )
    mainland_names = knowledge_package_of(report).manifest.disclosure_basis.mainland_standard_names
    standard_by_name = {
        name: code for code, name in (mainland_names.by_code() if mainland_names else {}).items()
    }
    disclosure_values: dict[str, object] = {}
    appendix_values: dict[str, object] = {}

    for row in REPORT_BASICS_ROWS:
        row_key = report_basics_row_key(row)
        if row_key not in seen:
            continue
        row_number = locations[row_key]
        raw = raw_values[row_key]
        text = _text(raw)
        if row.kind == "field":
            result.fields[row.key] = _parse_field_value(
                report=report,
                key=row.key,
                raw=raw,
                row_number=row_number,
                errors=errors,
            )
        elif row.kind == "intake":
            item = next(item for item in report.intakeItems if item.key == row.key)
            if text:
                if item.maxChars and len(text) > item.maxChars:
                    errors.append(
                        _cell_error(
                            code="report_basics_answer_too_long",
                            row=row_number,
                            message=f"「{item.prompt}」超过 {item.maxChars} 字上限",
                        )
                    )
                    continue
                if item.minChars and len(text) < item.minChars:
                    errors.append(
                        _cell_error(
                            code="report_basics_answer_too_short",
                            row=row_number,
                            message=(
                                f"「{item.prompt}」不足 {item.minChars} 字；"
                                "留空表示暂不填写"
                            ),
                        )
                    )
                    continue
                result.intake_answers[row.key] = StoredIntakeAnswer(answer=text)
        elif row.kind == "disclosure":
            if row.key == "mainlandStandard":
                if not text:
                    errors.append(
                        _cell_error(
                            code="report_basics_standard_missing",
                            row=row_number,
                            message="请选择大陆准则",
                        )
                    )
                elif text not in standard_by_name:
                    errors.append(
                        _cell_error(
                            code="report_basics_standard_invalid",
                            row=row_number,
                            message="大陆准则必须是模板提供的选项之一",
                        )
                    )
                else:
                    disclosure_values["mainlandStandard"] = standard_by_name[text]
        elif row.kind == "appendix":
            if row.key == "externalAssuranceReport.isIncluded":
                if text and text not in {BOOLEAN_YES, BOOLEAN_NO}:
                    errors.append(
                        _cell_error(
                            code="report_basics_boolean_invalid",
                            row=row_number,
                            message=f"「{row.label}」只能选择 {BOOLEAN_YES} 或 {BOOLEAN_NO}",
                        )
                    )
                else:
                    appendix_values["isIncluded"] = text == BOOLEAN_YES
            elif row.key == "readerFeedbackContactInformation.email":
                if text and not is_valid_contact_email(text):
                    errors.append(
                        _cell_error(
                            code="report_basics_email_invalid",
                            row=row_number,
                            message="读者反馈邮箱不是有效邮箱地址",
                        )
                    )
                else:
                    appendix_values["email"] = text or None
            else:
                appendix_values[row.key.split(".", 1)[1]] = text or None

    _validate_against_field_contract(
        report=report,
        parsed_fields=result.fields,
        locations=locations,
        errors=errors,
    )
    if errors:
        raise StructuredInputWorkbookError(errors)

    existing_profile = report.disclosureProfile or DisclosureProfile()
    result.disclosure_profile = DisclosureProfile(
        mainlandStandard=disclosure_values.get(
            "mainlandStandard", existing_profile.mainlandStandard
        ),
        # 港交所守则与其他参考文件当前不在模板收集面内：工作簿无从表达它们，
        # 导入必须原样保留报告现值，否则整批替换会静默清空既有报告的编制依据。
        includesHongKongExchangeGuide=existing_profile.includesHongKongExchangeGuide,
        additionalDisclosureReferences=list(
            existing_profile.additionalDisclosureReferences
        ),
    )
    result.appendix_package = AppendixPackage(
        externalAssuranceReport=ExternalAssuranceReport(
            isIncluded=bool(appendix_values.get("isIncluded", False)),
            fileLabel=appendix_values.get("fileLabel"),
        ),
        readerFeedbackContactInformation=ReaderFeedbackContactInformation(
            address=appendix_values.get("address"),
            email=appendix_values.get("email"),
            phone=appendix_values.get("phone"),
        ),
    )
    return result
