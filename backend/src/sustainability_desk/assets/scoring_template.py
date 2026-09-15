# ABOUTME: 动态双重重要性评分模板，按当前报告配置投影官方议题与评分尺度为可填写 xlsx。
# ABOUTME: 模板不拥有议题或评分规则；所有行、适用性和数值边界都从既有合同实时派生。
# ABOUTME(en): Dynamic double materiality scoring template, projecting official topics and the score scale into xlsx.
# ABOUTME(en): The template owns neither topics nor scoring rules; rows, applicability and bounds derive from contracts.
from __future__ import annotations

import io
from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.contract.loader import load_package_contract
from sustainability_desk.contract.models import AssessmentScoreScale, Report
from sustainability_desk.contract.structured_inputs import (
    StructuredInputContext,
    append_metadata_sheet,
    assessment_context_fingerprint,
    context_metadata_values,
)
from sustainability_desk.contract.topic_registry import applicable_scoring_topics, load_topic_contract

_BACKEND = Path(__file__).resolve().parents[3]
_ROOT = _BACKEND.parent

_PINE = "2C6E49"
_PINE_SUBTLE = "E7F0E9"
_SURFACE = "F6F7F5"
_BORDER = "E2E6E1"
_BORDER_STRONG = "C7CCC6"
_FOREGROUND = "19211D"
_MUTED = "677069"


def _score_scale(report: Report) -> AssessmentScoreScale:
    scale = (
        report.assessmentScoreScale
        or load_package_contract(knowledge_package_of(report)).assessmentScoreScale
    )
    if scale is None:
        raise ValueError("报告模板未声明双重重要性评分尺度")
    return scale


def _score_text(value: float) -> str:
    return f"{value:g}"


def _decimal_places(step: float) -> int:
    fraction = str(step).split(".")[1]
    return len(fraction)


def _add_score_validation(
    worksheet,
    *,
    column: str,
    rows: Iterable[int],
    scale: AssessmentScoreScale,
) -> None:
    rows = tuple(rows)
    if not rows:
        return
    first_row = rows[0]
    # Excel MOD 直接取浮点步长会误拒合法评分（如 4.3 对 0.1）；先 ROUND 到整数空间再做整数 MOD，
    # 与解析端 parse 容错语义一致。
    factor = 10 ** _decimal_places(scale.multipleOf)
    step_units = round(scale.multipleOf * factor)
    cell = f"{column}{first_row}"
    formula = (
        f"AND({cell}>{_score_text(scale.minimumExclusive)},"
        f"{cell}<={_score_text(scale.maximum)},"
        f"ABS({cell}*{factor}-ROUND({cell}*{factor},0))<0.000000001,"
        f"MOD(ROUND({cell}*{factor},0),{_score_text(step_units)})=0)"
    )
    validation = DataValidation(type="custom", formula1=formula, allow_blank=False)
    validation.error = f"请填写大于 {_score_text(scale.minimumExclusive)}、不大于 {_score_text(scale.maximum)}，且按 {_score_text(scale.multipleOf)} 递增的数值。"
    validation.errorTitle = "评分不符合要求"
    validation.prompt = "仅可填写符合当前评分尺度的数字。"
    validation.promptTitle = "填写评分"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    worksheet.add_data_validation(validation)
    for row in rows:
        validation.add(f"{column}{row}")



SCORING_SHEET_TITLE = "重要性评分表"


def create_scoring_template(
    report: Report,
    *,
    context: StructuredInputContext,
) -> bytes:
    """按当前报告配置生成可填写、可回读的重要性评分表。"""
    fingerprint = assessment_context_fingerprint(report, context)
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_scoring_sheet(workbook, report)
    append_metadata_sheet(
        workbook,
        context_metadata_values(
            input_kind="assessment",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def append_scoring_sheet(workbook: Workbook, report: Report) -> None:
    """向工作簿追加评分业务工作表；元数据页由调用方按所属工作簿种类写入。"""
    scale = _score_scale(report)
    topics = applicable_scoring_topics(report)

    worksheet = workbook.create_sheet(SCORING_SHEET_TITLE)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.45
    worksheet.page_margins.bottom = 0.45
    worksheet.sheet_properties.tabColor = _PINE

    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 18
    worksheet.row_dimensions[1].height = 54
    worksheet.row_dimensions[2].height = 28
    worksheet.row_dimensions[4].height = 20
    worksheet.row_dimensions[5].height = 34
    worksheet.row_dimensions[6].height = 28
    worksheet.row_dimensions[7].height = 8
    worksheet.row_dimensions[8].height = 25

    worksheet.merge_cells("B2:C2")
    worksheet["B2"] = "重要性评分表"
    worksheet["B2"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=_FOREGROUND)
    worksheet["B2"].alignment = Alignment(vertical="center")

    worksheet.merge_cells("B4:C4")
    worksheet["B4"] = f"当前配置：{len(topics)} 个适用议题"
    worksheet["B4"].font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
    worksheet["B4"].alignment = Alignment(vertical="center")

    worksheet.merge_cells("B5:C5")
    worksheet["B5"] = (
        f"请仅填写浅绿色单元格。评分范围：({_score_text(scale.minimumExclusive)}, {_score_text(scale.maximum)}]；"
        f"步长：{_score_text(scale.multipleOf)}。"
    )
    worksheet["B5"].font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
    worksheet["B5"].alignment = Alignment(vertical="center", wrap_text=True)

    # 两个维度的释义与网页同源于合同评分尺度：只有列头「财务重要性/影响重要性」时，
    # 离线填表的人拿不到判断依据。Excel 无折叠形态，两段顺序拼接。
    worksheet.merge_cells("B6:C6")
    worksheet["B6"] = "　".join(
        part
        for part in (
            f"财务重要性：{scale.financialMaterialityDefinition}",
            scale.financialMaterialityExplanation,
            f"影响重要性：{scale.impactMaterialityDefinition}",
            scale.impactMaterialityExplanation,
        )
        if part
    )
    worksheet["B6"].font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
    worksheet["B6"].alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[6].height = 30


    header_row = 8
    worksheet.append([])
    worksheet.cell(header_row, 1, "议题名称")
    worksheet.cell(header_row, 2, "财务重要性")
    worksheet.cell(header_row, 3, "影响重要性")
    header = worksheet[f"A{header_row}:C{header_row}"]
    for row in header:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=_PINE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                top=Side(style="medium", color=_PINE),
                bottom=Side(style="medium", color=_PINE),
            )

    topic_rows: list[int] = []
    current_dimension: str | None = None
    dimension_label = load_topic_contract(knowledge_package_of(report)).dimension_label
    row_number = header_row + 1
    for topic in topics:
        if topic.dimension != current_dimension:
            worksheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=3)
            cell = worksheet.cell(row_number, 1, dimension_label(topic.dimension))
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=_FOREGROUND)
            cell.fill = PatternFill("solid", fgColor=_SURFACE)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(top=Side(style="thin", color=_BORDER_STRONG))
            worksheet.row_dimensions[row_number].height = 21
            current_dimension = topic.dimension
            row_number += 1

        worksheet.cell(row_number, 1, topic.name)
        for column in range(1, 4):
            cell = worksheet.cell(row_number, column)
            cell.font = Font(name="Microsoft YaHei", size=10, color=_FOREGROUND)
            cell.alignment = Alignment(horizontal="left" if column == 1 else "center", vertical="center")
            cell.border = Border(bottom=Side(style="thin", color=_BORDER))
        for column in range(2, 4):
            cell = worksheet.cell(row_number, column)
            cell.fill = PatternFill("solid", fgColor=_PINE_SUBTLE)
            cell.number_format = "0.0"
        worksheet.row_dimensions[row_number].height = 22
        topic_rows.append(row_number)
        row_number += 1

    _add_score_validation(worksheet, column="B", rows=topic_rows, scale=scale)
    _add_score_validation(worksheet, column="C", rows=topic_rows, scale=scale)
    worksheet.freeze_panes = "A9"
    worksheet.print_area = f"A1:C{row_number - 1}"
