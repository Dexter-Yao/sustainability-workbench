# ABOUTME: 议题引导问题工作簿严格解析器，按隐藏键列的行协议逐行累计题目、选项与字数错误。
# ABOUTME: 仅整簿全部有效时返回答案字典（整批替换语义，允许空答案）；禁止部分导入。
# ABOUTME(en): Strict parser for the topic guidance question workbook, accumulating question, option and length errors
# ABOUTME(en): by the hidden key column row protocol. Returns answers only when the whole workbook is valid.
from __future__ import annotations

import io

import openpyxl

from sustainability_desk.assets.topic_questions_template import (
    GROUP_ROW_SUFFIX,
    MULTI_SELECT_NO,
    MULTI_SELECT_YES,
    OPTION_ROW_SUFFIX,
)
from sustainability_desk.assets.workbook_safety import (
    validate_workbook_dimensions,
    validate_xlsx_container,
)
from sustainability_desk.contract.models import IntakeItem, Report
from sustainability_desk.contract.stored_report_state import StoredIntakeAnswer
from sustainability_desk.contract.structured_inputs import (
    STRUCTURED_INPUT_METADATA_SHEET,
    StructuredInputCellError,
    StructuredInputContext,
    StructuredInputWorkbookError,
    context_metadata_values,
    topic_question_items,
    topic_questions_context_fingerprint,
    validate_metadata_sheet,
)


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _cell_error(
    *, code: str, sheet: str, row: int, column: str, message: str
) -> StructuredInputCellError:
    return StructuredInputCellError(
        code=code, sheet=sheet, row=row, column=column, message=message
    )


def parse_topic_questions_workbook(
    content: bytes,
    *,
    report: Report,
    context: StructuredInputContext,
    allowed_report_section_ids: frozenset[str] | None = None,
) -> dict[str, StoredIntakeAnswer]:
    """严格解析议题引导问题工作簿，返回模板范围内的整批答案（空答案的键不出现）。"""

    validate_xlsx_container(content)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as error:  # noqa: BLE001 — 统一投影为 typed 工作簿错误
        raise StructuredInputWorkbookError(
            [
                _cell_error(
                    code="workbook_invalid",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"无法读取议题信息工作簿：{error}",
                )
            ]
        ) from error
    validate_workbook_dimensions(workbook)

    fingerprint = topic_questions_context_fingerprint(
        report,
        context,
        allowed_report_section_ids=allowed_report_section_ids,
    )
    errors = validate_metadata_sheet(
        workbook,
        context_metadata_values(
            input_kind="topic_questions",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    catalog = topic_question_items(
        report,
        allowed_report_section_ids=allowed_report_section_ids,
    )
    catalog_by_key: dict[str, IntakeItem] = {item.key: item for item in catalog}

    raw_answers: dict[str, str] = {}
    raw_supplements: dict[str, str] = {}
    multi_selected: dict[str, list[str]] = {}
    locations: dict[str, tuple[str, int]] = {}
    seen_keys: set[str] = set()

    for sheet_name in workbook.sheetnames:
        if sheet_name in {STRUCTURED_INPUT_METADATA_SHEET, "填写说明"}:
            continue
        worksheet = workbook[sheet_name]
        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True), start=1
        ):
            values = tuple(row) + (None,) * max(0, 7 - len(row))
            raw_key = _text(values[0])
            if not raw_key or raw_key == "键":
                continue
            if raw_key.endswith(GROUP_ROW_SUFFIX):
                continue
            if raw_key.endswith(OPTION_ROW_SUFFIX):
                key = raw_key[: -len(OPTION_ROW_SUFFIX)]
                item = catalog_by_key.get(key)
                if item is None or item.kind != "multi_select":
                    errors.append(
                        _cell_error(
                            code="topic_question_option_row_unknown",
                            sheet=sheet_name,
                            row=row_number,
                            column="A",
                            message=f"多选选项行对应的问题不存在：{key}",
                        )
                    )
                    continue
                option = _text(values[3])
                if option not in (item.options or []):
                    errors.append(
                        _cell_error(
                            code="topic_question_option_unknown",
                            sheet=sheet_name,
                            row=row_number,
                            column="D",
                            message=f"问题 {key} 不包含选项：{option or '（空）'}",
                        )
                    )
                    continue
                choice = _text(values[4])
                if choice not in {"", MULTI_SELECT_YES, MULTI_SELECT_NO}:
                    errors.append(
                        _cell_error(
                            code="topic_question_option_choice_invalid",
                            sheet=sheet_name,
                            row=row_number,
                            column="E",
                            message=(
                                f"选项「{option}」的填写内容只能是"
                                f" {MULTI_SELECT_YES} 或 {MULTI_SELECT_NO}"
                            ),
                        )
                    )
                    continue
                if choice == MULTI_SELECT_YES:
                    multi_selected.setdefault(key, []).append(option)
                continue

            key = raw_key
            item = catalog_by_key.get(key)
            if item is None:
                errors.append(
                    _cell_error(
                        code="topic_question_key_unknown",
                        sheet=sheet_name,
                        row=row_number,
                        column="A",
                        message=f"该问题不属于当前报告的议题范围：{key}",
                    )
                )
                continue
            if key in seen_keys:
                errors.append(
                    _cell_error(
                        code="topic_question_key_duplicate",
                        sheet=sheet_name,
                        row=row_number,
                        column="A",
                        message=f"问题重复出现：{key}",
                    )
                )
                continue
            seen_keys.add(key)
            locations[key] = (sheet_name, row_number)
            raw_answers[key] = _text(values[4])
            raw_supplements[key] = _text(values[5])

    for key, item in catalog_by_key.items():
        if key not in seen_keys:
            errors.append(
                _cell_error(
                    code="topic_question_key_missing",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"工作簿缺少问题行：{key}（请重新下载模板）",
                )
            )

    answers: dict[str, StoredIntakeAnswer] = {}
    for key, item in catalog_by_key.items():
        if key not in seen_keys:
            continue
        sheet_name, row_number = locations[key]
        supplement = raw_supplements.get(key) or None
        answer: str | list[str] | None
        if item.kind == "multi_select":
            selected = multi_selected.get(key, [])
            # 保持合同选项顺序，不受工作簿行序影响。
            answer = [
                option for option in (item.options or []) if option in selected
            ] or None
        else:
            text_value = raw_answers.get(key) or None
            if item.kind == "single_select":
                if text_value is not None and text_value not in (item.options or []):
                    errors.append(
                        _cell_error(
                            code="topic_question_answer_invalid",
                            sheet=sheet_name,
                            row=row_number,
                            column="E",
                            message=f"问题 {key} 的答案必须是模板提供的选项之一",
                        )
                    )
                    continue
                answer = text_value
            else:
                if text_value is not None:
                    if item.maxChars and len(text_value) > item.maxChars:
                        errors.append(
                            _cell_error(
                                code="topic_question_answer_too_long",
                                sheet=sheet_name,
                                row=row_number,
                                column="E",
                                message=(
                                    f"问题 {key} 的回答超过 {item.maxChars} 字上限"
                                ),
                            )
                        )
                        continue
                    if item.minChars and len(text_value) < item.minChars:
                        errors.append(
                            _cell_error(
                                code="topic_question_answer_too_short",
                                sheet=sheet_name,
                                row=row_number,
                                column="E",
                                message=(
                                    f"问题 {key} 的回答不足 {item.minChars} 字；"
                                    "留空表示暂不回答"
                                ),
                            )
                        )
                        continue
                answer = text_value
        if answer is not None or supplement is not None:
            answers[key] = StoredIntakeAnswer(answer=answer, supplement=supplement)

    if errors:
        raise StructuredInputWorkbookError(errors)
    return answers
