# ABOUTME: 双重重要性评分表解析，将官方议题名称和财务/影响得分解析为 typed assessment topic。
# ABOUTME: 评分表可带品牌说明区，但必须有清晰表头并覆盖当前报告配置下的全部适用议题。
# ABOUTME(en): Parses the double materiality scoring sheet, turning official topic names and financial/impact scores
# ABOUTME(en): into typed assessment topics. A branding block may precede the header; all applicable topics must appear.
from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from sustainability_desk.assets.workbook_safety import (
    validate_workbook_dimensions,
    validate_xlsx_container,
)
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.contract.loader import load_package_contract
from sustainability_desk.contract.materiality_scoring import validate_materiality_score
from sustainability_desk.contract.models import Report
from sustainability_desk.contract.structured_inputs import (
    StructuredInputCellError,
    StructuredInputContext,
    StructuredInputWorkbookError,
    assessment_context_fingerprint,
    context_metadata_values,
    validate_metadata_sheet,
)
from sustainability_desk.contract.topic_registry import applicable_scoring_topics, resolve_topic

# 识别得分列：同一表头行需分别含「议题」「财务」「影响」字样。
TOPIC_HEADER = "议题"
FIN_HEADER = "财务"
IMP_HEADER = "影响"
# 工作表择优关键词：优先标题含此类词者。
PREFER_TITLE = ("矩阵", "分数", "双重")


@dataclass
class ScoredTopic:
    assessmentTopicId: str
    financialScore: float
    impactScore: float


@dataclass
class ScoringResult:
    scored: list[ScoredTopic] = field(default_factory=list)
    unmatched: list[str] = field(default_factory=list)
    sheet: str = ""
    note: str = ""
    context_fingerprint: str = ""


def parse_scoring(
    content: bytes,
    *,
    report: Report,
    context: StructuredInputContext,
) -> ScoringResult:
    """解析评分表字节流，并按当前报告配置校验完整适用议题清单。"""
    import openpyxl

    validate_xlsx_container(content)
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as error:  # noqa: BLE001 — 统一投影为 typed 工作簿错误
        raise StructuredInputWorkbookError(
            [
                StructuredInputCellError(
                    code="workbook_invalid",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"无法读取评分工作簿：{error}",
                )
            ]
        ) from error
    validate_workbook_dimensions(wb)
    fingerprint = assessment_context_fingerprint(report, context)
    errors = validate_metadata_sheet(
        wb,
        context_metadata_values(
            input_kind="assessment",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    picked = _pick_sheet(wb)
    if picked is None:
        errors.append(
            StructuredInputCellError(
                code="scoring_header_missing",
                sheet="工作簿",
                row=1,
                column="A",
                message="未找到同一行含『议题』『财务』『影响』表头的工作表",
            )
        )
        raise StructuredInputWorkbookError(errors)
    ws, header_row, topic_col, fin_col, imp_col = picked
    result = ScoringResult(sheet=ws.title, context_fingerprint=fingerprint)
    package = knowledge_package_of(report)
    scale = report.assessmentScoreScale or load_package_contract(package).assessmentScoreScale
    if scale is None:
        raise ValueError("报告模板未声明双重重要性评分尺度")
    scored_by_id: dict[str, ScoredTopic] = {}

    for row_number, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        name = row[topic_col] if topic_col < len(row) else None
        if name is None or str(name).strip() == "":
            continue
        raw_fin = row[fin_col] if fin_col < len(row) else None
        raw_imp = row[imp_col] if imp_col < len(row) else None
        if raw_fin in (None, "") and raw_imp in (None, ""):
            continue
        label = str(name).strip()
        ref = resolve_topic(package, label)
        if ref is None:
            errors.append(
                StructuredInputCellError(
                    code="assessment_topic_unknown",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(topic_col + 1),
                    message=f"评分表包含非官方议题名称：{label}",
                )
            )
            continue
        fin = _num(raw_fin)
        imp = _num(raw_imp)
        if fin is None:
            errors.append(
                StructuredInputCellError(
                    code="financial_score_missing_or_invalid",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(fin_col + 1),
                    message=f"{label} 的财务重要性评分缺失或不是数字",
                )
            )
        if imp is None:
            errors.append(
                StructuredInputCellError(
                    code="impact_score_missing_or_invalid",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(imp_col + 1),
                    message=f"{label} 的影响重要性评分缺失或不是数字",
                )
            )
        if fin is None or imp is None:
            continue
        try:
            fin = validate_materiality_score(fin, scale)
        except ValueError as exc:
            errors.append(
                StructuredInputCellError(
                    code="financial_score_out_of_scale",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(fin_col + 1),
                    message=f"评分表包含不符合评分尺度的数值：{label}（{exc}）",
                )
            )
            continue
        try:
            imp = validate_materiality_score(imp, scale)
        except ValueError as exc:
            errors.append(
                StructuredInputCellError(
                    code="impact_score_out_of_scale",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(imp_col + 1),
                    message=f"评分表包含不符合评分尺度的数值：{label}（{exc}）",
                )
            )
            continue
        if ref.materialityDetermination.kind != "scored":
            errors.append(
                StructuredInputCellError(
                    code="assessment_topic_fixed",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(topic_col + 1),
                    message=f"评分表包含非官方评分议题：{label}（该议题由产品固定分类，无需评分）",
                )
            )
            continue
        if ref.id in scored_by_id:
            errors.append(
                StructuredInputCellError(
                    code="assessment_topic_duplicate",
                    sheet=ws.title,
                    row=row_number,
                    column=get_column_letter(topic_col + 1),
                    message=f"评分表包含重复议题：{label}",
                )
            )
            continue
        scored_by_id[ref.id] = ScoredTopic(
            assessmentTopicId=ref.id,
            financialScore=fin,
            impactScore=imp,
        )

    applicable = applicable_scoring_topics(report)
    required_ids = {topic.id for topic in applicable}
    parsed_ids = set(scored_by_id)
    excluded = [topic_id for topic_id in sorted(parsed_ids - required_ids)]
    missing = [
        topic.name
        for topic in applicable
        if topic.id not in parsed_ids
    ]
    if excluded:
        errors.append(
            StructuredInputCellError(
                code="assessment_topic_not_applicable",
                sheet=ws.title,
                row=header_row,
                column=get_column_letter(topic_col + 1),
                message=f"评分表包含当前报告配置不适用的议题：{'、'.join(excluded)}",
            )
        )
    if missing:
        errors.append(
            StructuredInputCellError(
                code="assessment_topic_missing",
                sheet=ws.title,
                row=header_row,
                column=get_column_letter(topic_col + 1),
                message=f"评分表缺少适用议题：{'、'.join(missing)}",
            )
        )
    if errors:
        raise StructuredInputWorkbookError(errors)

    result.scored = [scored_by_id[topic.id] for topic in applicable]
    return result


def _pick_sheet(wb):
    """择含议题+财务+影响表头的工作表，允许品牌说明区位于表头上方。"""
    candidates = []
    for ws in wb.worksheets:
        for header_row, header in enumerate(
            ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), start=1
        ):
            topic_col = fin_col = imp_col = None
            for j, cell in enumerate(header):
                text = str(cell or "")
                if TOPIC_HEADER in text and topic_col is None:
                    topic_col = j
                if FIN_HEADER in text and fin_col is None:
                    fin_col = j
                if IMP_HEADER in text and imp_col is None:
                    imp_col = j
            # 三个标记必须落在互不相同的列——真表头一词一列。说明区可能在一个合并单元格里
            # 同时提到这三个词（如维度释义行），单看"含关键词"会把说明行误认成表头。
            if (
                topic_col is not None
                and fin_col is not None
                and imp_col is not None
                and len({topic_col, fin_col, imp_col}) == 3
            ):
                candidates.append((ws, header_row, topic_col, fin_col, imp_col))
                break
    if not candidates:
        return None
    for ws, header_row, topic_col, fin_col, imp_col in candidates:
        if any(k in ws.title for k in PREFER_TITLE):
            return ws, header_row, topic_col, fin_col, imp_col
    return candidates[0]


def _num(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
