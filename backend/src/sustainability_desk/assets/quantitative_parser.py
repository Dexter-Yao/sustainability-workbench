# ABOUTME: ESG 定量工作簿严格解析器，逐行累计 key、单位、数值与无值原因错误。
# ABOUTME: 仅整份目录工作簿全部有效时返回 typed 草稿，禁止部分导入或回写副本元数据。
# ABOUTME(en): Strict parser for the ESG quantitative workbook, accumulating key, unit, value and no-value-reason errors
# ABOUTME(en): row by row. Returns a typed draft only when the whole catalog workbook is valid; no partial import.
from __future__ import annotations

import io

import openpyxl
from pydantic import ValidationError

from sustainability_desk.assets.quantitative_template import (
    GHG_STANDARD_CELL,
    GHG_STANDARD_OTHER_CELL,
    QUANTITATIVE_CONFIG_SHEET,
    QUANTITATIVE_HEADER_ROW,
)
from sustainability_desk.assets.workbook_safety import (
    validate_workbook_dimensions,
    validate_xlsx_container,
)
from sustainability_desk.contract.models import (
    QuantitativeMetricDraft,
    QuantitativeMetricsMeta,
    Report,
)
from sustainability_desk.contract.structured_inputs import (
    StructuredInputCellError,
    StructuredInputContext,
    StructuredInputWorkbookError,
    quantitative_metrics_context_fingerprint,
    quantitative_metadata_values,
    validate_metadata_sheet,
)
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.quantitative_metrics import (
    QuantitativeMetricDef,
    derived_sum_metric_value,
    quantitative_metric_display_label,
    quantitative_metrics_by_key,
    all_quantitative_metrics,
    validate_complete_quantitative_metrics,
)


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _apply_derived_sums(
    parsed: dict[str, QuantitativeMetricDraft],
    catalog_by_key: dict[str, QuantitativeMetricDef],
) -> None:
    """就地把目录声明为求和派生的指标改写为求得值；来源未填齐则记为尚未收集。"""

    values_by_key = {key: draft.value for key, draft in parsed.items()}
    for key, metric in catalog_by_key.items():
        if not metric.sumOfMetricKeys or key not in parsed:
            continue
        total = derived_sum_metric_value(metric, values_by_key)
        draft = parsed[key]
        parsed[key] = draft.model_copy(
            update={
                "value": total,
                "noValueReason": None if total else (draft.noValueReason or "not_collected"),
            }
        )


def _cell_error(
    *,
    code: str,
    sheet: str,
    row: int,
    column: str,
    message: str,
) -> StructuredInputCellError:
    return StructuredInputCellError(
        code=code,
        sheet=sheet,
        row=row,
        column=column,
        message=message,
    )


def parse_quantitative_workbook(
    content: bytes,
    *,
    report: Report,
    context: StructuredInputContext,
    allowed_metric_keys: frozenset[str] | None = None,
) -> QuantitativeMetricsMeta:
    """严格解析当前报告范围内的指标及温室气体核算标准。"""

    validate_xlsx_container(content)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as error:  # noqa: BLE001 — 统一投影为 typed 工作簿错误
        raise StructuredInputWorkbookError(
            [
                _cell_error(
                    code="workbook_invalid",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"无法读取定量工作簿：{error}",
                )
            ]
        ) from error
    validate_workbook_dimensions(workbook)

    fingerprint = quantitative_metrics_context_fingerprint(report, context)
    errors = validate_metadata_sheet(
        workbook,
        quantitative_metadata_values(
            report=report,
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    package = knowledge_package_of(report)
    catalog = tuple(
        metric
        for metric in all_quantitative_metrics(package)
        if allowed_metric_keys is None or metric.key in allowed_metric_keys
    )
    catalog_by_key = {metric.key: metric for metric in catalog}
    all_catalog_by_key = quantitative_metrics_by_key(package)
    parsed: dict[str, QuantitativeMetricDraft] = {}
    locations: dict[str, tuple[str, int]] = {}
    seen_keys: set[str] = set()
    standard: str | None = None
    standard_other: str | None = None

    if QUANTITATIVE_CONFIG_SHEET not in workbook.sheetnames:
        errors.append(
            _cell_error(
                code="quantitative_config_sheet_missing",
                sheet=QUANTITATIVE_CONFIG_SHEET,
                row=1,
                column="A",
                message="定量工作簿缺少填写说明与核算标准配置区",
            )
        )
    else:
        config_sheet = workbook[QUANTITATIVE_CONFIG_SHEET]
        standard = _text(config_sheet[GHG_STANDARD_CELL].value) or None
        standard_other = _text(
            config_sheet[GHG_STANDARD_OTHER_CELL].value
        ) or None

    for sheet_name in dict.fromkeys(metric.sheet for metric in catalog):
        if sheet_name not in workbook.sheetnames:
            errors.append(
                _cell_error(
                    code="quantitative_sheet_missing",
                    sheet=sheet_name,
                    row=1,
                    column="A",
                    message=f"定量工作簿缺少工作表：{sheet_name}",
                )
            )
            continue
        worksheet = workbook[sheet_name]
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=QUANTITATIVE_HEADER_ROW + 1,
                values_only=True,
            ),
            start=QUANTITATIVE_HEADER_ROW + 1,
        ):
            values = tuple(row) + (None,) * max(0, 8 - len(row))
            key = _text(values[0])
            if not key:
                if any(_text(value) for value in values[1:8]):
                    errors.append(
                        _cell_error(
                            code="quantitative_key_missing",
                            sheet=sheet_name,
                            row=row_number,
                            column="A",
                            message="定量指标行缺少指标 key",
                        )
                    )
                continue
            if key in seen_keys:
                errors.append(
                    _cell_error(
                        code="quantitative_key_duplicate",
                        sheet=sheet_name,
                        row=row_number,
                        column="A",
                        message=f"定量指标 key 重复：{key}",
                    )
                )
                continue
            seen_keys.add(key)
            locations[key] = (sheet_name, row_number)
            metric = catalog_by_key.get(key)
            if metric is None:
                detail = (
                    "该指标不属于指标目录"
                    if key in all_catalog_by_key
                    else "该指标 key 不存在"
                )
                errors.append(
                    _cell_error(
                        code="quantitative_key_unknown",
                        sheet=sheet_name,
                        row=row_number,
                        column="A",
                        message=f"{detail}：{key}",
                    )
                )
                continue
            if metric.sheet != sheet_name:
                errors.append(
                    _cell_error(
                        code="quantitative_key_wrong_sheet",
                        sheet=sheet_name,
                        row=row_number,
                        column="A",
                        message=f"指标 {key} 应位于工作表 {metric.sheet}",
                    )
                )
            unit = _text(values[3])
            metric_name = _text(values[2])
            expected_name = quantitative_metric_display_label(metric)
            if metric_name != expected_name:
                errors.append(
                    _cell_error(
                        code="quantitative_metric_name_mismatch",
                        sheet=sheet_name,
                        row=row_number,
                        column="C",
                        message=f"指标 {key} 的官方名称必须为：{expected_name}",
                    )
                )
            if unit != metric.unit:
                errors.append(
                    _cell_error(
                        code="quantitative_unit_mismatch",
                        sheet=sheet_name,
                        row=row_number,
                        column="D",
                        message=f"指标 {key} 的单位必须为：{metric.unit}",
                    )
                )
            # 派生行的「数值」格是模板写下的自动求和说明，不是用户输入；一律丢弃，
            # 交由 _apply_derived_sums 从来源指标求得，避免把提示文字当成非法数值报错。
            value = None if metric.sumOfMetricKeys else (_text(values[4]) or None)
            no_value_reason = _text(values[5]) or None
            try:
                parsed[key] = QuantitativeMetricDraft(
                    value=value,
                    noValueReason=no_value_reason,
                    department=_text(values[6]) or None,
                    note=_text(values[7]) or None,
                )
            except ValidationError as error:
                code = "quantitative_draft_invalid"
                if value is not None and no_value_reason is not None:
                    code = "quantitative_value_choice_invalid"
                elif no_value_reason is not None:
                    code = "quantitative_no_value_reason_invalid"
                errors.append(
                    _cell_error(
                        code=code,
                        sheet=sheet_name,
                        row=row_number,
                        column="E",
                        message=f"指标 {key} 填写不合法：{error.errors()[0]['msg']}",
                    )
                )

    # 派生指标在解析边界就地求和：模板已将其标注为自动得出，用户即便手填也不采信，
    # 否则同一事实会有工作簿与求和两个来源。与在线录入共用同一目录声明。
    _apply_derived_sums(parsed, catalog_by_key)
    metrics = QuantitativeMetricsMeta(
        metrics=parsed,
        greenhouseGasAccountingStandard=standard,
        greenhouseGasAccountingStandardOther=standard_other,
    )
    semantic_issues = validate_complete_quantitative_metrics(
        metrics,
        package=package,
        allowed_metric_keys=allowed_metric_keys,
    )
    # 工作簿是往返产物：每个指标行由导出侧按目录写入，行缺失意味着文件被改坏或
    # 传错，与「用户选择不填某项」无关（后者在页面上是合法的，行仍在、值为空）。
    # 故完整性由本解析器按目录自行判定，不复用语义校验——语义校验只管已作答的项。
    expected_keys = {
        metric.key
        for metric in all_quantitative_metrics(package)
        if allowed_metric_keys is None or metric.key in allowed_metric_keys
    }
    missing_count = len(expected_keys - set(parsed))
    if missing_count:
        errors.append(
            _cell_error(
                code="quantitative_keys_missing",
                sheet="工作簿",
                row=1,
                column="A",
                message=f"定量工作簿缺少或未有效填写 {missing_count} 个指标 key",
            )
        )
    for issue in semantic_issues:
        if issue.metric_key is not None:
            sheet, row = locations.get(issue.metric_key, ("工作簿", 1))
            column = "E" if issue.metric_key in locations else "A"
        else:
            config_cell = (
                GHG_STANDARD_OTHER_CELL if "other" in issue.code else GHG_STANDARD_CELL
            )
            sheet = QUANTITATIVE_CONFIG_SHEET
            row = int(config_cell[1:])
            column = config_cell[0]
        errors.append(
            _cell_error(
                code=issue.code,
                sheet=sheet,
                row=row,
                column=column,
                message=issue.message,
            )
        )
    if errors:
        raise StructuredInputWorkbookError(errors)
    return metrics
