# ABOUTME: 基础资料工作簿生成器：企业及报告基本信息、公司简介与治理三题、披露准则与附录联系的单表投影。
# ABOUTME: 行计划是唯一编排源并与合同 user_input 字段全集互相核对；模板预填当前值，导入为整批替换。
# ABOUTME(en): Report basics workbook generator: entity and report info, company profile, governance questions,
# ABOUTME(en): disclosure requirements and appendix contacts on one sheet. Prefilled; import replaces the whole batch.
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from sustainability_desk.contract.product_name import workbook_sheet_title
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.contract.models import Report
from sustainability_desk.contract.structured_inputs import (
    StructuredInputContext,
    append_metadata_sheet,
    canonical_json_sha256,
    context_metadata_values,
)

_BACKEND = Path(__file__).resolve().parents[3]
_ROOT = _BACKEND.parent

REPORT_BASICS_SHEET = "基础资料"
REPORT_BASICS_HEADER_ROW = 7
REPORT_BASICS_COLUMNS: tuple[str, ...] = (
    "键",
    "信息类别",
    "填写项目",
    "填写要求",
    "填写格式 / 可选值",
    "填写内容",
    "填写提示",
)
BOOLEAN_YES = "是"
BOOLEAN_NO = "否"
# 长选项清单的隐藏承载区：从 J 列（第 10 列）起，每个长 enum 占一列。
_OPTION_AREA_FIRST_COLUMN = 10
_OPTION_AREA_TOP = REPORT_BASICS_HEADER_ROW + 1


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)

_PINE = "2C6E49"
_PINE_SUBTLE = "E7F0E9"
_BORDER = "E2E6E1"
_MUTED = "677069"
_FOREGROUND = "19211D"

type RowKind = Literal["field", "intake", "disclosure", "appendix"]


@dataclass(frozen=True)
class ReportBasicsRow:
    """基础资料工作簿的一行编排；field/intake 行的定义细节取自合同，不在此复制。"""

    section: str
    kind: RowKind
    key: str
    # 仅 disclosure/appendix 行需要自带展示与格式（它们不在 fields 合同内）。
    label: str | None = None
    format_hint: str | None = None
    options: tuple[str, ...] | None = None
    hint: str | None = None


# 行计划（业务编排 SSOT）：分区与顺序对齐 /intake/info 页；
# field/intake 行的 label、类型、选项、必填性一律取自合同。
REPORT_BASICS_ROWS: tuple[ReportBasicsRow, ...] = (
    ReportBasicsRow("公司主体", "field", "company_registered_name"),
    ReportBasicsRow("公司主体", "field", "company_short_name"),
    ReportBasicsRow("公司主体", "field", "industry_major_category"),
    ReportBasicsRow("公司主体", "field", "industry_division"),
    ReportBasicsRow("公司简介", "intake", "company_profile"),
    ReportBasicsRow("公司简介", "intake", "company_certifications"),
    ReportBasicsRow("报告期", "field", "reporting_year"),
    ReportBasicsRow("报告期", "field", "report_period_start"),
    ReportBasicsRow("报告期", "field", "report_period_end"),
    ReportBasicsRow("报告期", "field", "consolidation_scope"),
    ReportBasicsRow("报告期", "field", "consolidation_scope_note"),
    ReportBasicsRow(
        "披露准则",
        "disclosure",
        "mainlandStandard",
        label="大陆准则",
        format_hint="单选：在「填写内容」下拉选择",
        # Options come from the package's disclosure basis at render time (see _row_options).
    ),
    ReportBasicsRow("适用范围", "field", "has_technology_ethics_sensitive_activity"),
    ReportBasicsRow("发布与审批", "field", "report_publication_channel"),
    ReportBasicsRow("发布与审批", "field", "report_publication_website_url"),
    ReportBasicsRow("发布与审批", "field", "report_approval_year"),
    ReportBasicsRow("发布与审批", "field", "report_approval_month"),
    ReportBasicsRow("发布与审批", "field", "report_approval_body"),
    ReportBasicsRow(
        "外部鉴证",
        "appendix",
        "externalAssuranceReport.isIncluded",
        label="本报告包含外部鉴证报告",
        format_hint=f"单选：{BOOLEAN_YES} / {BOOLEAN_NO}",
        options=(BOOLEAN_YES, BOOLEAN_NO),
    ),
    ReportBasicsRow(
        "外部鉴证",
        "appendix",
        "externalAssuranceReport.fileLabel",
        label="外部鉴证报告文件名",
        format_hint="文本",
        hint="仅当包含外部鉴证报告时填写。",
    ),
    ReportBasicsRow("外部鉴证", "field", "assurance_provider_name"),
    ReportBasicsRow("外部鉴证", "field", "assurance_standard"),
    ReportBasicsRow(
        "读者反馈联系",
        "appendix",
        "readerFeedbackContactInformation.email",
        label="读者反馈邮箱",
        format_hint="邮箱地址",
    ),
    ReportBasicsRow(
        "读者反馈联系",
        "appendix",
        "readerFeedbackContactInformation.address",
        label="公司地址",
        format_hint="文本",
    ),
    ReportBasicsRow(
        "读者反馈联系",
        "appendix",
        "readerFeedbackContactInformation.phone",
        label="联系电话",
        format_hint="文本",
    ),
    ReportBasicsRow("公司治理", "field", "board_meeting_count"),
    ReportBasicsRow("公司治理", "field", "board_attendance_rate"),
    ReportBasicsRow("公司治理", "intake", "articles"),
    ReportBasicsRow("可持续发展管理", "intake", "sustainability_strategy"),
    ReportBasicsRow("可持续发展管理", "intake", "sustainability_governance_structure"),
    ReportBasicsRow("可持续发展管理", "intake", "sustainability_governance_duties"),
)

_FIELD_FORMAT_HINTS: dict[str, str] = {
    "string": "文本",
    "number": "数字",
    "year": "四位年份",
    "month": "年月（YYYY-MM）",
    "date": "日期（YYYY-MM-DD）",
    "email": "邮箱地址",
    "percent": "百分比数值（不带 % 号）",
    "url": "网址",
}


def report_basics_row_key(row: ReportBasicsRow) -> str:
    """隐藏键列协议：kind:key。"""

    return f"{row.kind}:{row.key}"


def validate_report_basics_row_plan(report: Report) -> None:
    """行计划必须与合同 user_input 字段全集一致；合同演进漏行在生成时 fail-loud。"""

    planned = {row.key for row in REPORT_BASICS_ROWS if row.kind == "field"}
    contract_fields = {
        key
        for key, field in report.fields.items()
        if field.source == "user_input"
    }
    missing = contract_fields - planned
    extra = planned - contract_fields
    if missing or extra:
        raise ValueError(
            f"基础资料行计划与合同字段不一致：缺少 {sorted(missing)}，多余 {sorted(extra)}"
        )
    planned_intake = {row.key for row in REPORT_BASICS_ROWS if row.kind == "intake"}
    contract_intake = {item.key for item in report.intakeItems}
    if planned_intake != contract_intake:
        raise ValueError(
            "基础资料行计划与合同报告级题不一致："
            f"计划 {sorted(planned_intake)}，合同 {sorted(contract_intake)}"
        )


def report_basics_context_fingerprint(
    report: Report,
    context: StructuredInputContext,
) -> str:
    """指纹化基础资料模板真实依赖：行计划、字段类型与合法选项。"""

    rows_payload = []
    for row in REPORT_BASICS_ROWS:
        if row.kind == "field":
            field = report.fields[row.key]
            rows_payload.append(
                {
                    "key": report_basics_row_key(row),
                    "type": field.type,
                    "options": field.options,
                    "required": field.required,
                }
            )
        elif row.kind == "intake":
            item = next(item for item in report.intakeItems if item.key == row.key)
            rows_payload.append(
                {
                    "key": report_basics_row_key(row),
                    "minChars": item.minChars,
                    "maxChars": item.maxChars,
                }
            )
        else:
            options = _row_options(report, row)
            rows_payload.append(
                {
                    "key": report_basics_row_key(row),
                    "options": list(options) if options else None,
                }
            )
    payload = {
        "reportId": str(context.reportId),
        "contractVersion": context.contractVersion,
        "compiledSemanticsVersion": context.compiledSemanticsVersion,
        "kind": "report_basics",
        "rows": rows_payload,
    }
    return canonical_json_sha256(payload)


def _mainland_standard_names(report: Report) -> dict[str, str]:
    names = knowledge_package_of(report).manifest.disclosure_basis.mainland_standard_names
    return names.by_code() if names else {}


def _row_options(report: Report, row: ReportBasicsRow) -> tuple[str, ...] | None:
    """Static row options, except the mainland standard names, which the package basis owns."""

    if row.kind == "disclosure" and row.key == "mainlandStandard":
        return tuple(_mainland_standard_names(report).values()) or None
    return row.options


def _current_value(report: Report, row: ReportBasicsRow) -> str | int | float | None:
    if row.kind == "field":
        field = report.fields.get(row.key)
        return field.value if field is not None else None
    if row.kind == "intake":
        item = next(
            (item for item in report.intakeItems if item.key == row.key), None
        )
        return item.answer if item is not None and isinstance(item.answer, str) else None
    profile = report.disclosureProfile
    appendix = report.appendixPackage
    if row.kind == "disclosure":
        if profile is None:
            return None
        if row.key == "mainlandStandard":
            return _mainland_standard_names(report).get(profile.mainlandStandard)
    if row.kind == "appendix":
        if appendix is None:
            return None
        if row.key == "externalAssuranceReport.isIncluded":
            return (
                BOOLEAN_YES
                if appendix.externalAssuranceReport.isIncluded
                else BOOLEAN_NO
            )
        if row.key == "externalAssuranceReport.fileLabel":
            return appendix.externalAssuranceReport.fileLabel
        contact = appendix.readerFeedbackContactInformation
        attribute = row.key.split(".", 1)[1]
        return getattr(contact, attribute, None)
    return None


def _guidance_text(report: Report, path: str) -> str:
    """从合同 inputGuidance 投影 Excel 的「填写提示」，不在此另写一份文案。

    网页按 design.md §2.2.1 分两处呈现（helpText 常驻、termExplanation 折进 ⓘ），
    Excel 无折叠形态，故顺序拼接为一格。不在此维护私有 hints dict：与合同对同一
    字段各写一份是网页与 Excel 文案漂移的来源。
    """

    guidance = (report.inputGuidance or {}).get(path)
    if guidance is None:
        return ""
    return "　".join(
        part for part in (guidance.helpText, guidance.termExplanation) if part
    )


def _row_presentation(report: Report, row: ReportBasicsRow) -> tuple[str, str, str, str]:
    """返回（填写项目、填写要求、格式/可选值、填写提示）。"""

    if row.kind == "field":
        field = report.fields[row.key]
        requirement = "必填" if field.required else "选填"
        if field.type == "enum" and field.options:
            format_hint = "单选：在「填写内容」下拉选择"
        else:
            format_hint = _FIELD_FORMAT_HINTS.get(field.type, "文本")
        return (
            field.label,
            requirement,
            format_hint,
            _guidance_text(report, f"fields.{row.key}.value"),
        )
    if row.kind == "intake":
        item = next(item for item in report.intakeItems if item.key == row.key)
        limits: list[str] = []
        if item.minChars:
            limits.append(f"至少 {item.minChars} 字")
        if item.maxChars:
            limits.append(f"最多 {item.maxChars} 字")
        format_hint = f"长文本（{'，'.join(limits)}）" if limits else "长文本"
        requirement = "必填" if item.collectionPriority == "core" else "选填"
        item_hint = "　".join(p for p in (item.hint, item.termExplanation) if p)
        return item.prompt, requirement, format_hint, item_hint
    return row.label or row.key, "选填", row.format_hint or "文本", row.hint or ""


def create_report_basics_template(
    report: Report,
    *,
    context: StructuredInputContext,
) -> bytes:
    """从合同与当前报告状态投影基础资料工作簿；不得维护静态字段副本。"""

    fingerprint = report_basics_context_fingerprint(report, context)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_instructions(workbook)
    append_report_basics_sheet(workbook, report)
    append_metadata_sheet(
        workbook,
        context_metadata_values(
            input_kind="report_basics",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def append_report_basics_sheet(workbook: Workbook, report: Report) -> None:
    """向工作簿追加基础资料业务工作表；说明页与元数据页由调用方写入。"""

    validate_report_basics_row_plan(report)
    worksheet = workbook.create_sheet(REPORT_BASICS_SHEET)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = _PINE
    worksheet.freeze_panes = f"A{REPORT_BASICS_HEADER_ROW + 1}"
    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["A"].hidden = True
    worksheet.column_dimensions["B"].width = 14
    worksheet.column_dimensions["C"].width = 44
    worksheet.column_dimensions["D"].width = 10
    worksheet.column_dimensions["E"].width = 34
    worksheet.column_dimensions["F"].width = 46
    worksheet.column_dimensions["G"].width = 48

    worksheet.row_dimensions[1].height = 48
    worksheet.merge_cells("B3:G3")
    title_cell = worksheet.cell(3, 2, "企业及报告基本信息")
    title_cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color=_FOREGROUND)
    worksheet.merge_cells("B4:G4")
    hint_cell = worksheet.cell(
        4, 2, "标注「必填」的项目必须填写；其余可按实际情况填写。只需修改「填写内容」列。"
    )
    hint_cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED)

    for column, label in enumerate(REPORT_BASICS_COLUMNS, start=1):
        cell = worksheet.cell(REPORT_BASICS_HEADER_ROW, column, label)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_PINE)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        top=Side(style="thin", color=_BORDER),
        bottom=Side(style="thin", color=_BORDER),
    )
    row_number = REPORT_BASICS_HEADER_ROW
    option_column = _OPTION_AREA_FIRST_COLUMN - 1
    for row in REPORT_BASICS_ROWS:
        row_number += 1
        label, requirement, format_hint, hint = _row_presentation(report, row)
        values = (
            report_basics_row_key(row),
            row.section,
            label,
            requirement,
            format_hint,
            _current_value(report, row),
            hint,
        )
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = Font(
                name="Microsoft YaHei",
                size=10 if column != 1 else 9,
                bold=column == 3,
                color=_FOREGROUND if column != 1 else _MUTED,
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in {3, 5, 6, 7},
            )
            cell.border = thin_border
            if column == 6:
                cell.fill = PatternFill("solid", fgColor=_PINE_SUBTLE)

        options: tuple[str, ...] | None = None
        if row.kind == "field":
            field = report.fields[row.key]
            if field.type == "enum" and field.options:
                options = tuple(str(option) for option in field.options)
        elif _row_options(report, row):
            options = _row_options(report, row)
        if options:
            joined = ",".join(options)
            if len(joined) + 2 <= 255:
                validation = DataValidation(
                    type="list",
                    formula1=f'"{joined}"',
                    allow_blank=True,
                )
            else:
                # Excel 内联 list 校验公式上限 255 字符；长选项清单（行业门类、准则全称）
                # 落在本表隐藏选项区，下拉引用同表区域（各 Excel 版本均支持）。
                option_column += 1
                letter = _column_letter(option_column)
                for offset, option in enumerate(options):
                    worksheet.cell(_OPTION_AREA_TOP + offset, option_column, option)
                worksheet.column_dimensions[letter].hidden = True
                validation = DataValidation(
                    type="list",
                    formula1=(
                        f"${letter}${_OPTION_AREA_TOP}:"
                        f"${letter}${_OPTION_AREA_TOP + len(options) - 1}"
                    ),
                    allow_blank=True,
                )
            validation.error = "请选择模板提供的选项"
            validation.errorTitle = "选项不合法"
            validation.showErrorMessage = True
            worksheet.add_data_validation(validation)
            validation.add(f"F{row_number}")
            display = " / ".join(options)
            worksheet.cell(
                row_number,
                5,
                (
                    "单选：" + display
                    if len(display) <= 60
                    else f"单选：在「填写内容」下拉选择（共 {len(options)} 项）"
                ),
            ).alignment = Alignment(vertical="top", wrap_text=True)



def _add_instructions(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet("填写说明")
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 110

    worksheet.row_dimensions[1].height = 54

    worksheet["A3"] = workbook_sheet_title("企业及报告基本信息表")
    worksheet["A3"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=_PINE)
    worksheet.row_dimensions[3].height = 30

    instructions = (
        "本表覆盖企业及报告基本信息、公司简介与治理制度描述；只需填写「填写内容」列，其余列为只读说明。",
        "单选项目在「填写内容」下拉选择；「其他参考文件」等多值项目在单元格内换行（Alt+Enter）分隔。",
        "标注「必填」的项目必须填写；日期使用 YYYY-MM-DD 格式，数值不带单位与千分位。",
        "导入时整簿替换当前内容：请先下载最新模板（已预填现有填写），修改后再导入。",
    )
    for row_offset, line in enumerate(instructions):
        row_number = 4 + row_offset
        cell = worksheet.cell(row_number, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
        worksheet.row_dimensions[row_number].height = 26
