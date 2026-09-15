# ABOUTME: 轻量版 ESG 定量信息工作簿生成器，从唯一权威指标目录投影三张可填写工作表。
# ABOUTME: 企业、期间、合并范围和单位均为只读副本；导入时必须与当前 Report 上下文一致。
# ABOUTME(en): Lightweight ESG quantitative workbook generator, projecting three fillable sheets from the catalog.
# ABOUTME(en): Entity, period, consolidation scope and units are read-only copies checked against Report on import.
from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from sustainability_desk.contract.product_name import workbook_sheet_title
from sustainability_desk.contract.models import Report
from sustainability_desk.contract.structured_inputs import (
    StructuredInputContext,
    append_metadata_sheet,
    quantitative_metrics_context_fingerprint,
    quantitative_metadata_values,
)
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.quantitative_metrics import (
    greenhouse_gas_accounting_standard_options,
    quantitative_metric_sheets,
    QuantitativeMetricDef,
    quantitative_metric_display_label,
    quantitative_metrics_by_key,
    quantitative_metrics_meta,
    quantitative_no_value_reasons,
    quantitative_no_value_reason_labels,
    all_quantitative_metrics,
)

_BACKEND = Path(__file__).resolve().parents[3]
_ROOT = _BACKEND.parent

QUANTITATIVE_HEADER_ROW = 8
QUANTITATIVE_CONFIG_SHEET = "填写说明"
GHG_STANDARD_CELL = "B8"
GHG_STANDARD_OTHER_CELL = "B9"
# 「指标说明」置于用户输入列之后：解析器按固定下标读取数值/无值原因/部门/备注
# （quantitative_parser 的 values[4]–values[7]），在中间插列会改写导入契约。
# 说明列是只读投影，与网页端同源于 QuantitativeMetricDef.metricDefinition。
QUANTITATIVE_COLUMNS: tuple[str, ...] = (
    "指标 key",
    "分类路径",
    "指标名称",
    "单位",
    "数值",
    "无值原因",
    "数据提供部门",
    "备注",
    "指标说明",
)
NO_VALUE_REASONS = quantitative_no_value_reasons()

_PINE = "2C6E49"
_PINE_SUBTLE = "E7F0E9"
_SURFACE = "F6F7F5"
_BORDER = "E2E6E1"
_MUTED = "677069"
_FOREGROUND = "19211D"


def _metric_guidance_text(metric: QuantitativeMetricDef) -> str:
    """把网页端的两段说明合并为 Excel 单元格文本。

    网页按 design.md §2.2.1 分层呈现（常驻段 + ⓘ 折叠段），Excel 没有折叠形态，
    故顺序拼接为一段。两端同源于目录字段，不在此另写一份文案。
    """

    return "　".join(part for part in (metric.metricDefinition, metric.termExplanation) if part)


def _field_text(report: Report, key: str) -> str:
    field = report.fields.get(key)
    return "" if field is None or field.value is None else str(field.value)


def _consolidation_scope_display(report: Report) -> str:
    """合并范围只读口径：枚举选项直接展示；「特殊口径」以说明原文为实质内容。"""
    scope = _field_text(report, "consolidation_scope")
    note = _field_text(report, "consolidation_scope_note")
    if scope == "特殊口径" and note:
        return f"{scope}（{note}）"
    return scope


def _display_metadata(report: Report) -> tuple[str, ...]:
    return (
        f"企业名称：{_field_text(report, 'company_registered_name')}",
        (
            "报告期间："
            f"{_field_text(report, 'report_period_start')} 至 "
            f"{_field_text(report, 'report_period_end')}"
        ),
        f"合并范围：{_consolidation_scope_display(report)}",
    )



def _add_instructions(
    workbook: Workbook,
    report: Report,
    *,
    config_sheet_title: str = QUANTITATIVE_CONFIG_SHEET,
) -> None:
    worksheet = workbook.create_sheet(config_sheet_title)
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 88

    worksheet.row_dimensions[1].height = 54

    worksheet["A3"] = workbook_sheet_title("定量信息表")
    worksheet["A3"].font = Font(
        name="Microsoft YaHei", size=16, bold=True, color=_PINE,
    )
    worksheet.row_dimensions[3].height = 30

    reason_options = "、".join(
        f"{code}（{label}）"
        for code, label in quantitative_no_value_reason_labels(knowledge_package_of(report)).items()
    )
    instructions = (
        "每项指标必须且只能填写「数值」或「无值原因」之一；数值不得包含单位或千分位。",
        f"无值原因可选：{reason_options}。",
        "企业名称、报告期间、合并范围和单位是当前报告的只读副本；请勿修改。",
    )
    for row_offset, line in enumerate(instructions):
        row_number = 4 + row_offset
        cell = worksheet.cell(row_number, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
        worksheet.row_dimensions[row_number].height = 30

    quantitative_meta = quantitative_metrics_meta(report)
    worksheet["A8"] = "温室气体核算标准"
    worksheet[GHG_STANDARD_CELL] = (
        quantitative_meta.get("greenhouseGasAccountingStandard") or None
    )
    worksheet["A9"] = "选择「其他」时请说明"
    worksheet[GHG_STANDARD_OTHER_CELL] = (
        quantitative_meta.get("greenhouseGasAccountingStandardOther") or None
    )
    for coordinate in ("A8", "A9"):
        worksheet[coordinate].font = Font(
            name="Microsoft YaHei", size=10, bold=True, color=_PINE,
        )
    for coordinate in (GHG_STANDARD_CELL, GHG_STANDARD_OTHER_CELL):
        worksheet[coordinate].font = Font(name="Microsoft YaHei", size=10)
        worksheet[coordinate].fill = PatternFill("solid", fgColor=_PINE_SUBTLE)

    standard_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(greenhouse_gas_accounting_standard_options(knowledge_package_of(report)))}"',
        allow_blank=True,
    )
    standard_validation.error = "请选择模板提供的温室气体核算标准"
    standard_validation.errorTitle = "核算标准不合法"
    standard_validation.showErrorMessage = True
    standard_validation.promptTitle = "温室气体核算标准"
    standard_validation.prompt = "请选择企业采用的核算标准；选择「其他」时在下一行说明。"
    standard_validation.showInputMessage = True
    worksheet.add_data_validation(standard_validation)
    standard_validation.add(GHG_STANDARD_CELL)


def create_quantitative_template(
    report: Report,
    *,
    context: StructuredInputContext,
    allowed_metric_keys: frozenset[str] | None = None,
) -> bytes:
    """从当前报告范围的权威指标目录生成工作簿；不得维护静态行副本。"""

    fingerprint = quantitative_metrics_context_fingerprint(report, context)
    workbook = Workbook()
    workbook.remove(workbook.active)
    append_quantitative_sheets(
        workbook,
        report,
        allowed_metric_keys=allowed_metric_keys,
    )
    append_metadata_sheet(
        workbook,
        quantitative_metadata_values(
            report=report,
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def append_quantitative_sheets(
    workbook: Workbook,
    report: Report,
    *,
    allowed_metric_keys: frozenset[str] | None = None,
    config_sheet_title: str = QUANTITATIVE_CONFIG_SHEET,
) -> None:
    """向工作簿追加定量说明（含 GHG 配置格）与业务工作表；元数据页由调用方写入。"""

    package = knowledge_package_of(report)
    metrics = tuple(
        metric
        for metric in all_quantitative_metrics(package)
        if allowed_metric_keys is None or metric.key in allowed_metric_keys
    )
    _add_instructions(workbook, report, config_sheet_title=config_sheet_title)

    for sheet_name in quantitative_metric_sheets(package):
        sheet_metrics = [metric for metric in metrics if metric.sheet == sheet_name]
        if not sheet_metrics:
            continue
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_properties.tabColor = _PINE
        worksheet.freeze_panes = f"A{QUANTITATIVE_HEADER_ROW + 1}"
        worksheet.column_dimensions["A"].width = 32
        worksheet.column_dimensions["A"].hidden = True
        worksheet.column_dimensions["B"].width = 38
        worksheet.column_dimensions["C"].width = 36
        worksheet.column_dimensions["D"].width = 24
        worksheet.column_dimensions["E"].width = 18
        worksheet.column_dimensions["F"].width = 22
        worksheet.column_dimensions["G"].width = 24
        worksheet.column_dimensions["H"].width = 36
        worksheet.column_dimensions["I"].width = 60

        worksheet.row_dimensions[1].height = 48

        meta_lines = _display_metadata(report)
        for row_offset, line in enumerate(meta_lines):
            row_number = 2 + row_offset
            worksheet.merge_cells(
                start_row=row_number,
                start_column=1,
                end_row=row_number,
                end_column=8,
            )
            cell = worksheet.cell(row_number, 1, line)
            cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
            cell.alignment = Alignment(wrap_text=True)

        # Title row
        worksheet.merge_cells("A6:H6")
        title_cell = worksheet.cell(6, 1, f"{sheet_name} — ESG 定量信息")
        title_cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color=_FOREGROUND)
        title_cell.alignment = Alignment(vertical="center")
        worksheet.row_dimensions[6].height = 28

        thin_border = Border(
            top=Side(style="thin", color=_BORDER),
            bottom=Side(style="thin", color=_BORDER),
        )
        for column, label in enumerate(QUANTITATIVE_COLUMNS, start=1):
            cell = worksheet.cell(QUANTITATIVE_HEADER_ROW, column, label)
            cell.font = Font(
                name="Microsoft YaHei", size=10, bold=True, color="FFFFFF",
            )
            cell.fill = PatternFill("solid", fgColor=_PINE)
            cell.alignment = Alignment(
                horizontal="center", vertical="center",
            )
            cell.border = Border(
                top=Side(style="medium", color=_PINE),
                bottom=Side(style="medium", color=_PINE),
            )

        first_data_row = QUANTITATIVE_HEADER_ROW + 1
        for offset, metric in enumerate(sheet_metrics):
            row_number = first_data_row + offset
            values = (
                metric.key,
                # groupPath[0] 即 category（144 条全等），再拼 category 会让首层出现两次。
                # Excel 无分组标题承载首层，故保留完整 groupPath，只去掉多拼的那一次。
                " / ".join(metric.groupPath),
                quantitative_metric_display_label(metric),
                metric.unit,
                None,
                None,
                None,
                None,
                _metric_guidance_text(metric),
            )
            # 求和派生指标（如温室气体排放总量＝范围一＋范围二）由导入解析就地算出，
            # 模板不留可填空格：留白会让用户以为需要自己填，填了又不被采信。
            derived_sources = [
                quantitative_metric_display_label(source)
                for key in metric.sumOfMetricKeys
                if (source := quantitative_metrics_by_key(package).get(key)) is not None
            ]
            if derived_sources:
                values = (
                    *values[:4],
                    f"自动求和：{' ＋ '.join(derived_sources)}",
                    None,
                    None,
                    None,
                    values[8],
                )
            for column, value in enumerate(values, start=1):
                cell = worksheet.cell(row_number, column, value)
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=column in {2, 3, 8, 9},
                )
                cell.border = thin_border
                if column in {5, 6, 7, 8}:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor=_SURFACE if derived_sources and column in {5, 6} else _PINE_SUBTLE,
                    )
                if derived_sources and column == 5:
                    cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED, italic=True)

        last_data_row = first_data_row + len(sheet_metrics) - 1
        reason_validation = DataValidation(
            type="list",
            formula1=f'"{",".join(NO_VALUE_REASONS)}"',
            allow_blank=True,
        )
        reason_validation.error = "请选择模板提供的无值原因代码"
        reason_validation.errorTitle = "无值原因不合法"
        reason_validation.showErrorMessage = True
        worksheet.add_data_validation(reason_validation)
        reason_validation.add(f"F{first_data_row}:F{last_data_row}")
        worksheet.auto_filter.ref = (
            f"A{QUANTITATIVE_HEADER_ROW}:H{last_data_row}"
        )
