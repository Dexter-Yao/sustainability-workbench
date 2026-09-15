# ABOUTME: 统一填报工作簿：把基础资料、重要性评分、定量信息与议题问题四类模板聚合为单一 xlsx。
# ABOUTME: 导入侧按 sheet 归属重建各子工作簿（值拷贝 + 注入对应元数据页），使四个既有解析器零改动复用。
# ABOUTME(en): Unified intake workbook: aggregates basics, materiality scoring, quantitative and topic question sheets.
# ABOUTME(en): On import it rebuilds each sub-workbook by sheet ownership, so the four parsers are reused unchanged.
from __future__ import annotations

import io
from collections.abc import Mapping
from dataclasses import dataclass
from uuid import UUID

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from sustainability_desk.contract.product_name import workbook_sheet_title
from sustainability_desk.assets.quantitative_template import (
    GHG_STANDARD_CELL,
    GHG_STANDARD_OTHER_CELL,
    QUANTITATIVE_CONFIG_SHEET,
    QUANTITATIVE_HEADER_ROW,
    append_quantitative_sheets,
)
from sustainability_desk.quantitative_metrics import quantitative_metric_sheets
from sustainability_desk.assets.report_basics_template import (
    REPORT_BASICS_SHEET,
    append_report_basics_sheet,
    report_basics_context_fingerprint,
)
from sustainability_desk.assets.scoring_template import (
    SCORING_SHEET_TITLE,
    append_scoring_sheet,
)
from sustainability_desk.assets.topic_questions_template import (
    append_topic_questions_sheets,
)
from sustainability_desk.assets.workbook_safety import (
    validate_workbook_dimensions,
    validate_xlsx_container,
)
from sustainability_desk.contract.models import Report
from sustainability_desk.contract.structured_inputs import (
    STRUCTURED_INPUT_METADATA_SHEET,
    StructuredInputCellError,
    StructuredInputContext,
    StructuredInputWorkbookError,
    append_metadata_sheet,
    assessment_context_fingerprint,
    canonical_json_sha256,
    context_metadata_values,
    quantitative_metadata_values,
    quantitative_metrics_context_fingerprint,
    topic_questions_context_fingerprint,
    validate_metadata_sheet,
)
from sustainability_desk.contract.knowledge_packages import knowledge_package_of
from sustainability_desk.contract.topic_registry import load_topic_contract

UNIFIED_OVERVIEW_SHEET = "总览"
# 统一册内定量说明页改名以避免与总览语义混淆；重建子工作簿时映射回标准名供解析器消费。
UNIFIED_QUANTITATIVE_CONFIG_SHEET = "定量填写说明"

# 空白母版（不绑定任何报告的账户级模板）的规范身份：元数据 report_id 为全零 UUID。
# 导入时按工作簿自我声明的模板血统分支校验——两条血统都是全等校验，不放宽任何一侧。
BLANK_UNIFIED_WORKBOOK_REPORT_ID = UUID(int=0)

_PINE = "2C6E49"
_MUTED = "677069"


@dataclass(frozen=True)
class UnifiedWorkbookParts:
    """按归属拆出的子工作簿字节；评分/定量在整册中留空时为 None（跳过导入、保持现状）。"""

    report_basics: bytes
    scoring: bytes | None
    quantitative: bytes | None
    topic_questions: bytes


def unified_workbook_context_fingerprint(
    basics_report: Report,
    revision_report: Report,
    context: StructuredInputContext,
    *,
    allowed_metric_keys: frozenset[str] | None = None,
    allowed_report_section_ids: frozenset[str] | None = None,
) -> str:
    """统一指纹 = 四个子模板指纹的规范组合；任一子上下文漂移即整册过期。"""

    payload = {
        "kind": "unified_workbook",
        "reportBasics": report_basics_context_fingerprint(basics_report, context),
        "assessment": assessment_context_fingerprint(basics_report, context),
        "quantitativeMetrics": quantitative_metrics_context_fingerprint(
            basics_report, context
        ),
        "topicQuestions": topic_questions_context_fingerprint(
            revision_report,
            context,
            allowed_report_section_ids=allowed_report_section_ids,
        ),
        "allowedMetricKeys": (
            sorted(allowed_metric_keys) if allowed_metric_keys is not None else None
        ),
    }
    return canonical_json_sha256(payload)


def create_unified_workbook(
    basics_report: Report,
    revision_report: Report,
    *,
    context: StructuredInputContext,
    allowed_metric_keys: frozenset[str] | None = None,
    allowed_report_section_ids: frozenset[str] | None = None,
    climate_core_required: bool = False,
) -> bytes:
    """聚合四类模板为统一填报工作簿；各业务 sheet 由对应模板模块投影，不另造副本。"""

    fingerprint = unified_workbook_context_fingerprint(
        basics_report,
        revision_report,
        context,
        allowed_metric_keys=allowed_metric_keys,
        allowed_report_section_ids=allowed_report_section_ids,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_overview(workbook)
    append_report_basics_sheet(workbook, basics_report)
    append_scoring_sheet(workbook, basics_report)
    append_quantitative_sheets(
        workbook,
        basics_report,
        allowed_metric_keys=allowed_metric_keys,
        config_sheet_title=UNIFIED_QUANTITATIVE_CONFIG_SHEET,
    )
    append_topic_questions_sheets(
        workbook,
        revision_report,
        allowed_report_section_ids=allowed_report_section_ids,
        climate_core_required=climate_core_required,
    )
    append_metadata_sheet(
        workbook,
        context_metadata_values(
            input_kind="unified_workbook",
            context=context,
            context_fingerprint=fingerprint,
        ),
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _declared_metadata_report_id(workbook: openpyxl.Workbook) -> str | None:
    """读取工作簿自我声明的 report_id（模板血统判别用）；元数据页缺失时返回 None。"""

    if STRUCTURED_INPUT_METADATA_SHEET not in workbook.sheetnames:
        return None
    worksheet = workbook[STRUCTURED_INPUT_METADATA_SHEET]
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row and str(row[0] or "").strip() == "report_id":
            return "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
    return None


def split_unified_workbook(
    content: bytes,
    *,
    basics_report: Report,
    revision_report: Report,
    context: StructuredInputContext,
    allowed_metric_keys: frozenset[str] | None = None,
    allowed_report_section_ids: frozenset[str] | None = None,
    blank_expected_metadata: Mapping[str, str] | None = None,
) -> UnifiedWorkbookParts:
    """校验统一元数据并把整册拆为可被既有解析器消费的子工作簿。

    评分与定量属可选输入面：对应 sheet 全部留空时返回 None（导入跳过、保持现状）；
    基础资料与议题问题是整批替换语义（模板已预填现状），恒参与导入。

    模板血统二选一：工作簿元数据声明全零 report_id（空白母版）时按
    blank_expected_metadata 全等校验；否则按当前报告上下文全等校验。
    子工作簿元数据始终按目标报告上下文注入，与来源血统无关。
    """

    validate_xlsx_container(content)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            data_only=True,
            read_only=True,
        )
    except Exception as error:  # noqa: BLE001 — 统一投影为 typed 工作簿错误
        raise StructuredInputWorkbookError(
            [
                StructuredInputCellError(
                    code="workbook_invalid",
                    sheet="工作簿",
                    row=1,
                    column="A",
                    message=f"无法读取统一填报工作簿：{error}",
                )
            ]
        ) from error
    validate_workbook_dimensions(workbook)

    declared_report_id = _declared_metadata_report_id(workbook)
    if (
        blank_expected_metadata is not None
        and declared_report_id == str(BLANK_UNIFIED_WORKBOOK_REPORT_ID)
    ):
        expected_metadata: Mapping[str, str] = blank_expected_metadata
    else:
        fingerprint = unified_workbook_context_fingerprint(
            basics_report,
            revision_report,
            context,
            allowed_metric_keys=allowed_metric_keys,
            allowed_report_section_ids=allowed_report_section_ids,
        )
        expected_metadata = context_metadata_values(
            input_kind="unified_workbook",
            context=context,
            context_fingerprint=fingerprint,
        )
    errors = validate_metadata_sheet(workbook, expected_metadata)

    topic_sheet_titles = [
        section.title[:31]
        for section_id, section in load_topic_contract(
            knowledge_package_of(revision_report)
        ).reportSectionsById.items()
        if (
            allowed_report_section_ids is None
            or section_id in allowed_report_section_ids
        )
        and section.title[:31] in workbook.sheetnames
    ]
    required_sheets = {
        REPORT_BASICS_SHEET: "基础资料",
        SCORING_SHEET_TITLE: "重要性评分",
        UNIFIED_QUANTITATIVE_CONFIG_SHEET: "定量填写说明",
    }
    for sheet_name, label in required_sheets.items():
        if sheet_name not in workbook.sheetnames:
            errors.append(
                StructuredInputCellError(
                    code="unified_sheet_missing",
                    sheet=sheet_name,
                    row=1,
                    column="A",
                    message=f"统一填报工作簿缺少工作表：{sheet_name}（{label}）",
                )
            )
    if errors:
        raise StructuredInputWorkbookError(errors)

    quantitative_sheets = [
        sheet
        for sheet in quantitative_metric_sheets(knowledge_package_of(revision_report))
        if sheet in workbook.sheetnames
    ]

    basics_bytes = _rebuild_sub_workbook(
        workbook,
        sheet_names=[REPORT_BASICS_SHEET],
        rename={},
        metadata_values=context_metadata_values(
            input_kind="report_basics",
            context=context,
            context_fingerprint=report_basics_context_fingerprint(
                basics_report, context
            ),
        ),
    )
    scoring_bytes = (
        _rebuild_sub_workbook(
            workbook,
            sheet_names=[SCORING_SHEET_TITLE],
            rename={},
            metadata_values=context_metadata_values(
                input_kind="assessment",
                context=context,
                context_fingerprint=assessment_context_fingerprint(
                    basics_report, context
                ),
            ),
        )
        if _scoring_sheet_has_scores(workbook)
        else None
    )
    quantitative_bytes = (
        _rebuild_sub_workbook(
            workbook,
            sheet_names=[UNIFIED_QUANTITATIVE_CONFIG_SHEET, *quantitative_sheets],
            rename={UNIFIED_QUANTITATIVE_CONFIG_SHEET: QUANTITATIVE_CONFIG_SHEET},
            metadata_values=quantitative_metadata_values(
                report=basics_report,
                context=context,
                context_fingerprint=quantitative_metrics_context_fingerprint(
                    basics_report, context
                ),
            ),
        )
        if _quantitative_sheets_have_values(workbook, quantitative_sheets)
        else None
    )
    topics_bytes = _rebuild_sub_workbook(
        workbook,
        sheet_names=topic_sheet_titles,
        rename={},
        metadata_values=context_metadata_values(
            input_kind="topic_questions",
            context=context,
            context_fingerprint=topic_questions_context_fingerprint(
                revision_report,
                context,
                allowed_report_section_ids=allowed_report_section_ids,
            ),
        ),
    )
    return UnifiedWorkbookParts(
        report_basics=basics_bytes,
        scoring=scoring_bytes,
        quantitative=quantitative_bytes,
        topic_questions=topics_bytes,
    )


def _rebuild_sub_workbook(
    source,
    *,
    sheet_names: list[str],
    rename: dict[str, str],
    metadata_values: dict[str, str],
) -> bytes:
    """按坐标值拷贝目标工作表并注入对应元数据页；样式无需保真（解析只读值）。"""

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in sheet_names:
        worksheet = source[sheet_name]
        target = workbook.create_sheet(rename.get(sheet_name, sheet_name))
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            for column_index, value in enumerate(row, start=1):
                if value is not None:
                    target.cell(row_index, column_index, value)
    append_metadata_sheet(workbook, metadata_values)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _scoring_sheet_has_scores(workbook) -> bool:
    worksheet = workbook[SCORING_SHEET_TITLE]
    for row in worksheet.iter_rows(min_row=9, min_col=2, max_col=3, values_only=True):
        if any(value is not None and str(value).strip() for value in row):
            return True
    return False


def _quantitative_sheets_have_values(workbook, quantitative_sheets: list[str]) -> bool:
    config = workbook[UNIFIED_QUANTITATIVE_CONFIG_SHEET]
    for coordinate in (GHG_STANDARD_CELL, GHG_STANDARD_OTHER_CELL):
        value = config[coordinate].value
        if value is not None and str(value).strip():
            return True
    for sheet_name in quantitative_sheets:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(
            min_row=QUANTITATIVE_HEADER_ROW + 1,
            min_col=5,
            max_col=8,
            values_only=True,
        ):
            value_cell, *rest = row
            # 派生行「数值」格是模板写下的自动求和提示文本，不是用户输入。
            if (
                value_cell is not None
                and str(value_cell).strip()
                and not str(value_cell).startswith("自动求和：")
            ):
                return True
            if any(value is not None and str(value).strip() for value in rest):
                return True
    return False


def _add_overview(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet(UNIFIED_OVERVIEW_SHEET)
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 110

    worksheet["A2"] = workbook_sheet_title("统一填报工作簿")
    worksheet["A2"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=_PINE)
    worksheet.row_dimensions[2].height = 30

    instructions = (
        "本工作簿汇总报告编制所需的全部填报内容，按顺序填写即可：",
        "1.「基础资料」：企业及报告基本信息、公司简介与治理制度描述（标注「必填」的项目必须填写）。",
        "2.「重要性评分表」：为各议题填写财务与影响重要性评分；不评分可整表留空，报告将完整覆盖全部适用议题。",
        "3.「定量填写说明」与「经济+环境 / 社会 / 治理」：ESG 定量指标；不填报可整体留空，一旦填写须每项给出数值或无值原因。",
        "4. 各议题工作表：回答议题引导问题，全部选填，回答越充分报告越具体。",
        "导入时整册替换当前填写内容（评分与定量留空视为暂不提交，保持系统内现状）。请先下载最新工作簿（已预填现有内容），修改后再导入。",
    )
    for row_offset, line in enumerate(instructions):
        row_number = 4 + row_offset
        cell = worksheet.cell(row_number, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Microsoft YaHei", size=10, color=_MUTED)
        worksheet.row_dimensions[row_number].height = 26
