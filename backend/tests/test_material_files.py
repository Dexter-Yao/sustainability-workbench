# ABOUTME: 资料上传边界与确定性文档解析测试。
# ABOUTME: 验证类型伪装、压缩包资源上限、扫描页 OCR 触发与降级，及 DOCX/XLSX 的可重复文本投影。
from __future__ import annotations

from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
from pypdf import PdfReader, PdfWriter

from sustainability_desk.material.intake.files import (
    MaterialFileError,
    MaterialZipLimitError,
    validate_material_file,
)
from sustainability_desk.material.intake.parsers import (
    _clean,
    _fingerprint,
    parse_docx,
    parse_pdf_text,
    parse_pptx,
    parse_xlsx,
    pdf_has_usable_text,
)
from sustainability_desk.material.intake.scanned_pdf_ocr import OcrPageText, ScannedPdfOcrError

from test_pdf_structure import _text_pdf


def _blank_pdf_bytes(page_count: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def _merge_image_onto_first_page(data: bytes) -> bytes:
    """把一张纯色图片合并进第一页，使该页含图片资源但保留原有文本层判定。"""
    image_buffer = BytesIO()
    Image.new("RGB", (24, 24), color=(10, 90, 60)).save(image_buffer, format="PDF")
    image_page = PdfReader(BytesIO(image_buffer.getvalue())).pages[0]
    writer = PdfWriter(clone_from=BytesIO(data))
    writer.pages[0].merge_page(image_page)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _scanned_page_pdf_bytes() -> bytes:
    """一页无文本层但含图片资源的 PDF，模拟扫描件页面。"""
    return _merge_image_onto_first_page(_blank_pdf_bytes())


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("环境管理", level=1)
    document.add_paragraph("公司已建立能源台账。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "用电量"
    table.cell(1, 1).text = "120 MWh"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "能源"
    sheet.append(["指标", "数值"])
    sheet.append(["用电量", 120])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _pptx_bytes(*, with_image_only_slide: bool = False) -> bytes:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank_layout)
    textbox = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(4_000_000), Emu(900_000))
    textbox.text_frame.text = "发展历程"
    table_shape = slide.shapes.add_table(
        2, 2, Emu(0), Emu(1_000_000), Emu(3_000_000), Emu(900_000)
    )
    table = table_shape.table
    table.cell(0, 0).text = "年份"
    table.cell(0, 1).text = "事件"
    table.cell(1, 0).text = "2020"
    table.cell(1, 1).text = "公司成立"
    slide.notes_slide.notes_text_frame.text = "内部评审通过"

    if with_image_only_slide:
        from io import BytesIO as _BytesIO

        from PIL import Image

        image_bytes = _BytesIO()
        Image.new("RGB", (4, 4), color="red").save(image_bytes, format="PNG")
        image_bytes.seek(0)
        image_slide = presentation.slides.add_slide(blank_layout)
        image_slide.shapes.add_picture(image_bytes, Emu(0), Emu(0))

    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_validate_material_file_uses_extension_mime_and_magic() -> None:
    checked = validate_material_file(
        filename="能源资料.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=_docx_bytes(),
    )
    assert checked.kind == "docx"
    assert checked.size_bytes > 0
    assert len(checked.sha256) == 64

    with pytest.raises(MaterialFileError, match="文件内容与扩展名不一致"):
        validate_material_file(
            filename="伪装.pdf",
            content_type="application/pdf",
            data=_docx_bytes(),
        )

    with pytest.raises(MaterialFileError, match="转换为 \\.docx"):
        validate_material_file(filename="旧资料.doc", content_type="application/msword", data=b"old")


def test_validate_material_file_rejects_overlong_filename_first() -> None:
    with pytest.raises(MaterialFileError, match="文件名不得超过"):
        validate_material_file(
            filename=f"{'a' * 500}.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=b"PK\x03\x04",
        )


def test_validate_material_file_rejects_unsafe_zip_before_parser() -> None:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"x" * 500_000)
        archive.writestr("word/document.xml", b"<w:document/>")

    with pytest.raises(MaterialZipLimitError, match="压缩比"):
        validate_material_file(
            filename="bomb.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=stream.getvalue(),
            max_zip_ratio=10,
        )


def test_pdf_preflight_uses_real_parser_and_rejects_over_40_pages() -> None:
    known = _blank_pdf_bytes(41)
    with pytest.raises(MaterialFileError, match="不得超过 40 页"):
        validate_material_file(filename="long.pdf", content_type="application/pdf", data=known)

    scan = validate_material_file(
        filename="scan.pdf", content_type="application/pdf", data=_blank_pdf_bytes()
    )
    assert scan.pdf_page_count == 1
    assert scan.requires_attention is False
    assert pdf_has_usable_text(_blank_pdf_bytes()) is False


def test_pdf_text_parser_rejects_blank_pdf() -> None:
    with pytest.raises(ValueError, match="未提取到可读文本"):
        parse_pdf_text(_blank_pdf_bytes(), source_id=uuid4(), source_label="scan")


def test_pdf_text_parser_ocrs_image_only_page_and_merges_fragment() -> None:
    data = _scanned_page_pdf_bytes()

    with patch(
        "sustainability_desk.material.intake.parsers.ocr_pdf_pages"
    ) as mock_ocr:
        mock_ocr.return_value = {
            0: OcrPageText(page_number=1, text="识别出的扫描文字", mean_confidence=0.97)
        }
        result = parse_pdf_text(data, source_id=uuid4(), source_label="scan")

    mock_ocr.assert_called_once_with(data, [0])
    assert result.fragments[0].text == "识别出的扫描文字"
    assert result.processing_steps[0].status == "succeeded"
    assert result.processing_steps[0].quality_flags == []


def test_pdf_text_parser_degrades_when_ocr_fails() -> None:
    data = _scanned_page_pdf_bytes()

    with patch(
        "sustainability_desk.material.intake.parsers.ocr_pdf_pages"
    ) as mock_ocr:
        mock_ocr.return_value = {0: ScannedPdfOcrError("OCR 子进程超时")}
        with pytest.raises(ValueError, match="未提取到可读文本"):
            parse_pdf_text(data, source_id=uuid4(), source_label="scan")

    mock_ocr.assert_called_once_with(data, [0])


def test_pdf_text_parser_flags_warning_when_ocr_finds_no_text_but_other_pages_have_text() -> None:
    """OCR 未识别到文字时，若文档其余页仍有原生文本，解析必须以 warning 降级而非整体失败。"""
    data = _merge_image_onto_first_page(
        _text_pdf(["", "second page has native text"])
    )

    with patch(
        "sustainability_desk.material.intake.parsers.ocr_pdf_pages"
    ) as mock_ocr:
        mock_ocr.return_value = {
            0: OcrPageText(page_number=1, text="", mean_confidence=None)
        }
        result = parse_pdf_text(data, source_id=uuid4(), source_label="mixed")

    mock_ocr.assert_called_once_with(data, [0])
    assert result.processing_steps[0].status == "warning"
    assert "pdf_scanned_page_ocr_empty" in result.processing_steps[0].quality_flags
    assert [fragment.text for fragment in result.fragments] == [
        "second page has native text"
    ]


def test_parser_text_boundary_replaces_malformed_surrogates() -> None:
    cleaned = _clean("已提取\ud800文本")

    assert "\ud800" not in cleaned
    assert _fingerprint(cleaned)


def test_parse_docx_projects_paragraphs_and_tables() -> None:
    result = parse_docx(_docx_bytes(), source_id=uuid4(), source_label="A")
    assert result.source_label == "A"
    assert [fragment.kind for fragment in result.fragments] == ["text", "text", "table"]
    assert result.fragments[0].locator.kind == "docx_paragraph"
    assert "公司已建立能源台账" in result.render_text()
    assert "用电量 | 120 MWh" in result.render_text()


def test_parse_xlsx_projects_sheet_rows_without_formula_execution() -> None:
    result = parse_xlsx(_xlsx_bytes(), source_id=uuid4(), source_label="B")
    assert result.fragments[0].locator.kind == "xlsx_range"
    assert result.fragments[0].locator.cell_range == "A1:B2"
    assert result.fragments[0].rows[1] == ["用电量", "120"]
    assert result.render_text() == "指标 | 数值\n用电量 | 120"


def test_validate_material_file_accepts_pptx_and_rejects_legacy_ppt() -> None:
    checked = validate_material_file(
        filename="发展历程.pptx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
        data=_pptx_bytes(),
    )
    assert checked.kind == "pptx"

    with pytest.raises(MaterialFileError, match="转换为 \\.pptx"):
        validate_material_file(
            filename="旧演示.ppt",
            content_type="application/vnd.ms-powerpoint",
            data=b"old",
        )


def test_parse_pptx_projects_slide_text_table_notes_and_image_placeholder() -> None:
    result = parse_pptx(
        _pptx_bytes(with_image_only_slide=True), source_id=uuid4(), source_label="C"
    )
    assert result.source_label == "C"
    assert [fragment.kind for fragment in result.fragments] == [
        "text",
        "table",
        "text",
        "text",
    ]
    assert all(
        fragment.locator.kind == "pptx_slide" for fragment in result.fragments
    )
    assert result.fragments[0].text == "发展历程"
    assert result.fragments[1].rows[1] == ["2020", "公司成立"]
    assert result.fragments[2].text == "备注：内部评审通过"
    assert result.fragments[3].text == "（本页为图片内容，未提取文字）"
