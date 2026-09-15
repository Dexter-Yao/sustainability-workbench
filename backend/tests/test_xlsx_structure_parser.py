# ABOUTME: 验证 XLSX 结构解析保留绝对位置、公式、空行、隐藏/合并状态与对象覆盖率。
# ABOUTME: 边界测试锁定 20,000 行和 100,000 单元格无截断，并在越界前拒绝展开。
from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image
import pytest

from sustainability_desk.material.intake import xlsx_structure
from sustainability_desk.material.intake.parsed_material import (
    ParsedChartNode,
    ParsedImageNode,
    ParsedNamedRangeNode,
    ParsedSheetNode,
    ParsedSheetRowNode,
    ParsedTableNode,
)
from sustainability_desk.material.intake.xlsx_structure import (
    MAX_XLSX_CELLS,
    MAX_XLSX_ROWS,
    XlsxStructureParseError,
    parse_xlsx_structure,
)

_LONG_CELL_TEXT = "供应链风险" * 6_000


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _structured_workbook() -> bytes:
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Visible"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    very_hidden = workbook.create_sheet("VeryHidden")
    very_hidden.sheet_state = "veryHidden"

    visible["A1"] = "指标"
    visible["B1"] = "=1+1"
    visible["B1"].number_format = "0.00"
    visible["D1"] = "类别"
    visible["E1"] = "数值"
    visible["D2"] = "范围一"
    visible["E2"] = 12
    visible["A4"] = "合并内容"
    visible["C4"] = _LONG_CELL_TEXT
    visible.merge_cells("A4:B4")
    visible.row_dimensions[4].hidden = True
    visible.column_dimensions["B"].hidden = True
    visible["A1"].comment = Comment("批注意见", "审阅人")

    table = Table(displayName="ClimateTable", ref="D1:E2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    visible.add_table(table)
    workbook.defined_names.add(
        DefinedName("ClimateRange", attr_text="'Visible'!$D$1:$E$2")
    )

    chart = BarChart()
    chart.title = "定量指标图"
    chart.add_data(
        Reference(visible, min_col=5, min_row=1, max_row=2),
        titles_from_data=True,
    )
    chart.set_categories(Reference(visible, min_col=4, min_row=2, max_row=2))
    visible.add_chart(chart, "G1")

    image_buffer = BytesIO()
    Image.new("RGB", (2, 2), color="green").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    visible.add_image(WorkbookImage(image_buffer), "G5")
    return _workbook_bytes(workbook)


def test_parser_preserves_xlsx_structure_without_executing_formulas() -> None:
    parsed = parse_xlsx_structure(
        _structured_workbook(),
        source_id=uuid4(),
        source_label="结构测试.xlsx",
    )

    sheets = [node for node in parsed.nodes if isinstance(node, ParsedSheetNode)]
    assert [(node.label, node.sheet_state) for node in sheets] == [
        ("Visible", "visible"),
        ("Hidden", "hidden"),
        ("VeryHidden", "veryHidden"),
    ]
    assert sheets[0].cell_range == "A1:E4"
    assert sheets[0].row_count == 4
    assert sheets[0].column_count == 5
    assert sheets[0].cell_count == 20
    assert len(sheets[0].row_node_ids) == 4
    rows = [
        node
        for node in parsed.nodes
        if isinstance(node, ParsedSheetRowNode)
        and node.locator.sheet_name == "Visible"
    ]
    assert [row.row_index for row in rows] == [1, 2, 3, 4]
    assert all(cell.value_kind == "blank" for cell in rows[2].cells)

    formula = rows[0].cells[1]
    assert formula.cell_reference == "B1"
    assert formula.value_kind == "formula"
    assert formula.formula == "=1+1"
    assert formula.text == "=1+1"
    assert formula.cached_value is None
    assert formula.native_data_type == "f"
    assert formula.number_format == "0.00"
    assert formula.hidden_column is True

    anchor = rows[3].cells[0]
    merged_child = rows[3].cells[1]
    assert rows[3].hidden is True
    assert anchor.hidden_row is True
    assert anchor.merged_anchor == "A4"
    assert anchor.merged_range == "A4:B4"
    assert anchor.column_span == 2
    assert merged_child.merged_anchor == "A4"
    assert merged_child.merged_range == "A4:B4"
    assert rows[3].cells[2].text == _LONG_CELL_TEXT
    assert len(rows[3].cells[2].text) > 20_000
    assert rows[1].cells[4].value_kind == "number"
    assert rows[1].cells[4].native_data_type == "n"
    assert rows[1].cells[4].text == "12"

    tables = [node for node in parsed.nodes if isinstance(node, ParsedTableNode)]
    named_ranges = [
        node for node in parsed.nodes if isinstance(node, ParsedNamedRangeNode)
    ]
    charts = [node for node in parsed.nodes if isinstance(node, ParsedChartNode)]
    images = [node for node in parsed.nodes if isinstance(node, ParsedImageNode)]
    assert [(node.title, node.locator.cell_range) for node in tables] == [
        ("ClimateTable", "D1:E2")
    ]
    assert [(node.name, node.definition) for node in named_ranges] == [
        ("ClimateRange", "'Visible'!$D$1:$E$2")
    ]
    assert charts[0].title == "定量指标图"
    assert any(
        reference.replace("$", "").endswith("!E2")
        for reference in charts[0].source_ranges
    )
    assert len(images) == 1
    assert images[0].media_sha256 is not None
    assert images[0].media_extraction_status == "complete"

    capabilities = {
        capability.capability: capability
        for capability in parsed.coverage.capabilities
    }
    assert capabilities["defined_tables"].source_unit_count == 1
    assert capabilities["defined_tables"].parsed_unit_count == 1
    assert capabilities["named_ranges"].parsed_unit_count == 1
    assert capabilities["charts"].parsed_unit_count == 1
    assert capabilities["images"].source_unit_count == 1
    assert capabilities["images"].parsed_unit_count == 1
    assert capabilities["comments"].source_unit_count == 1
    assert capabilities["comments"].parsed_unit_count == 0
    assert parsed.coverage.disposition == "incomplete_usable"
    assert parsed.coverage.allows_negative_evidence is False


def test_failed_image_byte_extraction_is_explicit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(xlsx_structure, "_image_sha256", lambda image: None)

    parsed = parse_xlsx_structure(
        _structured_workbook(),
        source_id=uuid4(),
        source_label="图片提取失败.xlsx",
    )

    images = [node for node in parsed.nodes if isinstance(node, ParsedImageNode)]
    image_coverage = next(
        capability
        for capability in parsed.coverage.capabilities
        if capability.capability == "images"
    )
    assert images[0].media_sha256 is None
    assert images[0].media_extraction_status == "failed"
    assert image_coverage.source_unit_count == 1
    assert image_coverage.parsed_unit_count == 0
    assert any(
        gap.code == "xlsx_image_resource_extraction_failed"
        for gap in parsed.coverage.gaps
    )


def _boundary_workbook(*, row: int, column: int) -> bytes:
    workbook = Workbook()
    workbook.active.cell(row=row, column=column, value="边界末端")
    return _workbook_bytes(workbook)


def test_column_formatting_beyond_content_does_not_expand_cell_extent() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    for row_index in range(1, 96):
        for column_index in range(1, 7):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=f"{row_index}:{column_index}",
            )
    formatted_columns = worksheet.column_dimensions["O"]
    formatted_columns.min = 15
    formatted_columns.max = 16_384
    formatted_columns.width = 12
    formatted_columns.hidden = True

    parsed = parse_xlsx_structure(
        _workbook_bytes(workbook),
        source_id=uuid4(),
        source_label="远端列格式.xlsx",
    )

    rows = [node for node in parsed.nodes if isinstance(node, ParsedSheetRowNode)]
    sheets = [node for node in parsed.nodes if isinstance(node, ParsedSheetNode)]
    assert sheets[0].cell_range == "A1:F95"
    assert sheets[0].row_count == 95
    assert sheets[0].column_count == 6
    assert sheets[0].cell_count == 570
    assert sheets[0].row_node_ids == tuple(row.node_id for row in rows)
    assert len(rows) == 95
    assert sum(len(row.cells) for row in rows) == 570
    assert rows[-1].cells[-1].cell_reference == "F95"
    assert parsed.coverage.table_cells_parsed == 570
    assert all(
        not cell.hidden_column for row in rows for cell in row.cells
    )


def test_formula_cache_scan_tolerates_sparse_empty_cells() -> None:
    workbook = Workbook()
    workbook.active["E2"] = "=SUM(A1:A2)"

    parsed = parse_xlsx_structure(
        _workbook_bytes(workbook),
        source_id=uuid4(),
        source_label="稀疏公式.xlsx",
    )

    rows = [node for node in parsed.nodes if isinstance(node, ParsedSheetRowNode)]
    formula = rows[1].cells[4]
    assert formula.cell_reference == "E2"
    assert formula.value_kind == "formula"
    assert formula.formula == "=SUM(A1:A2)"
    assert formula.cached_value is None


def test_parser_accepts_twenty_thousand_rows_and_one_hundred_thousand_cells() -> None:
    parsed = parse_xlsx_structure(
        _boundary_workbook(row=MAX_XLSX_ROWS, column=5),
        source_id=uuid4(),
        source_label="边界.xlsx",
    )

    rows = [node for node in parsed.nodes if isinstance(node, ParsedSheetRowNode)]
    assert len(rows) == MAX_XLSX_ROWS
    assert sum(len(row.cells) for row in rows) == MAX_XLSX_CELLS
    assert rows[-1].row_index == MAX_XLSX_ROWS
    assert rows[-1].cells[-1].cell_reference == "E20000"
    assert rows[-1].cells[-1].text == "边界末端"
    assert parsed.coverage.table_cells_parsed == MAX_XLSX_CELLS
    assert parsed.status == "complete"


@pytest.mark.parametrize(
    ("row", "column", "message"),
    [
        (MAX_XLSX_ROWS + 1, 1, "原始行数超过"),
        (10_001, 10, "原始单元格位置超过"),
        (95, 16_384, "原始单元格位置超过"),
    ],
)
def test_parser_rejects_workbooks_beyond_row_or_cell_bounds(
    row: int,
    column: int,
    message: str,
) -> None:
    with pytest.raises(XlsxStructureParseError, match=message):
        parse_xlsx_structure(
            _boundary_workbook(row=row, column=column),
            source_id=uuid4(),
            source_label="越界.xlsx",
        )
