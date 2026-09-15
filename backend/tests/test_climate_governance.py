# ABOUTME: 气候治理支柱测试——planner 据评估把 climate_change 挂入环境大章、guard、渲染保真、callout 零泄漏。
from pathlib import Path
from uuid import uuid4

from docx import Document

from sustainability_desk.contract.fill import fill_report
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import AssessmentResult
from sustainability_desk.export.docx_renderer import load_values_flat, render_docx
from sustainability_desk.planner import assemble_report, load_topic_intake
from assessment_fixtures import complete_assessment
from knowledge_package_fixtures import SSE_PACKAGE
from sustainability_desk.planner import load_topic_templates_from

BACKEND = Path(__file__).resolve().parents[1]
CONTRACT = SSE_PACKAGE.report_contract_path
TOPIC_DIR = SSE_PACKAGE.topic_sections_dir
VALUES = SSE_PACKAGE.sample_values_path


def _assess(materiality: str) -> AssessmentResult:
    return complete_assessment(
        load_contract(CONTRACT),
        materialities={"climate_change": materiality},
    )


def test_climate_template_loads_with_topic_id():
    """climate_change.yaml 加载为 contentScopeId=climate_change 的模板。"""
    templates = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)
    assert "climate_change" in templates
    assert templates["climate_change"].headingLevel == 2


def test_climate_assembled_into_environmental_module():
    """climate_change 实质 → 挂入环境可持续模块，含 climate.gov 治理支柱。"""
    contract = load_contract(CONTRACT)
    templates = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)
    intake = load_topic_intake(SSE_PACKAGE)
    res = assemble_report(contract, templates, assessment=_assess("dual"), intake_items=intake)
    env = next(s for s in res.report.sections if s.key == "environmental_sustainability")
    climate = next(c for c in (env.children or []) if c.key == "climate_change")
    assert any(c.key == "climate.gov" for c in (climate.children or []))


def test_climate_non_material_still_appears():
    """climate_change 非重要 → 仍作为适用议题进入环境大章。"""
    contract = load_contract(CONTRACT)
    templates = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)
    res = assemble_report(contract, templates, assessment=_assess("non"))
    env = next(s for s in res.report.sections if s.key == "environmental_sustainability")
    assert any(c.key == "climate_change" for c in (env.children or []))


def test_climate_governance_renders_without_callout_leak(base_template, out_dir):
    """治理支柱渲染：H3「（一）治理」+ 叙述正文在；准则/溯源 callout 不入正文。"""
    contract = load_contract(CONTRACT)
    templates = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)
    intake = load_topic_intake(SSE_PACKAGE)
    assembled = assemble_report(contract, templates, assessment=_assess("dual"), intake_items=intake).report
    report = fill_report(assembled, load_values_flat(VALUES))
    out = render_docx(report, base_template, out_dir / "climate_check.docx")
    doc = Document(str(out))
    head_style = {p.text: p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")}
    full = "\n".join(p.text for p in doc.paragraphs)
    assert head_style.get("（一）治理") == "Heading 3"
    # climate.gov_regulatory_note 准则说明块已随 callout 概念移除；其独有片段不应出现在正文。
    # about.basis 固定段含「ISSB」属合法正文，故不能用裸字符串断言；改为检查该块专有前缀。
    assert "依据 ISSB《国际财务报告可持续披露准则第2号" not in full, "气候准则说明不应进入正文"


def test_prompt_config_includes_climate_block_identity_only():
    """/api/prompt-config 只投影前端实际消费的可生成块身份。"""
    from sustainability_desk.api.app import prompt_config
    from sustainability_desk.api.auth import AuthenticatedUser

    # 该端点要求登录；此处只验证投影内容，身份取一个已验证主体即可。
    cfg = prompt_config(AuthenticatedUser(subject=uuid4(), email=None))
    blocks = {b["id"]: b for b in cfg["blocks"]}
    # 该端点无包参数：它按设计投影账户级默认包，而默认包随目标市场变动（气候治理块在
    # sse_zh_hans 叫 climate.gov_structure，在 hkex 两包叫 climate.gov_oversight）。
    # 断言的是投影形态——只出块身份，不出别的——故从默认包自身取块，不硬写某包的块名。
    assert blocks, "prompt-config 未投影任何可生成块"
    assert all(block == {"id": block_id} for block_id, block in blocks.items())
    assert "intake_items" not in cfg
    # 准则批注是大陆三所特性（用户在三个交易所间选 disclosureProfile）；港交所披露依据固定，
    # 其包内 entries 恒为空。故只断言键存在，具体条目由 SSE 包自己的批注测试覆盖。
    assert isinstance(cfg["user_visible_disclosure_clause_annotations"], list)


def test_generation_template_includes_topic_blocks():
    """generate 路径的 _template(SSE_PACKAGE) 须含 topic_sections 块，否则议题块 find_block 抛 KeyError、无法生成。"""
    from sustainability_desk.llm.generate import _template

    block = _template(SSE_PACKAGE).find_block("climate.gov_structure")  # 不应抛 KeyError
    assert block.blockType == "generative"


def test_climate_governance_posture_switches_between_facts_and_high_level_goal():
    """build_model_context 据清单作答状态切换：有内容组织事实，无内容使用统一高层目标。"""
    from sustainability_desk.llm.generate import _template
    from sustainability_desk.llm.prompts import build_model_context
    tmpl = _template(SSE_PACKAGE)
    block = tmpl.find_block("climate.gov_structure")
    templates = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)
    intake = load_topic_intake(SSE_PACKAGE)
    vals = load_values_flat(VALUES)

    # 有内容（填了清单 Q1）
    base_h = assemble_report(load_contract(CONTRACT), templates, assessment=_assess("dual"), intake_items=intake).report
    inst_have = fill_report(base_h, vals)
    ctx_have = build_model_context(block, tmpl, inst_have)
    assert ctx_have.evidence.intake_facts, "填了清单应注入 materials"
    assert ctx_have.evidence_posture.level == "block_facts"

    # 没内容（去掉 intake 作答）
    base_n = assemble_report(load_contract(CONTRACT), templates, assessment=_assess("dual"), intake_items=intake).report
    inst_none = fill_report(base_n, {k: vals[k] for k in vals if k != "intake"})
    ctx_none = build_model_context(block, tmpl, inst_none)
    assert not ctx_none.evidence.intake_facts, "没内容 materials 应为空"
    assert ctx_none.evidence_posture.level == "context_only"
    assert not ctx_none.writing_granularity
