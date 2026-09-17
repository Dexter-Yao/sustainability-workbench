# ABOUTME: schema→docx 渲染器测试——前四章实例渲染、标题分级、字段填值、条件显隐、无残留。
from datetime import datetime, timezone
import re
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from sustainability_desk.contract.loader import load_contract
from sustainability_desk.export.toc import page_layout_renderer
from sustainability_desk.contract.models import (
    CustomEngagementMethod,
    ExternalAssuranceReport,
    ReaderFeedbackContactInformation,
)
from sustainability_desk.contract.export_json import export_instance_json
from sustainability_desk.contract.stakeholder_engagement import (
    apply_stakeholder_engagement_projection,
    reconcile_stakeholder_engagement_profile,
)
from sustainability_desk.export.docx_renderer import (
    build_document_render_plan,
    render_docx,
    render_final_docx,
)
from sustainability_desk.export.format_profile import load_format_profile
from sustainability_desk.export.normalize_template import normalize_template
from knowledge_package_fixtures import SSE_PACKAGE


def _heading_style(head_style: dict[str, str], suffix: str) -> str | None:
    """按标题主体（去编号前缀）匹配 Heading 样式——编号 label 由 number_sections 注入，断言不锁死前缀。"""
    return next((style for text, style in head_style.items() if text.endswith(suffix)), None)


def test_render_full_instance(template_docx, sample_values_yaml, out_dir):
    base = normalize_template(template_docx, out_dir / "base.docx", profile=load_format_profile(SSE_PACKAGE))
    instance_path = export_instance_json(SSE_PACKAGE, out_dir / "instance.json")
    report = load_contract(instance_path)
    report = report.model_copy(
        update={
            "stakeholderEngagement": reconcile_stakeholder_engagement_profile(report)
        }
    )
    report = apply_stakeholder_engagement_projection(report)
    partner = next(
        entry
        for entry in report.stakeholderEngagement.entries
        if entry.stakeholderType == "partners"
    )
    partner.customMethods.append(
        CustomEngagementMethod(kind="collaboration_activity", label="区域协作计划")
    )
    out = render_final_docx(report, base, out_dir / "report_instance.docx")

    doc = Document(str(out))
    head_style = {p.text: p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")}
    full = "\n".join(p.text for p in doc.paragraphs)

    # 前章标题齐全（编号前缀由 number_sections 注入，按标题主体匹配）
    for h in ("关于本报告", "公司治理", "可持续发展治理"):
        assert _heading_style(head_style, h) is not None, f"缺章标题 {h}"
    # 可持续发展治理的三个固定 H2 直接平级渲染。
    assert _heading_style(head_style, "可持续发展治理架构") == "Heading 2"
    assert _heading_style(head_style, "利益相关方沟通") == "Heading 2"
    assert _heading_style(head_style, "议题重要性评估") == "Heading 2"
    assert _heading_style(head_style, "组织架构设置") is None
    assert _heading_style(head_style, "组织各层级主要职责分工") is None
    # 标题内字段引用（titleContent）插值为实例值「关于晟原精密」，而非纯文本骨架「关于公司」
    assert _heading_style(head_style, "关于晟原精密") == "Heading 1"
    assert "关于公司" not in full, "标题骨架不应出现在渲染结果"
    # 空标题分组节（承载审议批准段）不得产生空标题段
    assert all(p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading"))
    assert "审议并批准发布" in full
    assert "本报告于2026年4月经由董事会审议并批准发布。" in full
    assert "2026年2026-04月" not in full
    # 字段填入实例值（非占位）
    assert "晟原精密电子股份有限公司" in full
    assert "【" not in full, "存在未解析字段占位"

    # 源模板可能携带未使用的图表及其本地 XLSX 外部关系；交付物不得保留，
    # 否则 Word 打开时会提示更新其他文件的字段。
    with ZipFile(out) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        # 目录终稿只允许局部替换 PAGEREF 的缓存。不得经通用 XML serializer
        # 改写 Word 的命名空间前缀，否则 mc:Ignorable 会引用不存在的前缀，
        # Microsoft Word 将修复文档并破坏目录锚点。
        assert "<w:document" in document_xml
        assert "<ns0:document" not in document_xml
        declared_prefixes = set(re.findall(r'xmlns:([A-Za-z0-9_]+)=', document_xml))
        ignorable = re.search(r'mc:Ignorable="([^"]+)"', document_xml)
        assert ignorable is not None
        assert set(ignorable.group(1).split()) <= declared_prefixes
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            rels = archive.read(name).decode("utf-8")
            assert 'TargetMode="External"' not in rels, name
        assert "word/charts/chart1.xml" not in archive.namelist()
        # 交付件一律禁止 Word 打开时刷新**全部**字段：那会弹出“字段可能引用其他
        # 文件”的通用安全提示。目录项的 PAGEREF 指向同文档书签，不在此列。
        assert "w:updateFields" not in settings_xml
        # 装了版式引擎时页码已写进展示缓存并清掉脏标记；没装则保留脏标记，
        # 由阅读器打开时据书签自行解析。两种都是正确终态，故按可用性断言——
        # 写死其中一种会让这条用例在另一种环境里假失败。
        if page_layout_renderer() is not None:
            assert 'w:dirty="true"' not in document_xml
        else:
            assert 'w:dirty="true"' in document_xml
    sdts = doc.element.body.findall(qn("w:sdt"))
    toc_sdt = next(
        sdt
        for sdt in sdts
        if any(
            "PAGEREF GS_TOC_" in (node.text or "")
            for node in sdt.findall(".//" + qn("w:instrText"))
        )
    )
    cover_text = "".join(
        node.text or "" for node in doc.element.body.findall(".//" + qn("w:t"))
    )
    toc_text = "".join(node.text or "" for node in toc_sdt.findall(".//" + qn("w:t")))
    toc_entry_styles = {
        paragraph.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val"))
        for paragraph in toc_sdt.findall(".//" + qn("w:p"))
        if paragraph.find(qn("w:pPr")) is not None
        and paragraph.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
    }
    assert "晟原精密电子股份有限公司" in cover_text
    assert "可持续发展报告" in cover_text
    assert "2025年度" in cover_text
    assert "202X" not in cover_text and "用户输入年份" not in cover_text
    assert doc.core_properties.title == "2025年度可持续发展报告"
    topic_headings = {
        heading
        for heading, style in head_style.items()
        if style in {"Heading 1", "Heading 2"}
    }
    assert all(heading in toc_text for heading in topic_headings)
    assert toc_entry_styles <= {"TOC", "TOC1", "TOC2", "TOCHeading"}
    toc_entries = [
        paragraph
        for paragraph in toc_sdt.findall(".//" + qn("w:p"))
        if paragraph.find(qn("w:pPr")) is not None
        and paragraph.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
        and paragraph.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val"))
        in {"TOC1", "TOC2"}
    ]
    pageref_instructions = [
        "".join(node.text or "" for node in entry.findall(".//" + qn("w:instrText")))
        for entry in toc_entries
    ]
    assert len(toc_entries) == len(topic_headings)
    assert all(instruction.startswith(" PAGEREF GS_TOC_") for instruction in pageref_instructions)
    toc_page_values = [
        (entry.findall(".//" + qn("w:t"))[-1].text or "")
        for entry in toc_entries
    ]
    # 封面和目录不占用正文页码；“关于本报告”是正文节的第 1 页。
    assert toc_page_values[0] == "1"
    assert all(value.isdigit() for value in toc_page_values)
    toc_bookmarks = {
        marker.get(qn("w:name"))
        for marker in doc.element.findall(".//" + qn("w:bookmarkStart"))
        if (marker.get(qn("w:name")) or "").startswith("GS_TOC_")
    }
    assert {
        instruction.split()[1]
        for instruction in pageref_instructions
    } == toc_bookmarks
    assert "用生态韧性赋能可持续运营" not in toc_text
    assert "可持续发展战略" in toc_text
    assert len(doc.sections) == 2
    # 正式稿与审阅稿一样全篇承载署名行（封面、目录、正文），封面仍不显示页码。
    cover_footer = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "Sustainability Workbench" in cover_footer
    assert "github.com/Dexter-Yao/sustainability-workbench" in cover_footer
    assert "第 " not in cover_footer
    page_number_type = doc.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert page_number_type is not None
    assert page_number_type.get(qn("w:start")) == "1"

    # 条件显隐：无鉴证→鉴证段隐藏；无目标表→目标节隐藏
    assert "报告鉴证" not in full
    assert "可持续发展目标与关键进展" not in full
    # 无残留模板标记
    assert "{{" not in full and "{%" not in full
    # 备注（元指令）不进正式 Word（归 Plate 引导）
    assert "可多做几个版本" not in full and "（说明：" not in full
    # 表格渲染（荣誉/利益相关方/议题/IRO，目标表隐藏）
    assert len(doc.tables) >= 3
    stakeholder_tables = [
        table
        for table in doc.tables
        if table.rows and table.rows[0].cells[0].text == "利益相关方"
    ]
    assert len(stakeholder_tables) == 1
    stakeholder_text = "\n".join(
        cell.text
        for row in stakeholder_tables[0].rows
        for cell in row.cells
    )
    assert "社区与公众" in stakeholder_text
    assert "供应商及合作伙伴" not in stakeholder_text
    assert "供应商" in stakeholder_text and "合作伙伴" in stakeholder_text
    assert "数据安全与客户隐私保护" in stakeholder_text
    assert "反商业贿赂与反贪污" in stakeholder_text
    assert "区域协作计划" in stakeholder_text
    # 有序/无序列表渲染为 Word 编号（numPr），非 unicode 符号
    listed = [
        p for p in doc.paragraphs
        if (pPr := p._p.find(qn("w:pPr"))) is not None and pPr.find(qn("w:numPr")) is not None
    ]
    assert len(listed) >= 6, f"列表项过少: {len(listed)}"
    # 瘦身
    assert out.stat().st_size < 1_000_000


def test_table_list_cell_rendered_with_pause_mark(base_template, out_dir):
    """多选单元格（list）按顿号渲染，不出现 Python 列表字面量 ['..']。"""

    from docx import Document

    from sustainability_desk.contract.loader import load_contract
    from sustainability_desk.contract.models import Block, GsColDef, GsTable, Section
    from sustainability_desk.contract.table_ops import data_row
    from sustainability_desk.export.docx_renderer import render_docx

    contract = SSE_PACKAGE.report_contract_path
    report = load_contract(contract)
    cols = [GsColDef(key="vc", header="价值链", cellType="multi_select", options=["上游价值链", "公司运营"])]
    blk = Block(id="probe.tbl", type="table", blockType="generative", source="ai",
                table=GsTable(colDefs=cols,
                              children=[data_row(cols, {"vc": ["上游价值链", "公司运营"]}, state="ready")]))
    report.sections.append(Section(key="probe", title="探针", headingLevel=1, blocks=[blk]))
    out = render_docx(report, base_template, out_dir / "list_cell.docx")
    full = "\n".join(p.text for t in Document(str(out)).tables for r in t.rows for c in r.cells for p in c.paragraphs)
    assert "上游价值链、公司运营" in full
    assert "['" not in full and "']" not in full


def test_customer_comment_docx_uses_real_comments_without_changing_body(
    base_template,
    out_dir,
):
    """批注版与普通版同源渲染；说明进入 comments.xml 而非正文。"""
    from sustainability_desk.contract.models import Block, Inline, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        sections=[
            Section(
                key="sample",
                title="测试章节",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sample.body",
                        type="paragraph",
                        blockType="generative",
                        source="ai",
                        content=[Inline(kind="text", text="第一段内容。\n第二段内容。")],
                    )
                ],
            )
        ],
    )
    plan = build_document_render_plan(report)
    normal = render_docx(report, base_template, out_dir / "normal.docx", render_plan=plan)
    annotated = render_docx(
        report,
        base_template,
        out_dir / "annotated.docx",
        render_plan=plan,
        comment_text_by_anchor={anchor: f"说明：{anchor}" for anchor in plan.unit_anchor_ids},
        delivery_variant="review",
        review_generated_at=datetime.now(timezone.utc),
    )
    normal_text = "\n".join(p.text for p in Document(normal).paragraphs)
    annotated_text = "\n".join(p.text for p in Document(annotated).paragraphs)
    for expected_body in ("测试章节", "第一段内容。", "第二段内容。"):
        assert expected_body in normal_text
        assert expected_body in annotated_text
    assert "审阅版" not in annotated_text
    # 审阅稿封面不放品牌字标图片，文档里不应有嵌入图。
    assert len(Document(annotated).inline_shapes) == 0
    with ZipFile(annotated) as archive:
        assert "word/comments.xml" in archive.namelist()
        assert "说明：block:sample.body:paragraph:0" in archive.read("word/comments.xml").decode("utf-8")
    with ZipFile(normal) as archive:
        assert "word/comments.xml" not in archive.namelist()


def test_review_cover_anchors_report_overview_and_uses_frozen_footer(
    base_template,
    out_dir,
):
    """审阅版整体说明只附着在封面企业名称，页脚读取固定 revision 时间。"""
    from lxml import etree

    from sustainability_desk.contract.models import Block, Field, Inline, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        fields={
            "company_registered_name": Field(
                key="company_registered_name",
                label="公司注册名称",
                type="string",
                source="user_input",
                value="测试企业股份有限公司",
            ),
            "reporting_year": Field(
                key="reporting_year",
                label="报告年度",
                type="string",
                source="user_input",
                value="2025",
            ),
        },
        sections=[
            Section(
                key="sample",
                title="测试章节",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sample.body",
                        type="paragraph",
                        blockType="generative",
                        source="ai",
                        content=[Inline(kind="text", text="正文第一段。")],
                    )
                ],
            )
        ],
    )
    generated_at = datetime(2026, 8, 1, 12, 34, tzinfo=timezone.utc)
    output = render_final_docx(
        report,
        base_template,
        out_dir / "review.docx",
        render_plan=build_document_render_plan(report),
        comment_text_by_anchor={"block:sample.body:paragraph:0": "正文说明。"},
        delivery_variant="review",
        review_generated_at=generated_at,
        cover_comment_text="本次报告说明：本报告基于已提供信息形成。",
    )
    document = Document(output)
    assert len(document.sections) == 2
    # 制作说明是出处声明，封面与目录节同样承载（读者翻开第一页即可见）；
    # 但这些节不显示页码，页码只属于正文节。
    front_footer_text = "\n".join(
        paragraph.text for paragraph in document.sections[0].footer.paragraphs
    )
    assert "本报告由开源工具 Sustainability Workbench 根据用户提供的资料编制" in front_footer_text
    assert "生成于 2026年08月01日" in front_footer_text
    assert "第 " not in front_footer_text
    footer_text = "\n".join(
        paragraph.text for paragraph in document.sections[-1].footer.paragraphs
    )
    assert "本报告由开源工具 Sustainability Workbench 根据用户提供的资料编制" in footer_text
    assert "生成于 2026年08月01日" in footer_text
    # 页脚是署名，不写时分——日期足以分辨版本；「贵司」类 To B 敬语一律不用。
    assert "20:34" not in footer_text
    assert "贵司" not in footer_text
    assert "气候议题试用版" not in footer_text
    assert "中国标准时间" not in footer_text
    assert len([paragraph for paragraph in document.sections[-1].footer.paragraphs if paragraph.text.strip()]) == 1
    page_number_type = document.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert page_number_type is not None
    assert page_number_type.get(qn("w:start")) == "1"

    title_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "可持续发展报告"
    )
    assert title_paragraph._p.find(qn("w:pPr")).find(qn("w:pBdr")) is None

    with ZipFile(output) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        comments_xml = etree.fromstring(archive.read("word/comments.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    overview_comment = next(
        item
        for item in comments_xml.xpath("//w:comment", namespaces=namespaces)
        if "本次报告说明" in "".join(item.xpath(".//w:t/text()", namespaces=namespaces))
    )
    overview_id = overview_comment.get(qn("w:id"))
    company_paragraph = next(
        paragraph
        for paragraph in document_xml.xpath("//w:p", namespaces=namespaces)
        if "测试企业股份有限公司" in "".join(paragraph.xpath(".//w:t/text()", namespaces=namespaces))
    )
    assert company_paragraph.xpath(
        f".//w:commentRangeStart[@w:id='{overview_id}']", namespaces=namespaces
    )
    body_paragraph = next(
        paragraph
        for paragraph in document_xml.xpath("//w:p", namespaces=namespaces)
        if "正文第一段。" in "".join(paragraph.xpath(".//w:t/text()", namespaces=namespaces))
    )
    assert not body_paragraph.xpath(
        f".//w:commentRangeStart[@w:id='{overview_id}']", namespaces=namespaces
    )
def test_unbacked_template_image_slot_is_omitted_from_public_word(
    base_template,
    out_dir,
):
    """输入提示占位不得泄露到对外交付 Word。"""
    from sustainability_desk.contract.models import Block, ImageModel, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="图片 slot 测试",
        sections=[
            Section(
                key="probe",
                title="探针",
                headingLevel=1,
                blocks=[
                    Block(
                        id="probe.image_slot",
                        type="image",
                        blockType="slot",
                        source="template",
                        image=ImageModel(
                            caption="治理架构图",
                            placeholder="可在后续版本补充。",
                        ),
                    )
                ],
            )
        ],
    )

    plan = build_document_render_plan(report)
    output = render_docx(
        report,
        base_template,
        out_dir / "unbacked-image-slot.docx",
        render_plan=plan,
    )

    assert not any(unit.block_id == "probe.image_slot" for unit in plan.units)
    full_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "图片占位" not in full_text
    assert "治理架构图" not in full_text


def test_assessment_projection_does_not_overwrite_generated_iro_rows(base_template, out_dir):
    """Word 必须投影 revision 已保存的 IRO 表，不能以空 assessment.iroItems 覆盖它。"""
    from sustainability_desk.contract.models import (
        AssessmentResult,
        Block,
        GsColDef,
        GsTable,
        Report,
        ScoredAssessmentResult,
        Section,
    )
    from sustainability_desk.contract.models import RowExpansion, RowExpansionUnit
    from sustainability_desk.contract.table_ops import build_expanded_rows

    columns = [
        GsColDef(key="iro_topic", header="议题", cellType="text"),
        GsColDef(key="iro_desc", header="影响、风险与机遇描述", cellType="ai_text"),
        GsColDef(
            key="iro_class",
            header="分类",
            cellType="multi_select",
            options=["潜在正面影响", "潜在负面影响", "机遇", "风险"],
        ),
        GsColDef(key="iro_value_chain", header="影响范围", cellType="multi_select"),
        GsColDef(key="iro_time_horizon", header="影响周期", cellType="multi_select"),
    ]
    expansion = RowExpansion(
        sharedColumnKeys=["iro_topic", "iro_value_chain", "iro_time_horizon"],
        units=[
            RowExpansionUnit(key="impact", label="影响描述", columnKeys=["iro_desc", "iro_class"]),
            RowExpansionUnit(
                key="risk_opportunity",
                label="风险与/或机遇影响描述",
                columnKeys=["iro_desc", "iro_class"],
            ),
        ],
    )
    generated_rows = build_expanded_rows(
        columns,
        expansion,
        [
            {
                "iro_topic": "应对气候变化",
                "iro_value_chain": ["公司运营", "下游价值链"],
                "iro_time_horizon": ["中期", "长期"],
                "impact": {
                    "iro_desc": "制造运营可能带来潜在气候影响。",
                    "iro_class": ["潜在负面影响"],
                },
                "risk_opportunity": {
                    "iro_desc": "低碳产品需求增长带来潜在市场机遇。",
                    "iro_class": ["机遇"],
                },
            }
        ],
        state="ready",
    )
    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="IRO 投影回归测试",
        assessment=AssessmentResult(
            reportingYear=2025,
            topics=[
                ScoredAssessmentResult(
                    determination="scored",
                    assessmentTopicId="climate_change",
                    materiality="dual",
                    financialScore=4.8,
                    impactScore=4.6,
                )
            ],
        ),
        sections=[
            Section(
                key="probe",
                title="探针",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sm.iro_table",
                        type="table",
                        blockType="constrained",
                        source="ai",
                        state="ready",
                        table=GsTable(
                            caption="财务/双重重要性议题的影响、风险与机遇",
                            rowSource="assessment_iro",
                            colDefs=columns,
                            children=generated_rows,
                        ),
                    )
                ],
            )
        ],
    )

    output = render_docx(report, base_template, out_dir / "generated-iro.docx")
    table = next(
        table
        for table in Document(str(output)).tables
        if table.rows and table.rows[0].cells[0].text == "应对气候变化"
    )
    table_text = "\n".join(
        cell.text for row in table.rows for cell in row.cells
    )
    assert "制造运营可能带来潜在气候影响。" in table_text
    assert "低碳产品需求增长带来潜在市场机遇。" in table_text
    assert "公司运营、下游价值链" in table_text


def test_omitted_table_is_not_rendered_as_an_empty_word_table(base_template, out_dir):
    """已持久化的合同省略必须在 Word 中消失，不留下只有表头的空壳。"""
    from sustainability_desk.contract.models import Block, GsColDef, GsTable, GsTableCell, GsTableRow, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="t",
        sections=[Section(
            key="probe",
            title="探针",
            headingLevel=1,
            blocks=[Block(
                id="probe.omitted_table",
                type="table",
                blockType="generative",
                source="ai",
                state="omitted",
                table=GsTable(
                    caption="不应显示的空表",
                    colDefs=[GsColDef(key="name", header="名称", cellType="text")],
                    children=[GsTableRow(
                        headerRow=True,
                        children=[GsTableCell(type="th", colKey="name", value="名称")],
                    )],
                ),
            )],
        )],
    )

    output = render_docx(report, base_template, out_dir / "omitted-table.docx")
    document = Document(str(output))
    full_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )

    assert "不应显示的空表" not in full_text
    assert "名称" not in full_text


def test_empty_kpi_projection_omits_its_heading_caption_and_table(base_template, out_dir):
    """无定量值时 KPI 是整体省略，不在附录留下空标题或仅表头。"""
    from sustainability_desk.contract.loader import load_contract
    from sustainability_desk.export.docx_renderer import render_docx

    contract = SSE_PACKAGE.report_contract_path
    output = render_docx(
        load_contract(contract),
        base_template,
        out_dir / "empty-kpi.docx",
    )
    document = Document(str(output))
    heading_text = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    table_text = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]

    assert not any("ESG关键绩效表" in text for text in heading_text)
    assert "ESG 关键绩效表" not in table_text


def test_appendix_and_each_direct_child_start_on_new_page(base_template, out_dir):
    """附录标题与直属附录子节均分页，避免标题孤立在正文末页。"""
    from sustainability_desk.contract.models import Block, Inline, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="附录分页测试",
        sections=[
            Section(
                key="body",
                title="正文",
                headingLevel=1,
                blocks=[
                    Block(
                        id="body.text",
                        type="paragraph",
                        blockType="fixed",
                        source="template",
                        content=[Inline(kind="text", text="正文内容")],
                    )
                ],
            ),
            Section(
                key="report_appendix",
                title="附录",
                headingLevel=1,
                children=[
                    Section(
                        key="appendix.metrics",
                        title="一、关键绩效表",
                        headingLevel=2,
                        blocks=[
                            Block(
                                id="appendix.metrics.text",
                                type="paragraph",
                                blockType="fixed",
                                source="template",
                                content=[Inline(kind="text", text="指标内容")],
                            )
                        ],
                    ),
                    Section(
                        key="appendix.index",
                        title="二、索引表",
                        headingLevel=2,
                        blocks=[
                            Block(
                                id="appendix.index.text",
                                type="paragraph",
                                blockType="fixed",
                                source="template",
                                content=[Inline(kind="text", text="索引内容")],
                            )
                        ],
                    ),
                ],
            ),
        ],
    )

    output = render_docx(report, base_template, out_dir / "appendix-page-breaks.docx")
    headings = [
        paragraph
        for paragraph in Document(str(output)).paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2"}
    ]

    def heading(suffix: str):
        return next(paragraph for paragraph in headings if paragraph.text.endswith(suffix))

    assert heading("附录").paragraph_format.page_break_before is True
    assert heading("关键绩效表").paragraph_format.page_break_before is True
    assert heading("索引表").paragraph_format.page_break_before is True


def test_risk_matrix_table_exports_fixed_widths_and_repeating_header(base_template, out_dir):
    """风险矩阵表导出固定列宽，并为表头设置跨页重复。"""
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET

    from sustainability_desk.contract.loader import load_contract
    from sustainability_desk.contract.models import Block, GsColDef, GsTable, GsTableCell, GsTableRow, Section
    from sustainability_desk.contract.table_ops import data_row
    from sustainability_desk.export.docx_renderer import render_docx

    contract = SSE_PACKAGE.report_contract_path
    report = load_contract(contract)
    cols = [
        GsColDef(key="risk_category", header="类别", cellType="text"),
        GsColDef(key="risk_name", header="类型", cellType="text"),
        GsColDef(key="risk_description", header="具体内容", cellType="ai_text"),
        GsColDef(key="financial_impact", header="财务影响", cellType="multi_select", options=["成本增加"]),
        GsColDef(key="time_horizon", header="时间", cellType="multi_select", options=["短期"]),
        GsColDef(key="response_measures", header="应对措施", cellType="ai_text"),
    ]
    header = GsTableRow(headerRow=True, children=[
        GsTableCell(type="th", value="类别"),
        GsTableCell(type="th", value="类型"),
        GsTableCell(type="th", value="具体内容"),
        GsTableCell(type="th", value="财务影响"),
        GsTableCell(type="th", value="时间"),
        GsTableCell(type="th", value="应对措施"),
    ])
    blk = Block(id="probe.risk_matrix", type="table", blockType="generative", source="ai",
                table=GsTable(
                    layoutProfile="risk_response_matrix",
                    columnWidthWeights={
                        "risk_category": 9,
                        "risk_name": 12,
                        "risk_description": 26,
                        "financial_impact": 13,
                        "time_horizon": 10,
                        "response_measures": 30,
                    },
                    colDefs=cols,
                    children=[header, data_row(cols, {
                        "risk_category": "风险",
                        "risk_name": "法规政策风险",
                        "risk_description": "法规政策变化可能带来合规管理压力。",
                        "financial_impact": ["成本增加"],
                        "time_horizon": ["短期"],
                        "response_measures": "公司持续跟踪法规要求并完善日常管理。",
                    }, state="ready")],
                ))
    report.sections.append(Section(key="probe", title="探针", headingLevel=1, blocks=[blk]))

    out = render_docx(report, base_template, out_dir / "risk_matrix.docx")

    with ZipFile(out) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert root.findall(".//w:tblLayout[@w:type='fixed']", ns)
    assert root.findall(".//w:tblHeader", ns)
    grid_widths = [node.attrib[f"{{{ns['w']}}}w"] for node in root.findall(".//w:tblGrid/w:gridCol", ns)]
    assert len(set(grid_widths)) > 1
    rfonts = root.findall(".//w:rFonts", ns)
    assert any(node.attrib.get(f"{{{ns['w']}}}eastAsia") == "宋体" for node in rfonts)
    assert any(node.attrib.get(f"{{{ns['w']}}}ascii") == "SimSun" for node in rfonts)


def test_iro_table_uses_declared_widths_and_keeps_business_rows_intact(base_template, out_dir):
    """IRO 表的长文本列应获足够宽度，且业务行不得在跨页时留下空白续行。"""
    import xml.etree.ElementTree as ET
    from zipfile import ZipFile

    from sustainability_desk.contract.models import Report, Section
    from sustainability_desk.contract.table_ops import data_row

    contract = SSE_PACKAGE.report_contract_path
    source = load_contract(contract)
    block = source.find_block("sm.iro_table").model_copy(deep=True)
    assert block.table is not None
    assert block.table.columnWidthWeights is not None
    assert block.table.columnWidthWeights["iro_desc"] > block.table.columnWidthWeights["iro_topic"]
    # 描述列的写作口径下沉到各披露子行（影响 / 风险机遇），列级不再承载统一 genHint。
    expansion = block.generation.rowExpansion
    assert block.table.colDefs[1].genHint is None
    hints = {unit.key: unit.genHintOverride["iro_desc"] for unit in expansion.units}
    assert "对外部环境与社会造成的影响" in hints["impact"]
    assert "对公司自身构成的风险或带来的机遇" in hints["risk_opportunity"]
    block.table.children.append(
        data_row(
            block.table.colDefs,
            {
                "iro_topic": "应对气候变化",
                "iro_desc": "制造过程中的能源使用与转型要求可能带来潜在影响、风险与机遇。",
                "iro_class": ["潜在负面影响", "风险", "机遇"],
                "iro_value_chain": ["上游价值链", "公司运营", "下游价值链"],
                "iro_time_horizon": ["短期", "中期", "长期"],
            },
            state="ready",
        )
    )
    # IRO 表以「存在财务/双重重要性议题」为显隐条件，导出探针须提供满足该条件的评估结果。
    from sustainability_desk.contract.models import AssessmentResult, ScoredAssessmentResult

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="IRO 导出版式测试",
        assessment=AssessmentResult(
            reportingYear=2025,
            topics=[
                ScoredAssessmentResult(
                    determination="scored",
                    assessmentTopicId="climate_change",
                    materiality="dual",
                    financialScore=4.6,
                    impactScore=4.5,
                )
            ],
        ),
        sections=[Section(key="probe", title="探针", headingLevel=1, blocks=[block])],
    )
    output = render_docx(report, base_template, out_dir / "iro-layout.docx")

    with ZipFile(output) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    table = next(
        table
        for table in root.findall(".//w:tbl", ns)
        if "应对气候变化" in "".join(table.itertext())
    )
    widths = [int(node.attrib[f"{{{ns['w']}}}w"]) for node in table.findall("./w:tblGrid/w:gridCol", ns)]
    assert widths[1] > widths[0]
    rows = table.findall("./w:tr", ns)
    assert rows[-1].find("./w:trPr/w:cantSplit", ns) is not None

def test_optional_appendix_chapters_hidden_when_values_missing(template_docx, sample_values_yaml, out_dir):
    """读者反馈/外部鉴证纯选填：缺值整章省略，部分填写只渲染已填行。"""
    base = normalize_template(template_docx, out_dir / "base.docx", profile=load_format_profile(SSE_PACKAGE))
    instance_path = export_instance_json(SSE_PACKAGE, out_dir / "instance.json")
    report = load_contract(instance_path)
    report = report.model_copy(
        update={"stakeholderEngagement": reconcile_stakeholder_engagement_profile(report)}
    )
    report = apply_stakeholder_engagement_projection(report)

    report.appendixPackage.readerFeedbackContactInformation = ReaderFeedbackContactInformation()
    report.appendixPackage.externalAssuranceReport = ExternalAssuranceReport(isIncluded=True)
    out = render_final_docx(report, base, out_dir / "appendix_hidden.docx")
    doc = Document(str(out))
    head_style = {p.text: p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")}
    paras = [p.text for p in doc.paragraphs]
    assert _heading_style(head_style, "读者反馈") is None
    assert _heading_style(head_style, "外部鉴证声明") is None
    assert not any(t.startswith("联系邮箱") for t in paras)

    report.appendixPackage.readerFeedbackContactInformation = ReaderFeedbackContactInformation(
        email="esg@example.com"
    )
    out = render_final_docx(report, base, out_dir / "appendix_partial.docx")
    doc = Document(str(out))
    head_style = {p.text: p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")}
    paras = [p.text for p in doc.paragraphs]
    assert _heading_style(head_style, "读者反馈") is not None
    assert "联系邮箱：esg@example.com" in paras
    assert not any(t.startswith("公司地址：") for t in paras)
    assert not any(t.startswith("联系电话：") for t in paras)



def _notice_fixture():
    """构造一份最小资料处理说明（不经冻结谱系，直接给渲染器）。"""
    from sustainability_desk.export.format_profile import load_format_profile
    from sustainability_desk.report_review_packages import (
        MaterialProcessingNotice,
        MaterialProcessingNoticeFile,
        MaterialProcessingNoticeItem,
    )

    notice_texts = load_format_profile(SSE_PACKAGE).delivery_texts.material_notice
    MATERIAL_PROCESSING_NOTICE_HEADING = notice_texts.heading
    MATERIAL_PROCESSING_NOTICE_LEAD = notice_texts.lead
    MATERIAL_ATTENTION_NOTICE_HEADING = notice_texts.files_heading
    MATERIAL_ATTENTION_NOTICE_LEAD = notice_texts.files_lead

    return MaterialProcessingNotice(
        heading=MATERIAL_PROCESSING_NOTICE_HEADING,
        lead=MATERIAL_PROCESSING_NOTICE_LEAD,
        files_heading=MATERIAL_ATTENTION_NOTICE_HEADING,
        files_lead=MATERIAL_ATTENTION_NOTICE_LEAD,
        files=(
            MaterialProcessingNoticeFile(
                material_name="产品研发与创新部分资料整理.docx",
                items=(
                    MaterialProcessingNoticeItem(
                        message="研发人员占比数据存在显著差异。",
                        next_action="需与企业确认准确数据口径。",
                    ),
                ),
            ),
        ),
    )


def test_review_notice_page_sits_before_toc_without_taking_chapter_number(
    template_docx, sample_values_yaml, out_dir
) -> None:
    """说明页位于封面之后、目录之前，且不占章号、不进目录。"""
    base_template = normalize_template(template_docx, out_dir / "notice-base.docx", profile=load_format_profile(SSE_PACKAGE))
    report = load_contract(export_instance_json(SSE_PACKAGE, out_dir / "notice.json"))
    plan = build_document_render_plan(report)
    generated_at = datetime(2026, 8, 1, 12, 34, tzinfo=timezone.utc)
    output = render_final_docx(
        report,
        base_template,
        out_dir / "notice-review.docx",
        render_plan=plan,
        comment_text_by_anchor={anchor: "说明。" for anchor in plan.unit_anchor_ids},
        delivery_variant="review",
        review_generated_at=generated_at,
        cover_comment_text="本次报告说明。",
        material_processing_notice=_notice_fixture(),
    )
    document = Document(output)
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert "资料处理说明" in texts
    assert "产品研发与创新部分资料整理.docx" in texts
    assert any(text.startswith("1. 研发人员占比数据存在显著差异") for text in texts)
    assert any(text.startswith("提示") for text in texts)
    # 资料核对事项是说明页的小节，排在说明页引导句之后。
    assert "资料核对事项" in texts
    assert texts.index("资料处理说明") < texts.index("资料核对事项")
    body_children = list(document.element.body.iterchildren())
    toc_position = next(
        index for index, child in enumerate(body_children) if child.tag == qn("w:sdt")
    )
    notice_paragraph = next(
        paragraph for paragraph in document.paragraphs if paragraph.text == "资料核对事项"
    )
    assert body_children.index(notice_paragraph._p) < toc_position

    notice_index = texts.index("资料处理说明")
    first_chapter = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "关于本报告"
    )
    # 顺序：封面 → 说明页 → 目录 → 正文第一章。
    assert texts.index("可持续发展报告") < notice_index < first_chapter
    body_children = list(document.element.body.iterchildren())
    toc_position = next(
        index
        for index, child in enumerate(body_children)
        if child.tag == qn("w:sdt")
    )
    notice_position = next(
        index
        for index, child in enumerate(body_children)
        if child is document.paragraphs[notice_index]._p
    )
    assert notice_position < toc_position
    # 说明页标题不是 Heading 样式，故不进目录域、不占「第N章」编号。
    notice_paragraph = document.paragraphs[notice_index]
    assert not notice_paragraph.style.name.startswith("Heading")
    assert notice_paragraph._p.find(qn("w:pPr")) is None or (
        notice_paragraph._p.find(qn("w:pPr")).find(qn("w:numPr")) is None
    )


def _compliance_notice_fixture():
    from sustainability_desk.export.format_profile import load_format_profile
    from sustainability_desk.report_review_packages import (
        StandardsComplianceNotice,
        StandardsComplianceNoticeGroup,
        StandardsComplianceNoticeItem,
    )

    standards_texts = load_format_profile(SSE_PACKAGE).delivery_texts.standards_notice
    STANDARDS_COMPLIANCE_NOTICE_HEADING = standards_texts.heading
    STANDARDS_COMPLIANCE_NOTICE_LEAD = standards_texts.lead

    return StandardsComplianceNotice(
        heading=STANDARDS_COMPLIANCE_NOTICE_HEADING,
        lead=STANDARDS_COMPLIANCE_NOTICE_LEAD,
        summary="本报告已覆盖 150 项准则应披露要求。",
        groups=(
            StandardsComplianceNoticeGroup(
                heading="建议留意事项",
                lead="以下事项在定稿前值得关注，可通过补充资料或补充说明处理。",
                items=(
                    StandardsComplianceNoticeItem(
                        scope_label="水资源管理",
                        message="「水资源节约目标以及具体措施情况」因未获得对应资料，本报告未作披露。",
                        next_action="补充相应资料后重新生成，或在定稿时补充一句未披露原因的说明。",
                    ),
                ),
            ),
            StandardsComplianceNoticeGroup(
                heading="准则允许的省略",
                lead="以下内容依据准则规定未纳入本报告，列出供了解，无须处理。",
                items=(
                    StandardsComplianceNoticeItem(
                        scope_label="应对气候变化",
                        message="气候情景分析属准则鼓励披露事项，本报告未开展情景分析，未作相关披露。",
                    ),
                ),
            ),
        ),
    )


def test_standards_compliance_notice_follows_material_notice_before_toc(
    template_docx, sample_values_yaml, out_dir
) -> None:
    """准则对照说明页排在资料处理说明页之后、目录之前，且不占章号、不进目录。"""
    import pytest

    base_template = normalize_template(template_docx, out_dir / "compliance-base.docx", profile=load_format_profile(SSE_PACKAGE))
    report = load_contract(
        export_instance_json(SSE_PACKAGE, out_dir / "compliance.json")
    )
    plan = build_document_render_plan(report)
    output = render_final_docx(
        report,
        base_template,
        out_dir / "compliance-review.docx",
        render_plan=plan,
        delivery_variant="review",
        review_generated_at=datetime(2026, 9, 1, 8, 0, tzinfo=timezone.utc),
        material_processing_notice=_notice_fixture(),
        standards_compliance_notice=_compliance_notice_fixture(),
    )
    document = Document(output)
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert "准则对照说明" in texts
    assert any("不构成鉴证意见或法律意见" in text for text in texts)
    assert any(text.startswith("1. 【水资源管理】") for text in texts)
    assert any("气候情景分析属准则鼓励披露事项" in text for text in texts)

    # 顺序：封面 → 资料处理说明 → 准则对照说明 → 目录 → 正文第一章。
    first_chapter = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "关于本报告"
    )
    assert texts.index("资料处理说明") < texts.index("准则对照说明") < first_chapter

    body_children = list(document.element.body.iterchildren())
    toc_position = next(
        index for index, child in enumerate(body_children) if child.tag == qn("w:sdt")
    )
    compliance_paragraph = next(
        paragraph for paragraph in document.paragraphs if paragraph.text == "准则对照说明"
    )
    assert body_children.index(compliance_paragraph._p) < toc_position
    # 标题不是 Heading 样式，故不进目录域、不占「第N章」编号。
    assert not compliance_paragraph.style.name.startswith("Heading")

    # 两个说明页同属一个版式家族：段内行距均须显式压到条目行距，
    # 不得继承正文 1.5 倍——否则续行与条目间距等宽，一条事项的续行会被读成下一条。
    profile = load_format_profile(SSE_PACKAGE)
    notice_line_spacing = profile.material_processing_notice.line_spacing_multiple
    assert notice_line_spacing < profile.body.line_spacing_multiple
    for heading_text in ("资料处理说明", "准则对照说明"):
        page_start = texts.index(heading_text)
        page_paragraphs = [
            paragraph
            for paragraph in document.paragraphs[page_start:]
            if paragraph.text.strip()
        ][:6]
        for paragraph in page_paragraphs:
            assert paragraph.paragraph_format.line_spacing == pytest.approx(
                notice_line_spacing
            ), f"{heading_text}：{paragraph.text[:20]}"


def test_standard_docx_rejects_standards_compliance_notice(
    template_docx, sample_values_yaml, out_dir
) -> None:
    """正式稿绝不携带准则对照说明页。"""
    import pytest

    base_template = normalize_template(template_docx, out_dir / "reject-compliance-base.docx", profile=load_format_profile(SSE_PACKAGE))
    report = load_contract(
        export_instance_json(SSE_PACKAGE, out_dir / "reject-compliance.json")
    )
    with pytest.raises(ValueError, match="准则对照说明页"):
        render_final_docx(
            report,
            base_template,
            out_dir / "reject-compliance.docx",
            standards_compliance_notice=_compliance_notice_fixture(),
        )


def test_standard_docx_rejects_material_processing_notice(
    template_docx, sample_values_yaml, out_dir
) -> None:
    """正式稿绝不携带说明页——与批注、审阅页脚同一道守卫。"""
    import pytest

    base_template = normalize_template(template_docx, out_dir / "reject-base.docx", profile=load_format_profile(SSE_PACKAGE))
    report = load_contract(export_instance_json(SSE_PACKAGE, out_dir / "reject.json"))
    with pytest.raises(ValueError, match="资料处理说明页"):
        render_final_docx(
            report,
            base_template,
            out_dir / "reject-standard.docx",
            material_processing_notice=_notice_fixture(),
        )


def test_footnote_lands_in_footnotes_part_and_is_identical_across_both_drafts(
    base_template,
    out_dir,
):
    """脚注保真：正文只留上标引用，说明文本进 footnotes.xml；正式稿与审阅稿必须一字不差。

    两稿由同一 DocumentRenderPlan 渲染，脚注若在其中一稿漂移，读者拿到的限定条件
    就与审阅时看到的不一致——这是交付物层面的事实分叉，非样式问题。
    """
    from zipfile import ZipFile

    from sustainability_desk.contract.models import Block, Inline, Report, Section

    note = "鉴于公司主营业务未涉及科技伦理敏感领域，经评估，科技伦理不构成本年度实质性议题。"
    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        sections=[
            Section(
                key="sample",
                title="测试章节",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sample.coverage",
                        type="paragraph",
                        blockType="fixed",
                        source="template",
                        content=[Inline(kind="text", text="本年度报告合计披露 21 项可持续发展议题。")],
                        footnote=[Inline(kind="text", text=note)],
                    )
                ],
            )
        ],
    )
    plan = build_document_render_plan(report)
    standard = render_docx(report, base_template, out_dir / "fn_standard.docx", render_plan=plan)
    review = render_docx(
        report,
        base_template,
        out_dir / "fn_review.docx",
        render_plan=plan,
        delivery_variant="review",
        review_generated_at=datetime.now(timezone.utc),
    )

    for produced in (standard, review):
        body = "\n".join(p.text for p in Document(produced).paragraphs)
        assert "本年度报告合计披露 21 项可持续发展议题。" in body
        assert note not in body, "脚注说明不得降级为正文段落"
        with ZipFile(produced) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            footnotes_xml = archive.read("word/footnotes.xml").decode("utf-8")
        assert document_xml.count("<w:footnoteReference") == 1
        assert note in footnotes_xml
        assert note not in document_xml

    with ZipFile(standard) as archive:
        standard_footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    with ZipFile(review) as archive:
        review_footnotes = archive.read("word/footnotes.xml").decode("utf-8")
    assert standard_footnotes == review_footnotes, "两稿脚注不一致"


def test_block_without_footnote_declaration_adds_no_reference(base_template, out_dir):
    """未声明脚注的块不得凭空产生引用——孤立上标在交付物里是可见缺陷。"""
    from zipfile import ZipFile

    from sustainability_desk.contract.models import Block, Inline, Report, Section

    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        sections=[
            Section(
                key="sample",
                title="测试章节",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sample.body",
                        type="paragraph",
                        blockType="fixed",
                        source="template",
                        content=[Inline(kind="text", text="普通正文段落。")],
                    )
                ],
            )
        ],
    )
    produced = render_docx(report, base_template, out_dir / "fn_none.docx")
    with ZipFile(produced) as archive:
        assert "<w:footnoteReference" not in archive.read("word/document.xml").decode("utf-8")


def _ai_disclosure_report():
    """构造承载 AI 标识断言的最小报告：封面只需企业名与报告年度。"""
    from sustainability_desk.contract.models import Block, Field, Inline, Report, Section

    return Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        fields={
            "company_registered_name": Field(
                key="company_registered_name",
                label="公司注册名称",
                type="string",
                source="user_input",
                value="测试企业股份有限公司",
            ),
            "reporting_year": Field(
                key="reporting_year",
                label="报告年度",
                type="string",
                source="user_input",
                value="2025",
            ),
        },
        sections=[
            Section(
                key="sample",
                title="测试章节",
                headingLevel=1,
                blocks=[
                    Block(
                        id="sample.body",
                        type="paragraph",
                        blockType="generative",
                        source="ai",
                        content=[Inline(kind="text", text="正文第一段。")],
                    )
                ],
            )
        ],
    )


def test_every_delivery_variant_carries_ai_disclosure_on_cover_and_in_metadata(
    base_template,
    out_dir,
):
    """两种交付形态一律承载 AI 生成合成内容标识。

    《人工智能生成合成内容标识办法》（2025-09-01 施行）第四条要求显式标识、第五条要求
    文件元数据隐式标识。标识义务看内容是否 AI 生成，与稿件类型（正式稿/审阅稿）无关，
    故两种形态逐一断言——版式调整时不得把这一行悄悄删掉。
    """
    profile = load_format_profile(SSE_PACKAGE)
    disclosure_text = profile.cover_full_report.ai_disclosure_text
    report = _ai_disclosure_report()
    generated_at = datetime(2026, 8, 1, 12, 34, tzinfo=timezone.utc)

    for delivery_variant in ("standard", "review"):
        output = render_final_docx(
            report,
            base_template,
            out_dir / f"disclosure-{delivery_variant}.docx",
            render_plan=build_document_render_plan(report),
            delivery_variant=delivery_variant,
            review_generated_at=(
                generated_at if delivery_variant == "review" else None
            ),
        )
        document = Document(output)
        body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert disclosure_text in body_text, (
            f"{delivery_variant} 封面缺少显式 AI 标识"
        )
        # 隐式标识三要素：生成合成标签、服务提供者、内容制作编号。
        comments = document.core_properties.comments or ""
        assert comments.startswith("AIGC;"), (
            f"{delivery_variant} 元数据缺少生成合成标签"
        )
        assert "Sustainability Workbench" in comments
        assert "测试企业股份有限公司-2025" in comments
        # 内部标识不进交付物：编号由封面可见事实派生，不含 report_id。
        assert "report_id" not in comments


def test_toc_page_numbers_resolve_without_a_layout_engine(base_template, out_dir):
    """没装 LibreOffice 时导出仍给出可用目录：PAGEREF 带 dirty 标记由阅读器自算页码。

    LibreOffice 曾是硬依赖，只为**预先算好**目录页码——缺它直接 TocFinalizationError。
    对自用场景，为这点确定性要求一整套 800 MB 办公套件不划算：目录项是指向同文档书签的
    PAGEREF 域，Word 与 LibreOffice 打开时会自行解析（实测页码与预计算结果逐条一致）。

    本用例钉住降级路径的两个前提：缺引擎时不抛错，且域上留有 dirty 标记——
    少了标记，阅读器会把占位的「1」当成已算好的页码，于是每条目录都显示第 1 页，
    那比报错更糟：它看起来是对的。
    """
    import sustainability_desk.export.docx_renderer as renderer_module

    report = _ai_disclosure_report()
    out_path = out_dir / "no-layout-engine.docx"

    with patch.object(renderer_module, "page_layout_renderer", return_value=None):
        renderer_module.render_final_docx(report, base_template, out_path)

    with ZipFile(out_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    page_refs = re.findall(r"PAGEREF GS_TOC", document_xml)
    assert page_refs, "目录项应带 PAGEREF 域"
    assert document_xml.count('w:dirty="true"') >= len(page_refs), (
        "每个目录 PAGEREF 都要带 dirty 标记，否则阅读器会沿用占位页码"
    )
