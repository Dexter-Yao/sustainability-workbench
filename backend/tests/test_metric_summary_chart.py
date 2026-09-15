# ABOUTME: ESG 定量指标摘要图测试——验证用户填写值可确定性渲染、空值隐藏、Word 导出嵌入。
# ABOUTME: 覆盖报告年度展示和 0 值有效性，避免把真实零值当作未填写。
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import Inches

from sustainability_desk.api.auth import AuthenticatedUser
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import (
    Block,
    DerivedVisualizationSpec,
    Field,
    ImageModel,
    QuantitativeMetricDraft,
    QuantitativeMetricsMeta,
    Report,
    ReportMeta,
    Section,
)
from sustainability_desk.export.docx_renderer import render_docx
from sustainability_desk.export.metric_summary_chart import (
    build_metric_summary_layout,
    filled_metric_entries,
    render_quantitative_metric_summary,
)
from knowledge_package_fixtures import SSE_PACKAGE, SSE_REPORT_PROFILE_ID

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
_CONTRACT = SSE_PACKAGE.report_contract_path
_METRIC_KEYS = ["economic_environment_r04", "economic_environment_r05", "economic_environment_r07"]
_NINE_METRIC_KEYS = [f"economic_environment_r{row:02d}" for row in range(4, 13)]


def _report(metrics: dict[str, str | None]) -> Report:
    return Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试报告",
        fields={
            "reporting_year": Field(
                key="reporting_year",
                label="报告年度",
                type="year",
                source="user_input",
                value=2025,
            )
        },
        sections=[],
        meta=ReportMeta(
            quantitativeMetrics=QuantitativeMetricsMeta(
                metrics={
                    key: QuantitativeMetricDraft(value=value)
                    for key, value in metrics.items()
                }
            )
        ),
    )


def test_metric_summary_renders_valid_png():
    png = render_quantitative_metric_summary(
        _report({
            "economic_environment_r04": "120.5",
            "economic_environment_r05": "86.2",
            "economic_environment_r07": "40.1",
        }),
        _METRIC_KEYS,
    )

    assert png is not None
    assert png[:8] == _PNG_MAGIC
    assert len(png) > 2000


def test_metric_summary_treats_zero_as_filled_value():
    entries = filled_metric_entries(
        _report({"economic_environment_r04": "0"}),
        ["economic_environment_r04"],
    )

    assert len(entries) == 1
    assert entries[0].value == "0"


def test_metric_summary_disambiguates_other_metric_label():
    entries = filled_metric_entries(
        _report({"economic_environment_r17": "3.2"}),
        ["economic_environment_r17"],
    )

    assert len(entries) == 1
    assert entries[0].label != "其他"
    assert entries[0].label.endswith("：其他")


def test_metric_summary_returns_none_when_all_values_empty():
    assert render_quantitative_metric_summary(
        _report({"economic_environment_r04": ""}),
        ["economic_environment_r04"],
    ) is None


def test_metric_summary_omits_unfilled_metrics_from_partially_filled_image():
    entries = filled_metric_entries(
        _report({
            "economic_environment_r04": "120.5",
            "economic_environment_r05": None,
            "economic_environment_r07": "40.1",
        }),
        _METRIC_KEYS,
    )

    assert [entry.value for entry in entries] == ["120.5", "40.1"]


@pytest.mark.parametrize(
    ("count", "expected_mode"),
    [(1, "highlight_cards"), (3, "highlight_cards"), (6, "compact_cards"), (9, "summary_table")],
)
def test_metric_summary_layout_parses_count_mode_year_and_leaf_labels(count, expected_mode):
    keys = _NINE_METRIC_KEYS[:count]
    layout = build_metric_summary_layout(
        _report({key: str(index + 1) for index, key in enumerate(keys)}),
        keys,
    )

    assert layout is not None
    assert layout.yearLabel == "2025年"
    assert layout.mode == expected_mode
    assert len(layout.entries) == count
    assert all(entry.group == "" for entry in layout.entries)
    assert all("｜" not in entry.label for entry in layout.entries)


@pytest.mark.parametrize("count", [1, 3, 6, 9])
def test_metric_summary_draws_report_year_once(monkeypatch, count):
    import sustainability_desk.export.metric_summary_chart as chart

    calls: list[str] = []
    original = chart._draw_year

    def record_year(ax, year_label, **kwargs):
        calls.append(year_label)
        return original(ax, year_label, **kwargs)

    monkeypatch.setattr(chart, "_draw_year", record_year)
    keys = _NINE_METRIC_KEYS[:count]
    png = chart.render_quantitative_metric_summary(
        _report({key: str(index + 1) for index, key in enumerate(keys)}),
        keys,
    )

    assert png is not None
    assert calls == ["2025年"]


def test_metric_summary_normalizes_unfilled_sentinel_at_report_boundary():
    report = _report({"economic_environment_r04": "未填写"})

    assert report.meta.quantitativeMetrics.metrics["economic_environment_r04"].value is None
    assert filled_metric_entries(report, ["economic_environment_r04"]) == []


def test_api_metric_summary_image_returns_png():
    import sustainability_desk.api.app as APP

    # economic_environment_r04 只存在于 SSE 指标目录；不绑包时端点按服务端默认 Profile
    # 取目录，该默认随目标市场变动，指标不在目录内即渲染为空并返回 204。
    req = APP.MetricSummaryImageRequest(
        report=_report({"economic_environment_r04": "120.5"}),
        spec=DerivedVisualizationSpec(
            kind="quantitative_metric_summary",
            metricKeys=["economic_environment_r04"],
        ),
        report_profile_id=SSE_REPORT_PROFILE_ID,
    )

    res = APP.quantitative_metric_summary_image(
        req, AuthenticatedUser(subject=uuid4(), email=None)
    )

    assert res.media_type == "image/png"
    assert res.body[:8] == _PNG_MAGIC


def test_api_metric_summary_image_returns_204_for_empty_values():
    import sustainability_desk.api.app as APP

    req = APP.MetricSummaryImageRequest(
        report=_report({"economic_environment_r04": ""}),
        spec=DerivedVisualizationSpec(
            kind="quantitative_metric_summary",
            metricKeys=["economic_environment_r04"],
        ),
        # 同上绑包：不绑时「指标不在默认目录」会顶替「值为空」满足下面这条断言。
        report_profile_id=SSE_REPORT_PROFILE_ID,
    )

    res = APP.quantitative_metric_summary_image(
        req, AuthenticatedUser(subject=uuid4(), email=None)
    )

    assert res.status_code == 204


def test_metric_summary_embedded_as_inline_image(base_template, out_dir):
    report = load_contract(_CONTRACT)
    report.fields["reporting_year"] = Field(
        key="reporting_year",
        label="报告年度",
        type="year",
        source="user_input",
        value=2025,
    )
    report.meta = ReportMeta(
        quantitativeMetrics=QuantitativeMetricsMeta(
            metrics={"economic_environment_r04": QuantitativeMetricDraft(value="120.5")}
        )
    )
    report.sections.append(
        Section(
            key="metric_probe",
            title="指标探针",
            headingLevel=1,
            blocks=[
                Block(
                    id="metric_probe.image",
                    type="image",
                    blockType="fixed",
                    source="derived",
                    image=ImageModel(
                        derivedVisualization=DerivedVisualizationSpec(
                            kind="quantitative_metric_summary",
                            metricKeys=["economic_environment_r04"],
                        )
                    ),
                )
            ],
        )
    )

    out = render_docx(report, base_template, out_dir / "metric_summary.docx")

    inline_shapes = Document(str(out)).inline_shapes

    assert len(inline_shapes) >= 1
    assert Inches(4.0) <= inline_shapes[-1].width <= Inches(6.5)
    assert inline_shapes[-1].height > Inches(0.5)


def test_metric_summary_does_not_consume_figure_numbering(base_template, out_dir):
    """指标摘要图不占「图N」、不进 figure_specs；
    其后的真图（素材图等）编号顺延不受影响，仍从 图1 起编。"""
    from io import BytesIO

    from PIL import Image

    from sustainability_desk.export.docx_renderer import ResolvedEvidenceImage
    from sustainability_desk.export.figure_projection import FigureSpec

    report = load_contract(_CONTRACT)
    report.fields["reporting_year"] = Field(
        key="reporting_year",
        label="报告年度",
        type="year",
        source="user_input",
        value=2025,
    )
    report.meta = ReportMeta(
        quantitativeMetrics=QuantitativeMetricsMeta(
            metrics={"economic_environment_r04": QuantitativeMetricDraft(value="120.5")}
        )
    )
    report.sections.append(
        Section(
            key="metric_probe",
            title="指标探针",
            headingLevel=1,
            blocks=[
                Block(
                    id="metric_probe.summary",
                    type="image",
                    blockType="fixed",
                    source="derived",
                    image=ImageModel(
                        derivedVisualization=DerivedVisualizationSpec(
                            kind="quantitative_metric_summary",
                            metricKeys=["economic_environment_r04"],
                        )
                    ),
                ),
                Block(
                    id="metric_probe.evidence",
                    type="image",
                    blockType="slot",
                    source="user_input",
                    image=ImageModel(layoutAssetSlot=True, layoutAssetIds=[uuid4()]),
                ),
            ],
        )
    )
    buffer = BytesIO()
    Image.new("RGB", (60, 40), (10, 120, 10)).save(buffer, format="PNG")
    resolved = {
        "metric_probe.evidence": (
            ResolvedEvidenceImage(width_cm=16.0,
                data=buffer.getvalue(), caption="治理架构图", alt_text="治理架构图"
            ),
        )
    }

    figure_specs: list[FigureSpec] = []
    out = render_docx(
        report,
        base_template,
        out_dir / "metric_summary_numbering.docx",
        resolved_images=resolved,
        figure_specs_out=figure_specs,
    )

    # 摘要图不进 figure_specs；素材图仍从 图1 起编，编号不因摘要图顺延。
    figure_only = [spec for spec in figure_specs if spec.numbered_kind == "figure"]
    assert [(spec.number_label, spec.block_id, spec.kind) for spec in figure_only] == [
        ("图1", "metric_probe.evidence", "evidence_image")
    ]
    figure_captions = [
        paragraph.text
        for paragraph in Document(str(out)).paragraphs
        if paragraph.style.name == "Caption Figure"
    ]
    assert figure_captions == ["图1　治理架构图"]


def test_metric_summary_render_failure_falls_back_to_placeholder(monkeypatch, base_template, out_dir):
    def fail_render(*_args, **_kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(
        "sustainability_desk.export.metric_summary_chart.render_quantitative_metric_summary",
        fail_render,
    )
    report = load_contract(_CONTRACT)
    report.meta = ReportMeta(
        quantitativeMetrics=QuantitativeMetricsMeta(
            metrics={"economic_environment_r04": QuantitativeMetricDraft(value="120.5")}
        )
    )
    report.sections.append(
        Section(
            key="metric_probe",
            title="指标探针",
            headingLevel=1,
            blocks=[
                Block(
                    id="metric_probe.image",
                    type="image",
                    blockType="fixed",
                    source="derived",
                    image=ImageModel(
                        derivedVisualization=DerivedVisualizationSpec(
                            kind="quantitative_metric_summary",
                            metricKeys=["economic_environment_r04"],
                            emptyBehavior="hide",
                        )
                    ),
                )
            ],
        )
    )

    out = render_docx(report, base_template, out_dir / "metric_summary_failed.docx")
    full = "\n".join(paragraph.text for paragraph in Document(str(out)).paragraphs)

    assert "指标摘要图生成失败" in full
