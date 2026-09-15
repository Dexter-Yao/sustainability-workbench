# ABOUTME: 导出基底测试——自有基底必须提供渲染器依赖的样式、目录控件与脚注部件，且样式归正合乎格式标准。
# ABOUTME: 基底由 export/base_template.py 生成；本文件同时守护「重生成结果与入库文件一致」。
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from sustainability_desk.export.base_template import build_base_template
from sustainability_desk.export.format_profile import load_format_profile
from sustainability_desk.export.normalize_template import normalize_template
from knowledge_package_fixtures import SSE_PACKAGE


def test_heading_styles_no_autonumber_left_aligned(template_docx, out_dir):
    base = normalize_template(template_docx, out_dir / "base.docx", profile=load_format_profile(SSE_PACKAGE))
    doc = Document(str(base))
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        st = doc.styles[name]
        ppr = st._element.find(qn("w:pPr"))
        assert ppr is None or ppr.find(qn("w:numPr")) is None, f"{name} 仍有自动编号（双重编号）"
        assert st.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT, f"{name} 未左对齐"


def test_heading_sizes_and_theme_color_normalized(template_docx, out_dir):
    """标题字号归正 16/14/12/12pt，主题色统一到格式 profile 主色（格式标准 §1、§2）。"""
    from sustainability_desk.export.format_profile import load_format_profile, rgb

    base = normalize_template(template_docx, out_dir / "heading-theme.docx", profile=load_format_profile(SSE_PACKAGE))
    doc = Document(str(base))
    primary = rgb(load_format_profile(SSE_PACKAGE).theme.primary)
    expected_sizes = {"Heading 1": 16.0, "Heading 2": 14.0, "Heading 3": 12.0, "Heading 4": 12.0}
    for name, size_pt in expected_sizes.items():
        style = doc.styles[name]
        assert style.font.size.pt == size_pt, f"{name} 字号未归正"
        assert style.font.color.rgb == primary, f"{name} 颜色未归主题主色"
        color_elem = style.element.get_or_add_rPr().find(qn("w:color"))
        assert color_elem is None or color_elem.get(qn("w:themeColor")) is None, (
            f"{name} 仍残留主题色引用"
        )


def test_base_template_provides_renderer_dependencies(template_docx):
    """渲染器按 styleId TOC{n} 写目录条目、按名解析脚注样式、依赖目录控件与 footnotes 部件。"""
    doc = Document(str(template_docx))
    style_ids = {style.style_id for style in doc.styles}
    assert {"TOC1", "TOC2", "Normal", "Heading1", "Heading2", "Heading3", "Caption", "TableGrid"} <= style_ids
    names = {style.name.lower() for style in doc.styles}
    assert {"footnote text", "footnote reference", "annotation text", "annotation reference"} <= names
    toc_controls = [
        sdt
        for sdt in doc.element.body.findall(qn("w:sdt"))
        if any("TOC" in (node.text or "") for node in sdt.findall(".//" + qn("w:instrText")))
    ]
    assert len(toc_controls) == 1
    with zipfile.ZipFile(template_docx) as z:
        names_in_package = set(z.namelist())
        assert "word/footnotes.xml" in names_in_package
        assert not any(name.startswith("word/media/") for name in names_in_package), "基底不得携带任何图片"
    section = doc.sections[0]
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)


def test_body_font_size_follows_profile(template_docx):
    from sustainability_desk.export.format_profile import load_format_profile

    doc = Document(str(template_docx))
    assert doc.styles["Normal"].font.size.pt == load_format_profile(SSE_PACKAGE).body.font_size_pt


def test_toc_scope_stops_at_topic_without_open_field_updates(template_docx, out_dir):
    base = normalize_template(template_docx, out_dir / "base.docx", profile=load_format_profile(SSE_PACKAGE))
    with zipfile.ZipFile(base) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
        settings = z.read("word/settings.xml").decode("utf-8")
    assert 'TOC \\o "1-2"' in doc_xml
    assert 'TOC \\o "1-4"' not in doc_xml
    assert "updateFields" not in settings


def test_regenerated_base_matches_committed_styles(template_docx, out_dir):
    """重生成的基底与入库文件在样式表与正文结构上一致，防止手改入库文件后与生成器漂移。"""
    regenerated = build_base_template(SSE_PACKAGE, out_dir / "regenerated-base.docx")
    with zipfile.ZipFile(template_docx) as committed, zipfile.ZipFile(regenerated) as fresh:
        for part in ("word/styles.xml", "word/document.xml", "word/footnotes.xml"):
            assert committed.read(part) == fresh.read(part), f"{part} 与生成器输出不一致，请重跑 base_template"
