# ABOUTME: 空白统一填报工作簿（账户级母版）的血统合同：全零 report_id 哨兵身份、可导入任意报告、
# ABOUTME: 未声明空白血统或未提供空白期望时维持按报告全等校验，不放宽任何一侧。
from __future__ import annotations

from io import BytesIO
from uuid import UUID

import openpyxl
import pytest

from sustainability_desk.assets.report_basics_parser import parse_report_basics_workbook
from sustainability_desk.assets.report_basics_template import REPORT_BASICS_SHEET
from sustainability_desk.assets.scoring_template import SCORING_SHEET_TITLE
from sustainability_desk.assets.topic_questions_parser import parse_topic_questions_workbook
from sustainability_desk.assets.unified_workbook import (
    BLANK_UNIFIED_WORKBOOK_REPORT_ID,
    UNIFIED_QUANTITATIVE_CONFIG_SHEET,
    split_unified_workbook,
)
from sustainability_desk.contract.contract_version import contract_version
from sustainability_desk.contract.fill import fill_report
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.report_revision import build_report_revision
from sustainability_desk.contract.stored_report_state import StoredReportStateV4
from sustainability_desk.contract.structured_inputs import (
    STRUCTURED_INPUT_METADATA_SHEET,
    StructuredInputContext,
    StructuredInputWorkbookError,
)
from sustainability_desk.structured_input_service import (
    blank_unified_workbook_expected_metadata,
    blank_unified_workbook_template,
)
from knowledge_package_fixtures import SSE_PACKAGE

_CONTRACT = SSE_PACKAGE.report_contract_path

# 目标报告上下文必须与空白母版身份不同才能证明血统分支；合同版本按真实版本
# （空白期望由真实合同派生，目标侧解析也要求版本一致）。
TARGET_CONTEXT = StructuredInputContext(
    reportId=UUID("00000000-0000-0000-0000-000000000009"),
    contractVersion=contract_version(SSE_PACKAGE),
    compiledSemanticsVersion="semantics-test",
)


def _target_reports():
    state = StoredReportStateV4(
        version=4,
        fields={
            "company_registered_name": "目标测试公司",
            "industry_major_category": "制造业",
            "reporting_year": 2025,
        },
    )
    basics = fill_report(
        load_contract(_CONTRACT),
        {
            "fields": {k: v for k, v in state.fields.items() if v is not None},
            "intake": {},
        },
    )
    revision = build_report_revision(state, tolerate_incomplete_assessment=True, package=SSE_PACKAGE)
    return basics, revision


def test_blank_template_declares_nil_report_identity() -> None:
    content = blank_unified_workbook_template(SSE_PACKAGE)
    workbook = openpyxl.load_workbook(BytesIO(content))
    assert REPORT_BASICS_SHEET in workbook.sheetnames
    assert SCORING_SHEET_TITLE in workbook.sheetnames
    assert UNIFIED_QUANTITATIVE_CONFIG_SHEET in workbook.sheetnames
    metadata = {
        str(row[0]): "" if row[1] is None else str(row[1])
        for row in workbook[STRUCTURED_INPUT_METADATA_SHEET].iter_rows(
            min_row=2, values_only=True
        )
        if row and row[0]
    }
    assert metadata["report_id"] == str(BLANK_UNIFIED_WORKBOOK_REPORT_ID)
    assert metadata["contract_version"] == contract_version(SSE_PACKAGE)
    assert metadata == blank_unified_workbook_expected_metadata(SSE_PACKAGE)


def test_blank_template_imports_into_any_report_with_blank_expectation() -> None:
    content = blank_unified_workbook_template(SSE_PACKAGE)
    basics, revision = _target_reports()
    parts = split_unified_workbook(
        content,
        basics_report=basics,
        revision_report=revision,
        context=TARGET_CONTEXT,
        blank_expected_metadata=blank_unified_workbook_expected_metadata(SSE_PACKAGE),
    )
    # 空白母版评分/定量恒为空 → 跳过导入、保持目标现状。
    assert parts.scoring is None
    assert parts.quantitative is None
    # 子工作簿元数据按目标报告上下文注入：既有解析器零改动消费。
    parsed_basics = parse_report_basics_workbook(
        parts.report_basics,
        report=basics,
        context=TARGET_CONTEXT,
    )
    # 空白母版基础资料整批替换语义：未填写即清空为「无值」。
    assert parsed_basics.fields.get("company_registered_name") in (None, "")
    answers = parse_topic_questions_workbook(
        parts.topic_questions,
        report=revision,
        context=TARGET_CONTEXT,
    )
    assert answers == {}


def test_blank_lineage_requires_explicit_expectation() -> None:
    """未提供空白期望（如未来某调用点遗漏）时，空白母版不得静默通过按报告校验。"""

    content = blank_unified_workbook_template(SSE_PACKAGE)
    basics, revision = _target_reports()
    with pytest.raises(StructuredInputWorkbookError):
        split_unified_workbook(
            content,
            basics_report=basics,
            revision_report=revision,
            context=TARGET_CONTEXT,
        )
