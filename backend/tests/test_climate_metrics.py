# ABOUTME: 指标与目标支柱测试——climate.metrics 装配、轻量版移除目标段和自填指标表、定量数据进入上下文。
from pathlib import Path

from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import Report
from sustainability_desk.llm.prompts import build_model_context
from sustainability_desk.planner import load_topic_intake
from assessment_fixtures import complete_assessment
from knowledge_package_fixtures import SSE_PACKAGE
from sustainability_desk.llm.prompt_profiles import prompt_labels
from sustainability_desk.planner import load_topic_templates_from

BACKEND = Path(__file__).resolve().parents[1]
TOPIC_DIR = SSE_PACKAGE.topic_sections_dir
CONTRACT = SSE_PACKAGE.report_contract_path


def _metrics():
    children = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)["climate_change"].children
    return next(c for c in children if c.key == "climate_change.metrics")


def _block(bid):
    return next(
        block
        for section in [_metrics(), *(_metrics().children or [])]
        for block in section.blocks
        if block.id == bid
    )


def test_metrics_after_iro():
    children = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)["climate_change"].children
    keys = [c.key for c in children]
    assert keys.index("climate.iro") < keys.index("climate_change.metrics")


def test_metrics_keeps_only_emissions_methodology_block():
    """轻量版指标与目标支柱只保留派生指标图和核算口径段。"""
    ids = {
        block.id
        for section in [_metrics(), *(_metrics().children or [])]
        for block in section.blocks
    }
    assert ids == {
        "climate_change.metrics_summary",
        "climate_change.metrics_narrative_body",
        "climate_change.layout_assets",
    }
    visual = _block("climate_change.metrics_summary")
    assert visual.type == "image"
    assert visual.image.derivedVisualization.kind == "quantitative_metric_summary"
    assert set(visual.image.derivedVisualization.metricKeys) == {
        "economic_environment_r04",
        "economic_environment_r05",
        "economic_environment_r07",
    }
    blk = _block("climate_change.metrics_narrative_body")
    assert blk.blockType == "constrained"
    assert blk.generation.inputs.evidence.intakeItems == []
    assert set(blk.generation.inputs.evidence.quantitativeMetrics or []) == {
        "economic_environment_r04",
        "economic_environment_r05",
        "economic_environment_r07",
    }


def test_no_metrics_no_data_placeholder_branch():
    """指标与目标支柱不再用 no-data 占位段替代 ESG 定量指标 prompt 投影。"""
    assert all(
        block.id != "climate.metrics_indicators_placeholder"
        for section in [_metrics(), *(_metrics().children or [])]
        for block in section.blocks
    )


def test_declared_quantitative_metrics_enter_prompt_without_metrics_supplement():
    """已声明 ESG 定量指标进入 ModelContext；空值仅作为指标维度参考。"""
    section = load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)["climate_change"]
    block = _block("climate_change.metrics_narrative_body")
    report = Report(knowledgePackageId=SSE_PACKAGE.id, title="t", sections=[section])

    ctx = build_model_context(block, report, report)
    from sustainability_desk.llm.prompts import evidence_materials

    text = "\n".join(material.text for material in evidence_materials(ctx.evidence, prompt_labels(SSE_PACKAGE)))

    assert len(ctx.evidence.metric_evidence) == 3
    assert "来源：ESG 定量数据表" in text
    assert "本指标仅作为披露维度参考" in text
    assert "填写值" not in text
    assert "未填写" not in text
    assert "economic_environment_r04" not in text


def test_emissions_methodology_unfilled_metrics_do_not_expand_into_accounting_system():
    block = _block("climate_change.metrics_narrative_body")
    posture = "\n".join(block.generation.simplifiedWritingGuidance or [])

    assert posture == ""
    assert block.generation.inputs.evidence.quantitativeMetrics


def _assess():
    return complete_assessment(
        load_contract(CONTRACT),
        materialities={"climate_change": "dual"},
    )


def test_export_renders_metrics_and_hides_table(base_template, out_dir):
    """导出：（四）指标与目标标题保留；气候自填指标表不再由议题正文渲染。"""
    from docx import Document
    from sustainability_desk.contract.fill import fill_report
    from sustainability_desk.export.docx_renderer import load_values_flat, render_docx
    from sustainability_desk.planner import assemble_report

    intake = load_topic_intake(SSE_PACKAGE)
    assembled = assemble_report(
        load_contract(CONTRACT),
        load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE),
        assessment=_assess(),
        intake_items=intake,
    ).report
    report = fill_report(assembled, load_values_flat(SSE_PACKAGE.sample_values_path))
    out = render_docx(report, base_template, out_dir / "metrics_check.docx")
    doc = Document(str(out))
    heads = {p.text: p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")}
    assert heads.get("（四）指标与目标") == "Heading 3"
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "应对气候变化关键绩效指标" not in full
    assert "应对气候变化指标体系建设" not in full
