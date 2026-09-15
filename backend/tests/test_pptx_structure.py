# ABOUTME: 验证 PPTX 结构解析的阅读顺序、表格、备注、图片占位与显式缺口。
# ABOUTME: 用 python-pptx 程序化构造最小 fixture，覆盖文本框、表格、图片与纯图片页。
from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from pptx import Presentation
from pptx.util import Emu

from sustainability_desk.material.intake.parsed_material import (
    ParsedImageNode,
    ParsedSlideNode,
    ParsedTableNode,
    ParsedTableRowNode,
    ParsedTextNode,
)
from sustainability_desk.material.intake.pptx_structure import parse_pptx_structure


def _pptx_with_text_table_notes_and_image() -> bytes:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank_layout)
    # 故意先添加位置靠下的表格、再添加位置靠上的文本框，
    # 验证解析器按 (top, left) 阅读顺序重排，而非按插入顺序。
    table_shape = slide.shapes.add_table(
        2, 2, Emu(0), Emu(2_000_000), Emu(3_000_000), Emu(900_000)
    )
    table = table_shape.table
    table.cell(0, 0).text = "年份"
    table.cell(0, 1).text = "事件"
    table.cell(1, 0).text = "2020"
    table.cell(1, 1).text = "公司成立"

    textbox = slide.shapes.add_textbox(
        Emu(0), Emu(0), Emu(4_000_000), Emu(900_000)
    )
    textbox.text_frame.text = "发展历程"

    slide.notes_slide.notes_text_frame.text = "内部评审通过"

    from PIL import Image

    image_bytes = BytesIO()
    Image.new("RGB", (4, 4), color="red").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    image_only_slide = presentation.slides.add_slide(blank_layout)
    image_only_slide.shapes.add_picture(image_bytes, Emu(0), Emu(0))

    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_pptx_orders_shapes_by_top_then_left_and_keeps_table_and_notes() -> None:
    parsed = parse_pptx_structure(
        _pptx_with_text_table_notes_and_image(),
        source_id=uuid4(),
        source_label="发展历程.pptx",
    )

    slide_nodes = [node for node in parsed.nodes if isinstance(node, ParsedSlideNode)]
    assert len(slide_nodes) == 2

    slide_one_children = sorted(
        (
            node
            for node in parsed.nodes
            if node.parent_node_id == slide_nodes[0].node_id
        ),
        key=lambda node: node.ordinal_path,
    )
    # 阅读顺序：先靠上的文本框，后靠下的表格——与插入顺序相反。
    assert isinstance(slide_one_children[0], ParsedTextNode)
    assert slide_one_children[0].text == "发展历程"
    assert isinstance(slide_one_children[1], ParsedTableNode)
    assert slide_one_children[1].column_count == 2
    notes_node = slide_one_children[2]
    assert isinstance(notes_node, ParsedTextNode)
    assert notes_node.locator.kind == "pptx_notes"
    assert notes_node.text == "内部评审通过"

    table_rows = [
        node
        for node in parsed.nodes
        if isinstance(node, ParsedTableRowNode)
    ]
    assert len(table_rows) == 2
    assert [cell.text for cell in table_rows[1].cells] == ["2020", "公司成立"]

    image_slide_children = [
        node
        for node in parsed.nodes
        if node.parent_node_id == slide_nodes[1].node_id
    ]
    assert len(image_slide_children) == 1
    assert isinstance(image_slide_children[0], ParsedImageNode)
    assert image_slide_children[0].media_extraction_status == "complete"
    assert image_slide_children[0].media_sha256 is not None

    gap_codes = {gap.code for gap in parsed.coverage.gaps}
    assert "pptx_image_semantics_unparsed" in gap_codes
    assert "pptx_slide_without_extractable_text" in gap_codes
    assert parsed.status == "complete_with_gaps"
    assert parsed.coverage.disposition == "incomplete_usable"


def test_pptx_registers_gap_for_unmodeled_shape_kinds() -> None:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(blank_layout)
    # 连接线没有文本框也没有表格，代表动画/连接线/OLE 等尚未建模的形状类别。
    slide.shapes.add_connector(1, Emu(0), Emu(0), Emu(1_000_000), Emu(1_000_000))

    stream = BytesIO()
    presentation.save(stream)

    parsed = parse_pptx_structure(stream.getvalue(), source_id=uuid4())

    gap_codes = {gap.code for gap in parsed.coverage.gaps}
    assert "pptx_shape_semantics_not_modeled" in gap_codes
