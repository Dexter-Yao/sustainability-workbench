# ABOUTME: 统一填报工作簿真实 xlsx 测试：四类 sheet 聚合、统一指纹元数据、拆分重建与空面跳过语义。
# ABOUTME: 拆出的子工作簿必须能被四个既有解析器零改动消费；评分/定量留空返回 None。
from __future__ import annotations

from io import BytesIO
from uuid import UUID

import openpyxl

from sustainability_desk.assets.quantitative_template import QUANTITATIVE_HEADER_ROW
from sustainability_desk.assets.quantitative_parser import parse_quantitative_workbook
from sustainability_desk.assets.report_basics_parser import parse_report_basics_workbook
from sustainability_desk.assets.report_basics_template import REPORT_BASICS_SHEET
from sustainability_desk.assets.scoring import parse_scoring
from sustainability_desk.assets.scoring_template import SCORING_SHEET_TITLE
from sustainability_desk.assets.topic_questions_parser import parse_topic_questions_workbook
from sustainability_desk.assets.unified_workbook import (
    UNIFIED_OVERVIEW_SHEET,
    UNIFIED_QUANTITATIVE_CONFIG_SHEET,
    create_unified_workbook,
    split_unified_workbook,
)
from sustainability_desk.contract.fill import fill_report
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.report_revision import build_report_revision
from sustainability_desk.contract.stored_report_state import StoredReportStateV4
from sustainability_desk.contract.structured_inputs import (
    STRUCTURED_INPUT_METADATA_SHEET,
    StructuredInputContext,
)
from knowledge_package_fixtures import SSE_PACKAGE

CONTEXT = StructuredInputContext(
    reportId=UUID("00000000-0000-0000-0000-000000000003"),
    contractVersion="cv-test",
    compiledSemanticsVersion="semantics-test",
)
_CONTRACT = SSE_PACKAGE.report_contract_path


def _state() -> StoredReportStateV4:
    return StoredReportStateV4(
        version=4,
        fields={
            "company_registered_name": "测试股份有限公司",
            "industry_major_category": "制造业",
            "reporting_year": 2025,
        },
    )


def _reports():
    state = _state()
    basics = fill_report(
        load_contract(_CONTRACT),
        {"fields": {k: v for k, v in state.fields.items() if v is not None}, "intake": {}},
    )
    revision = build_report_revision(state, tolerate_incomplete_assessment=True, package=SSE_PACKAGE)
    return basics, revision


def _unified_bytes() -> tuple[bytes, object, object]:
    basics, revision = _reports()
    content = create_unified_workbook(basics, revision, context=CONTEXT)
    return content, basics, revision


def test_unified_workbook_aggregates_all_input_sheets() -> None:
    content, _, _ = _unified_bytes()
    workbook = openpyxl.load_workbook(BytesIO(content))
    names = workbook.sheetnames
    assert names[0] == UNIFIED_OVERVIEW_SHEET
    assert REPORT_BASICS_SHEET in names
    assert SCORING_SHEET_TITLE in names
    assert UNIFIED_QUANTITATIVE_CONFIG_SHEET in names
    assert "经济+环境" in names and "社会" in names and "治理" in names
    assert "应对气候变化" in names
    assert STRUCTURED_INPUT_METADATA_SHEET in names
    # 定量说明页在统一册中改名，不与总览/议题说明冲突。
    assert "填写说明" not in names


def test_split_skips_empty_scoring_and_quantitative_and_feeds_parsers() -> None:
    content, basics, revision = _unified_bytes()
    parts = split_unified_workbook(
        content,
        basics_report=basics,
        revision_report=revision,
        context=CONTEXT,
    )
    # 评分与定量整表留空 → 跳过导入、保持系统内现状。
    assert parts.scoring is None
    assert parts.quantitative is None

    parsed_basics = parse_report_basics_workbook(
        parts.report_basics,
        report=basics,
        context=CONTEXT,
    )
    assert parsed_basics.fields["company_registered_name"] == "测试股份有限公司"

    answers = parse_topic_questions_workbook(
        parts.topic_questions,
        report=revision,
        context=CONTEXT,
    )
    assert answers == {}


def test_split_rebuilds_filled_scoring_and_quantitative_for_parsers() -> None:
    content, basics, revision = _unified_bytes()
    workbook = openpyxl.load_workbook(BytesIO(content))

    scoring_sheet = workbook[SCORING_SHEET_TITLE]
    for row in scoring_sheet.iter_rows(min_row=9):
        if row[0].value and not scoring_sheet.merged_cells.__contains__(row[0].coordinate):
            row[1].value = 4.5
            row[2].value = 4.0
    quant_sheet = workbook["经济+环境"]
    for row in quant_sheet.iter_rows(min_row=QUANTITATIVE_HEADER_ROW + 1):
        if row[0].value:
            row[5].value = "not_collected"
    for sheet_name in ("社会", "治理"):
        for row in workbook[sheet_name].iter_rows(min_row=QUANTITATIVE_HEADER_ROW + 1):
            if row[0].value:
                row[5].value = "not_collected"
    for row in workbook["应对气候变化"].iter_rows():
        if row[0].value == "climate.q_governance_policies":
            row[4].value = "不确定"
    output = BytesIO()
    workbook.save(output)

    parts = split_unified_workbook(
        output.getvalue(),
        basics_report=basics,
        revision_report=revision,
        context=CONTEXT,
    )
    assert parts.scoring is not None
    assert parts.quantitative is not None

    parsed_scoring = parse_scoring(parts.scoring, report=basics, context=CONTEXT)
    assert parsed_scoring.scored
    assert all(score.financialScore == 4.5 for score in parsed_scoring.scored)

    metrics = parse_quantitative_workbook(
        parts.quantitative,
        report=basics,
        context=CONTEXT,
    )
    assert len(metrics.metrics) == 149

    answers = parse_topic_questions_workbook(
        parts.topic_questions,
        report=revision,
        context=CONTEXT,
    )
    assert answers["climate.q_governance_policies"].answer == "不确定"
