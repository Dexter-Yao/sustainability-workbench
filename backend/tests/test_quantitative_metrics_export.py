# ABOUTME: ESG KPI 附录导出回归，验证 Word 始终从定量元数据重建而非使用浏览器表格缓存。
# ABOUTME: 覆盖历史快照含残留 table.children 时的 SSOT 边界，防止缓存行进入正式交付物。
from __future__ import annotations


from docx import Document

from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import GsTableCell, GsTableRow, QuantitativeMetricDraft, QuantitativeMetricsMeta, ReportMeta
from sustainability_desk.diagnostics import diagnose
from sustainability_desk.export.docx_renderer import render_docx
from knowledge_package_fixtures import SSE_PACKAGE


CONTRACT = SSE_PACKAGE.report_contract_path


def test_word_export_rebuilds_appendix_metrics_from_meta(base_template, out_dir) -> None:
    report = load_contract(CONTRACT)
    report.meta = ReportMeta(
        quantitativeMetrics=QuantitativeMetricsMeta(
            metrics={"economic_environment_r04": QuantitativeMetricDraft(value="12.5")}
        )
    )
    appendix = next(section for section in report.sections if section.key == "report_appendix")
    table = appendix.children[0].blocks[0].table
    table.children.append(
        GsTableRow(children=[GsTableCell(colKey="metric", value="历史缓存指标")])
    )

    out = render_docx(report, base_template, out_dir / "metrics_ssot.docx")
    cells = [cell.text for table in Document(str(out)).tables for row in table.rows for cell in row.cells]

    assert "历史缓存指标" not in cells
    assert "范围一：温室气体排放总量" in cells
    assert "12.5" in cells


def test_diagnosis_ignores_cached_appendix_rows() -> None:
    report = load_contract(CONTRACT)
    appendix = next(section for section in report.sections if section.key == "report_appendix")
    table = appendix.children[0].blocks[0].table
    table.children.append(
        GsTableRow(state="failed", children=[GsTableCell(colKey="metric", value="历史缓存指标")])
    )

    issues = diagnose(report).issues

    assert not any(issue.blockId == "appendix.esg_key_performance_metrics" for issue in issues)


def _report_with_note_metric():
    report = load_contract(CONTRACT)
    report.meta = ReportMeta(
        quantitativeMetrics=QuantitativeMetricsMeta(
            metrics={
                "social_r02": QuantitativeMetricDraft(
                    value="331", note="上海公司编制口径，正式交付前请客户确认更新。"
                )
            }
        )
    )
    return report


def test_standard_word_omits_remark_column_and_shows_topic_name_only(
    base_template, out_dir
) -> None:
    """口径备注是复核性说明只进审阅稿；正式稿 KPI 表 4 列，第一列仅议题名称。"""
    report = _report_with_note_metric()

    out = render_docx(report, base_template, out_dir / "metrics_standard.docx")
    doc = Document(str(out))
    kpi_table = next(
        table
        for table in doc.tables
        if any("议题类别" in cell.text for row in table.rows for cell in row.cells)
    )
    header_texts = [cell.text for cell in kpi_table.rows[0].cells]
    all_texts = [cell.text for row in kpi_table.rows for cell in row.cells]

    assert "口径备注" not in header_texts
    assert len(header_texts) == 4
    assert not any("请客户确认" in text for text in all_texts)
    assert "人力资本发展" in all_texts
    assert not any("人力资本发展-" in text for text in all_texts)


def test_review_word_keeps_remark_column(base_template, out_dir) -> None:
    from datetime import datetime, timezone

    report = _report_with_note_metric()

    out = render_docx(
        report,
        base_template,
        out_dir / "metrics_review.docx",
        delivery_variant="review",
        review_generated_at=datetime(2026, 8, 4, 12, 0, tzinfo=timezone.utc),
    )
    doc = Document(str(out))
    kpi_table = next(
        table
        for table in doc.tables
        if any("议题类别" in cell.text for row in table.rows for cell in row.cells)
    )
    header_texts = [cell.text for cell in kpi_table.rows[0].cells]
    all_texts = [cell.text for row in kpi_table.rows for cell in row.cells]

    assert "口径备注" in header_texts
    assert any("请客户确认" in text for text in all_texts)
