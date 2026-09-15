# ABOUTME: 基础资料工作簿真实 xlsx 往返测试：行计划与合同互核、类型 parse-first、准则与附录映射。
# ABOUTME: 解析失败必须返回全部 typed 单元格错误；成功导入为模板范围内整批替换。
from __future__ import annotations

from io import BytesIO
from uuid import UUID

import openpyxl
import pytest

from sustainability_desk.assets.report_basics_parser import parse_report_basics_workbook
from sustainability_desk.assets.report_basics_template import (
    REPORT_BASICS_SHEET,
    create_report_basics_template,
    validate_report_basics_row_plan,
)
from sustainability_desk.contract.fill import fill_report
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import Report
from sustainability_desk.contract.structured_inputs import (
    STALE_WORKBOOK_MESSAGE,
    StructuredInputContext,
    StructuredInputWorkbookError,
)
from knowledge_package_fixtures import SSE_PACKAGE

CONTEXT = StructuredInputContext(
    reportId=UUID("00000000-0000-0000-0000-000000000002"),
    contractVersion="cv-test",
    compiledSemanticsVersion="semantics-test",
)
_CONTRACT = SSE_PACKAGE.report_contract_path


def _report() -> Report:
    return fill_report(
        load_contract(_CONTRACT),
        {
            "fields": {
                "company_registered_name": "测试股份有限公司",
                "reporting_year": 2025,
            },
            "intake": {
                "company_profile": {
                    "answer": "公司专注节能设备研发制造，产品覆盖多个行业客户，" * 6,
                }
            },
        },
    )


def _set_value(workbook, key: str, value: object, column: int = 6) -> None:
    worksheet = workbook[REPORT_BASICS_SHEET]
    for row in worksheet.iter_rows():
        if row[0].value == key:
            row[column - 1].value = value
            return
    raise AssertionError(f"模板中找不到行：{key}")


def test_row_plan_matches_contract_user_input_fields() -> None:
    # 合同新增 user_input 字段而行计划未更新时必须 fail-loud，不允许模板静默漏字段。
    validate_report_basics_row_plan(_report())


def test_template_round_trip_parses_types_and_mappings() -> None:
    report = _report()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))

    assert workbook[REPORT_BASICS_SHEET].column_dimensions["A"].hidden is True
    # 模板预填当前值。
    rows = {
        row[0]: row
        for row in workbook[REPORT_BASICS_SHEET].iter_rows(values_only=True)
        if row[0]
    }
    assert rows["field:company_registered_name"][5] == "测试股份有限公司"

    _set_value(workbook, "field:company_short_name", "测试科技")
    _set_value(workbook, "field:industry_major_category", "制造业")
    _set_value(workbook, "field:report_period_start", "2025-01-01")
    _set_value(workbook, "field:report_approval_year", 2026)
    _set_value(workbook, "field:report_approval_month", "2026-06")
    _set_value(workbook, "field:board_attendance_rate", 98.5)
    _set_value(
        workbook,
        "disclosure:mainlandStandard",
        "《深圳证券交易所上市公司自律监管指引第17号——可持续发展报告（试行）》",
    )
    _set_value(workbook, "appendix:externalAssuranceReport.isIncluded", "是")
    _set_value(workbook, "appendix:externalAssuranceReport.fileLabel", "外部鉴证声明.pdf")
    _set_value(
        workbook,
        "appendix:readerFeedbackContactInformation.email",
        "esg@example.com",
    )
    output = BytesIO()
    workbook.save(output)

    parsed = parse_report_basics_workbook(
        output.getvalue(),
        report=report,
        context=CONTEXT,
    )
    assert parsed.fields["company_short_name"] == "测试科技"
    assert parsed.fields["industry_major_category"] == "制造业"
    assert parsed.fields["reporting_year"] == 2025
    assert parsed.fields["report_period_start"] == "2025-01-01"
    # month 合同口径是 YYYY-MM（正文引用与前端 <input type="month"> 同源）。
    assert parsed.fields["report_approval_month"] == "2026-06"
    assert parsed.fields["board_attendance_rate"] == 98.5
    assert parsed.disclosure_profile is not None
    assert parsed.disclosure_profile.mainlandStandard == "szse"
    assert parsed.appendix_package is not None
    assert parsed.appendix_package.externalAssuranceReport.isIncluded is True
    assert (
        parsed.appendix_package.externalAssuranceReport.fileLabel
        == "外部鉴证声明.pdf"
    )
    assert (
        parsed.appendix_package.readerFeedbackContactInformation.email
        == "esg@example.com"
    )
    # 公司简介预填值原样往返。
    assert "company_profile" in parsed.intake_answers


def test_invalid_values_produce_cell_errors() -> None:
    report = _report()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    _set_value(workbook, "field:reporting_year", "明年")
    _set_value(workbook, "field:report_period_end", "2025/12/31")
    _set_value(workbook, "field:industry_major_category", "自造行业")
    # month 填纯数字月份是历史高危形态：必须在导入边界拒绝，不允许流到装配/导出才爆。
    _set_value(workbook, "field:report_approval_month", 3)
    _set_value(workbook, "appendix:readerFeedbackContactInformation.email", "不是邮箱")
    _set_value(workbook, "disclosure:mainlandStandard", "自造准则")
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(StructuredInputWorkbookError) as excinfo:
        parse_report_basics_workbook(
            output.getvalue(),
            report=report,
            context=CONTEXT,
        )
    codes = {error.code for error in excinfo.value.errors}
    assert codes >= {
        "report_basics_number_invalid",
        "report_basics_date_invalid",
        "report_basics_enum_invalid",
        "report_basics_email_invalid",
        "report_basics_standard_invalid",
        "report_basics_month_invalid",
    }


def test_missing_row_is_rejected() -> None:
    report = _report()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    _set_value(workbook, "field:company_short_name", "unknown:row", column=1)
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(StructuredInputWorkbookError) as excinfo:
        parse_report_basics_workbook(
            output.getvalue(),
            report=report,
            context=CONTEXT,
        )
    codes = {error.code for error in excinfo.value.errors}
    assert "report_basics_key_unknown" in codes
    assert "report_basics_key_missing" in codes


def _report_with_legacy_disclosure_references() -> Report:
    """历史报告：收窄前已保存港交所守则与其他参考文件。"""
    return fill_report(
        load_contract(_CONTRACT),
        {
            "fields": {
                "company_registered_name": "测试股份有限公司",
                "reporting_year": 2025,
            },
            "intake": {
                "company_profile": {
                    "answer": "公司专注节能设备研发制造，产品覆盖多个行业客户，" * 6,
                }
            },
            "disclosureProfile": {
                "mainlandStandard": "sse",
                "includesHongKongExchangeGuide": True,
                "additionalDisclosureReferences": ["《企业可持续披露准则——基本准则（试行）》"],
            },
        },
    )


def test_disabled_disclosure_rows_absent_from_template() -> None:
    """收窄期：港交所守则与其他参考文件不再出现在收资表，Excel 侧无从表达它们。"""
    report = _report_with_legacy_disclosure_references()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    keys = {
        row[0]
        for row in workbook[REPORT_BASICS_SHEET].iter_rows(values_only=True)
        if row[0]
    }
    assert "disclosure:includesHongKongExchangeGuide" not in keys
    assert "disclosure:additionalDisclosureReferences" not in keys
    assert "disclosure:mainlandStandard" in keys


def test_workbook_import_preserves_disabled_disclosure_references() -> None:
    """禁用项不在模板范围内：整批替换必须原样保留既有值，不得静默清空编制依据。"""
    report = _report_with_legacy_disclosure_references()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    _set_value(
        workbook,
        "disclosure:mainlandStandard",
        "《深圳证券交易所上市公司自律监管指引第17号——可持续发展报告（试行）》",
    )
    output = BytesIO()
    workbook.save(output)

    parsed = parse_report_basics_workbook(output.getvalue(), report=report, context=CONTEXT)

    assert parsed.disclosure_profile is not None
    # 模板范围内的准则照常整批替换。
    assert parsed.disclosure_profile.mainlandStandard == "szse"
    # 禁用项不被整批替换清空。
    assert parsed.disclosure_profile.includesHongKongExchangeGuide is True
    assert parsed.disclosure_profile.additionalDisclosureReferences == [
        "《企业可持续披露准则——基本准则（试行）》"
    ]


def test_workbook_import_rejects_forged_disabled_disclosure_rows() -> None:
    """Excel 通道不得绕过界面禁用：手工补回隐藏键的伪造行按未知键整簿拒绝。"""
    report = _report_with_legacy_disclosure_references()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    worksheet = workbook[REPORT_BASICS_SHEET]
    worksheet.append(["disclosure:includesHongKongExchangeGuide", None, None, None, None, "否"])
    worksheet.append(["disclosure:additionalDisclosureReferences", None, None, None, None, "伪造准则"])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(StructuredInputWorkbookError) as excinfo:
        parse_report_basics_workbook(output.getvalue(), report=report, context=CONTEXT)

    codes = {error.code for error in excinfo.value.errors}
    assert codes == {"report_basics_key_unknown"}


def test_stale_workbook_reports_actionable_message_without_internal_identifiers() -> None:
    """旧模板导入必须给出用户可执行的中文提示，不泄露元数据键名、指纹与 report_id。"""
    report = _report()
    content = create_report_basics_template(report, context=CONTEXT)
    workbook = openpyxl.load_workbook(BytesIO(content))
    stale_context = StructuredInputContext(
        reportId=CONTEXT.reportId,
        contractVersion="cv-previous",
        compiledSemanticsVersion="semantics-previous",
    )
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(StructuredInputWorkbookError) as excinfo:
        parse_report_basics_workbook(
            output.getvalue(), report=report, context=stale_context
        )

    errors = excinfo.value.errors
    assert any(error.code == "metadata_mismatch" for error in errors)
    for error in errors:
        assert error.message == STALE_WORKBOOK_MESSAGE
        assert "重新下载模板" in error.message
        # 内部标识符一律不进用户可见文本。
        assert "context_fingerprint" not in error.message
        assert "contract_version" not in error.message
        assert "report_id" not in error.message
        assert str(CONTEXT.reportId) not in error.message
