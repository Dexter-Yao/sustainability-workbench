# ABOUTME: 素材图片承载块的 Word 渲染测试:多图按序嵌入、宽度按解析器给定、题注与替代文本落 docx。
# ABOUTME: renderer 只消费 ResolvedEvidenceImage,尺寸判断不在渲染端。
from __future__ import annotations

from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from PIL import Image

from sustainability_desk.contract.models import Block, ImageModel, Report, Section
from sustainability_desk.export.docx_renderer import (
    ResolvedEvidenceImage,
    build_document_render_plan,
    render_docx,
)
from knowledge_package_fixtures import SSE_PACKAGE


def _png(color: tuple[int, int, int], size: tuple[int, int] = (60, 40)) -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def _slot_report() -> Report:
    return Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="素材承载位渲染测试",
        sections=[
            Section(
                key="probe",
                title="探针",
                headingLevel=1,
                blocks=[
                    Block(
                        id="probe.layout_assets",
                        type="image",
                        blockType="slot",
                        source="user_input",
                        image=ImageModel(
                            layoutAssetSlot=True,
                            layoutAssetIds=[uuid4(), uuid4()],
                        ),
                    )
                ],
            )
        ],
    )


def test_layout_asset_group_renders_in_order_with_captions(base_template, out_dir) -> None:
    report = _slot_report()
    resolved = {
        "probe.layout_assets": (
            ResolvedEvidenceImage(
                data=_png((200, 30, 30)),
                caption="ISO 9001 质量管理体系认证证书",
                alt_text="红色测试图一",
                width_cm=11.0,
            ),
            ResolvedEvidenceImage(
                data=_png((30, 30, 200)),
                caption="高新技术企业证书",
                alt_text="蓝色测试图二",
                width_cm=16.0,
            ),
        )
    }
    plan = build_document_render_plan(report)
    output = render_docx(
        report,
        base_template,
        out_dir / "layout-assets.docx",
        resolved_images=resolved,
        render_plan=plan,
    )

    document = Document(str(output))
    texts = [paragraph.text for paragraph in document.paragraphs]
    first = texts.index("图1　ISO 9001 质量管理体系认证证书")
    second = texts.index("图2　高新技术企业证书")
    assert first < second, "多图必须按解析顺序渲染"
    caption_styles = {
        paragraph.text: paragraph.style.name
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Caption")
    }
    assert caption_styles["图1　ISO 9001 质量管理体系认证证书"] == "Caption Figure"
    assert caption_styles["图2　高新技术企业证书"] == "Caption Figure"

    with ZipFile(output) as bundle:
        media = [name for name in bundle.namelist() if name.startswith("word/media/")]
    assert len([name for name in media if "image" in name]) >= 2

    xml = document.element.xml
    assert 'descr="红色测试图一"' in xml
    assert 'descr="蓝色测试图二"' in xml
    # EMU: 11cm=3960000、16cm=5760000,宽度由解析器给定
    assert 'cx="3960000"' in xml
    assert 'cx="5760000"' in xml


def test_empty_layout_slot_stays_out_of_word(base_template, out_dir) -> None:
    report = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="空承载位",
        sections=[
            Section(
                key="probe",
                title="探针",
                headingLevel=1,
                blocks=[
                    Block(
                        id="probe.layout_assets",
                        type="image",
                        blockType="slot",
                        source="user_input",
                        image=ImageModel(layoutAssetSlot=True),
                    )
                ],
            )
        ],
    )
    plan = build_document_render_plan(report)
    output = render_docx(
        report,
        base_template,
        out_dir / "empty-layout-assets.docx",
        render_plan=plan,
    )
    assert not any(unit.block_id == "probe.layout_assets" for unit in plan.units)
    full_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "图片占位" not in full_text
