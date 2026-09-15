# ABOUTME: 验证 File Agent 只通过受限只读工具理解文件，并冻结多范围 FileMaterial。
# ABOUTME: 文件内容的细节由 File Agent 保留；下游 Mapping 不重新提取或改写它。
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Emu
from pydantic import ValidationError

from sustainability_desk.material.intake.file_agent_ai import (
    FileAgentDeps,
    FileDossierProposal,
    build_report_need_catalog,
)
from sustainability_desk.material.intake.file_agent_contract import (
    FileAgentProductTask,
    FileAgentToolRequest,
    FileDossierDraft,
    FileMaterialDraft,
    FileMaterialScope,
)
from sustainability_desk.material.intake.file_agent_workspace import FileAgentWorkspace, FileAgentWorkspaceError
from sustainability_desk.material.intake.models import UserFileDeclaration
from sustainability_desk.llm.prompt_profiles import load_prompt_profile
from knowledge_package_fixtures import SSE_PACKAGE

FILE_AGENT_TEXTS = load_prompt_profile(SSE_PACKAGE).agent_instructions.file_agent
FILE_AGENT_INSTRUCTIONS = FILE_AGENT_TEXTS.instructions


def _task() -> FileAgentProductTask:
    return FileAgentProductTask(
        report_id=uuid4(), report_subject_name="测试企业", report_period_label="2025 年度",
        material_scopes=(
            FileMaterialScope(alias="scope_climate", scope_id="report-section:climate_change", title="应对气候变化", kind="esg_topic"),
            FileMaterialScope(alias="scope_company", scope_id="report-area:company_intro", title="公司治理", kind="report_area"),
        ),
    )


def _workspace(tmp_path: Path) -> FileAgentWorkspace:
    document = Document()
    document.add_heading("气候治理", level=1)
    document.add_paragraph("董事会每年审议气候相关风险。")
    path = tmp_path / "climate.docx"
    document.save(path)
    return FileAgentWorkspace.open(
        source_id=uuid4(), source_path=path,
        declaration=UserFileDeclaration(description="公司提交的气候治理制度资料。", role="semantic_material", topic_tags=["气候变化"]),
        product_task=_task(),
    )


def _pptx_workspace(tmp_path: Path) -> FileAgentWorkspace:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank_layout)
    textbox = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(4_000_000), Emu(900_000))
    textbox.text_frame.text = "发展历程"
    table_shape = slide.shapes.add_table(
        2, 2, Emu(0), Emu(1_000_000), Emu(3_000_000), Emu(900_000)
    )
    table = table_shape.table
    table.cell(0, 0).text = "年份"
    table.cell(0, 1).text = "事件"
    table.cell(1, 0).text = "2020"
    table.cell(1, 1).text = "公司成立"
    slide.notes_slide.notes_text_frame.text = "内部评审通过"

    from PIL import Image

    image_bytes = BytesIO()
    Image.new("RGB", (4, 4), color="red").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    image_only_slide = presentation.slides.add_slide(blank_layout)
    image_only_slide.shapes.add_picture(image_bytes, Emu(0), Emu(0))

    path = tmp_path / "发展历程.pptx"
    presentation.save(path)
    return FileAgentWorkspace.open(
        source_id=uuid4(), source_path=path,
        declaration=UserFileDeclaration(description="公司提交的发展历程演示文稿。", role="semantic_material", topic_tags=["公司治理"]),
        product_task=_task(),
    )


def test_workspace_converts_pptx_with_reading_order_table_notes_and_image_placeholder(
    tmp_path: Path,
) -> None:
    workspace = _pptx_workspace(tmp_path)
    manifest = workspace.execute(FileAgentToolRequest(tool_name="convert_file"))
    opened = workspace.execute(
        FileAgentToolRequest(tool_name="read_file", path=manifest.files[0].path)
    )
    content = opened.content
    assert "## 第 1 页" in content
    assert content.index("发展历程") < content.index("| 年份")
    assert "2020" in content and "公司成立" in content
    assert "备注：内部评审通过" in content
    assert "## 第 2 页" in content
    assert "（本页为图片内容，未提取文字）" in content


def test_workspace_only_parses_after_authorized_conversion(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    assert workspace.parsed_material is None
    manifest = workspace.execute(FileAgentToolRequest(tool_name="convert_file"))
    opened = workspace.execute(FileAgentToolRequest(tool_name="read_file", path=manifest.files[0].path))
    assert "董事会每年审议气候相关风险。" in opened.content
    assert set(workspace.context.available_tools) == {"convert_file", "read_file"}


def test_file_dossier_draft_requires_materials_only_when_relevant() -> None:
    material = FileMaterialDraft(
        applicable_scope_aliases=("scope_climate",),
        content_markdown="董事会每年审议气候相关风险。",
    )
    assert FileDossierDraft(
        relevance="relevant", relevance_reason="气候治理资料。", materials=(material,)
    ).contract == "sustainability_desk.file_dossier_draft.v7"
    with pytest.raises(ValidationError):
        FileDossierDraft(relevance="relevant", relevance_reason="相关。")
    with pytest.raises(ValidationError):
        FileDossierDraft(relevance="not_relevant", relevance_reason="无关。", materials=(material,))


def test_file_dossier_proposal_parses_provider_json_fields_before_domain_validation() -> None:
    proposal = FileDossierProposal.model_validate(
        {
            "relevance": "relevant",
            "relevance_reason": "气候治理资料。",
            "materials": '[{"applicable_scope_aliases":["scope_climate"],"content_markdown":"董事会每年审议气候相关风险。"}]',
            "attention_items": '[{"code":"data_freshness","message":"文件未说明最新更新日期。","next_action":"结合后续资料审慎使用。"}]',
        }
    )
    draft = proposal.to_draft()
    assert draft.materials[0].applicable_scope_aliases == ("scope_climate",)
    assert draft.attention_items[0].code == "data_freshness"


def test_harness_freezes_multiple_scope_materials_after_full_read(tmp_path: Path) -> None:
    workspace = _workspace(tmp_path)
    deps = FileAgentDeps(workspace=workspace)
    draft = FileDossierDraft(
        relevance="relevant", relevance_reason="文件包含治理职责。",
        materials=(FileMaterialDraft(
            applicable_scope_aliases=("scope_climate", "scope_company"),
            content_markdown="董事会每年审议气候相关风险，并负责相关治理监督。",
        ),),
    )
    with pytest.raises(FileAgentWorkspaceError, match="转换并完整读取"):
        deps.finalize(draft)

    manifest = workspace.execute(FileAgentToolRequest(tool_name="convert_file"))
    workspace.execute(FileAgentToolRequest(tool_name="read_file", path=manifest.files[0].path))
    dossier = deps.finalize(draft)

    assert dossier.contract == "sustainability_desk.file_dossier.v7"
    assert dossier.materials[0].applicable_scope_ids == (
        "report-section:climate_change", "report-area:company_intro",
    )
    assert dossier.materials[0].content_markdown.startswith("董事会")


def test_file_agent_prompt_requires_detail_without_raw_locator_contract() -> None:
    catalog = build_report_need_catalog(_task().material_scopes)
    assert catalog["allowedMaterialScopes"][0]["scopeAlias"] == "scope_climate"
    assert "不得为了简短而省略" in FILE_AGENT_INSTRUCTIONS
    assert "字符坐标" in FILE_AGENT_INSTRUCTIONS
    assert "直接提交严格的 FileDossierProposal" in FILE_AGENT_INSTRUCTIONS


def _task_with_uncovered_scope() -> FileAgentProductTask:
    """目录含本次报告未覆盖的范围（如受限范围下的水资源议题）。"""

    return FileAgentProductTask(
        report_id=uuid4(), report_subject_name="测试企业", report_period_label="2025 年度",
        scope_label="轻量版气候章节试用",
        material_scopes=(
            FileMaterialScope(alias="scope_climate", scope_id="report-section:climate_change", title="应对气候变化", kind="esg_topic"),
            FileMaterialScope(
                alias="scope_water", scope_id="report-section:water_resource_management",
                title="水资源管理", kind="esg_topic", within_report_scope=False,
            ),
        ),
    )


def test_file_agent_catalog_keeps_uncovered_scopes_and_states_report_coverage(tmp_path: Path) -> None:
    """范围外议题留在目录里并标明不在本次报告内，模型据此如实归类而不是判不相关。"""
    from sustainability_desk.material.intake.file_agent_ai import _initial_prompt

    task = _task_with_uncovered_scope()
    catalog = build_report_need_catalog(task.material_scopes)
    assert [item["withinReportScope"] for item in catalog["allowedMaterialScopes"]] == [True, False]
    assert task.covered_scopes() == (task.material_scopes[0],)

    document = Document()
    document.add_paragraph("公司建立了水资源管理制度。")
    path = tmp_path / "water.docx"
    document.save(path)
    workspace = FileAgentWorkspace.open(
        source_id=uuid4(), source_path=path,
        declaration=UserFileDeclaration(description="公司提交的水资源管理制度资料。", role="semantic_material", topic_tags=["uncertain"]),
        product_task=task,
    )
    prompt = _initial_prompt(workspace, FILE_AGENT_TEXTS)
    assert '"allowedScopeAliases":["scope_climate","scope_water"]' in prompt
    assert '"coveredScopeAliases":["scope_climate"]' in prompt
    assert "不进入本次报告" in prompt
    assert "不得因为该范围不在本次报告内就判为 not_relevant" in FILE_AGENT_INSTRUCTIONS

    # Harness 照常把范围外 alias 解析成 compiled scope id：适用范围如实冻结，不丢信息。
    deps = FileAgentDeps(workspace=workspace)
    manifest = workspace.execute(FileAgentToolRequest(tool_name="convert_file"))
    workspace.execute(FileAgentToolRequest(tool_name="read_file", path=manifest.files[0].path))
    dossier = deps.finalize(FileDossierDraft(
        relevance="relevant", relevance_reason="文件只涉及水资源管理。",
        materials=(FileMaterialDraft(applicable_scope_aliases=("scope_water",), content_markdown="公司建立了水资源管理制度。"),),
    ))
    assert dossier.materials[0].applicable_scope_ids == ("report-section:water_resource_management",)


def test_product_task_requires_at_least_one_covered_scope() -> None:
    with pytest.raises(ValidationError, match="至少要有一个属于本次报告范围"):
        FileAgentProductTask(
            report_id=uuid4(),
            material_scopes=(
                FileMaterialScope(
                    alias="scope_water", scope_id="report-section:water_resource_management",
                    title="水资源管理", kind="esg_topic", within_report_scope=False,
                ),
            ),
        )
