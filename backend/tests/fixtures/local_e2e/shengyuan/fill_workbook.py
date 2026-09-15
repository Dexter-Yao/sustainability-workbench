# ABOUTME: 用晟原精密的事实基准填写统一填报工作簿，产出可直接导入的 xlsx。
# ABOUTME: 只定点写入可写单元格，不重建工作簿、不触碰隐藏元数据，保住导入侧血统校验。
"""填写统一填报工作簿（晟原精密合成 corpus）。

数值与文本的唯一来源是同目录 README.md 的「事实基准」；本脚本不得引入 README 未定义的新数字。
定量指标按 `backend/data/quantitative_metrics.json` 的官方 metricLabel 对位，
不沿用 eval fixture 的错位赋值（详见 README §2）。

用法（在 backend/ 下）：
    uv run python tests/fixtures/local_e2e/shengyuan/fill_workbook.py [--out 路径]
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from sustainability_desk.structured_input_service import blank_unified_workbook_template

FIXTURE_DIR = Path(__file__).resolve().parent
DEFAULT_OUT = FIXTURE_DIR / "shengyuan-unified-workbook.xlsx"

COMPANY_PROFILE = """晟原精密电子股份有限公司成立于 2008 年，是一家专注于精密电子元器件与连接器研发、制造与销售的企业，注册地与主生产基地均位于江苏省苏州市工业园区。公司成立初期以消费电子领域的精密冲压端子加工起步，2012 年前后建立自有模具开发能力，逐步由来料加工转向自主设计制造。2015 年起顺应下游电动化趋势进入新能源汽车零部件配套领域，2019 年通过 IATF 16949 汽车质量管理体系认证，形成消费电子与新能源汽车并重的业务结构。

公司现有产品覆盖高速连接器、精密冲压端子与结构件三条产品线，应用于消费电子、新能源汽车与工业控制领域。其中消费电子与新能源汽车为两大主要应用市场，合计约占营业收入的 75%，工业控制及其他应用约占 25%。客户群体主要为境内外整机厂与车企一级供应商，区域覆盖华东、华北与珠三角等制造业集聚区；2025 年度活跃客户约 68 家，前三大客户营业收入贡献度约 26%，客户集中度相对均衡。

公司核心能力集中在精密模具开发与高一致性批量制造：设有产品研发中心与工艺创新中心，配备模具研发实验室与连接器性能测试室，具备尺寸检测、盐雾试验、插拔寿命等检验能力。截至 2025 年末在册员工 860 人，其中研发团队约 90 人；2025 年研发投入 3,900 万元，重点投向新能源汽车用高温高可靠性连接器与精密冲压工艺自动化。

公司已建立并持续运行 ISO 9001 质量管理体系、ISO 14001 环境管理体系、ISO 45001 职业健康安全管理体系与 IATF 16949 汽车质量管理体系。公司秉持精密制造与绿色低碳并重的经营理念，在满足客户质量与交付要求的同时，推进能源效率提升、废弃物资源化与供应链可持续管理。"""

ARTICLES = """公司依照《中华人民共和国公司法》及公司章程建立法人治理结构，由股东会、董事会、监事会与高级管理层构成，各治理主体职责边界与运行安排在公司章程中予以明确。

股东会为公司权力机构，依法行使审议批准公司经营方针与投资计划、选举和更换董事与监事、审议批准年度财务预算方案与决算方案等职权。

董事会对股东会负责，负责决定公司经营计划与投资方案、制定年度财务预算与决算方案、聘任或解聘高级管理人员，并决定公司内部管理机构的设置。董事会下设战略委员会与审计委员会，就重大投资与财务事项提出专业意见。2025 年度董事会共召开会议 8 次，董事平均出席率 96%。

监事会对股东会负责，负责检查公司财务、对董事与高级管理人员执行职务的行为进行监督。

高级管理层由总经理及各分管副总经理组成，在董事会授权范围内负责日常经营管理，组织实施董事会决议，并向董事会报告经营情况与重大事项进展。

可持续发展相关事项的治理安排在上述法人治理结构之下另行设立，具体见可持续发展治理架构说明。"""

GOVERNANCE_STRUCTURE = """公司可持续发展（ESG）事项实行「董事会—可持续发展委员会—可持续发展办公室—专项工作组」四层治理架构。

第一层为董事会，是公司可持续发展事项的最高决策机构，每年听取一次可持续发展工作专项汇报，审议年度可持续发展报告，并对重大可持续发展目标与投入作出决策。

第二层为可持续发展委员会，于 2024 年 10 月依据《可持续发展委员会工作规则》（SY-GOV-001）设立，向董事会报告。委员会主任委员由总经理陆勤担任，委员包括分管生产、质量与 EHS 的副总经理周敏，以及质量、人力资源、采购、法务合规等相关部门负责人。2025 年度委员会共召开 3 次会议。

第三层为可持续发展办公室，为常设执行机构，挂靠总经理办公室，主任由何静担任，向可持续发展委员会报告。办公室负责日常协调、信息收集与报告编制。

第四层为四个专项工作组，向可持续发展办公室报告：环境与能源工作组（牵头人 EHS 主管汪磊）、员工与安全工作组（牵头人人力资源部部长韩雪）、供应链工作组（牵头人采购部部长徐涛）、合规与商业道德工作组（牵头人法务合规专员秦朗）。各工作组成员由相关部门业务骨干兼任。"""

GOVERNANCE_DUTIES = """董事会负责审定公司可持续发展战略方向与年度目标，审议年度可持续发展报告，决策重大环境与社会议题相关投入；每年听取一次专项汇报，重大事项由可持续发展委员会随时报告。

可持续发展委员会负责将董事会确定的方向转化为年度重点工作，审议议题重要性评估结果、年度目标与实施计划，协调跨部门资源，审核报告披露口径；委员会向董事会报告工作。

可持续发展办公室负责统筹日常工作：组织年度议题重要性评估、收集与核对各部门可持续发展数据、编制年度可持续发展报告、跟踪各工作组行动项进展；办公室按月向委员会汇总进展。

四个专项工作组按职责领域分工落实，按季度向可持续发展办公室提交进展。环境与能源工作组负责环境合规、能源与排放、水资源与废弃物管理相关工作的执行与数据归集；员工与安全工作组负责人力资本发展、职业健康与安全相关制度执行与数据归集；供应链工作组负责供应商准入、评价、尽职调查与退出机制的执行；合规与商业道德工作组负责反商业贿赂、反不正当竞争与广告宣传合规审查的执行。

可持续发展相关责任嵌入各业务条线的日常职责，不另设平行序列：相关工作由对应工作组按领域牵头协调，由业务条线在本职流程中执行。数据归集实行「谁产生、谁负责」原则，由各数据来源指定提供人对本条线数据的准确性负责。"""

BASICS: dict[int, object] = {
    8: "晟原精密电子股份有限公司",
    9: "晟原精密",
    10: "制造业",
    11: "计算机、通信和其他电子设备制造业",
    12: COMPANY_PROFILE,
    13: 2025,
    14: "2025-01-01",
    15: "2025-12-31",
    16: "仅公司本部（无并表子公司）",
    17: "报告范围为公司本部及苏州主生产基地，无并表子公司。",
    21: "是",
    22: "公司官网发布",
    23: "https://www.example-shengyuan.com",
    24: 2026,
    25: "2026-04",
    26: "董事会",
    31: "esg@example-shengyuan.com",
    32: "江苏省苏州市工业园区晟原路 18 号",
    33: "0512-00000000",
    34: 8,
    35: 96,
    36: ARTICLES,
    37: GOVERNANCE_STRUCTURE,
    38: GOVERNANCE_DUTIES,
}

# 行号 -> (财务重要性, 影响重要性)；行名与 topic_registry 官方议题名一一对应。
# 阈值 4.0/4.0（contract/assessment_classify.py:23）：两维均达标才是 dual 象限，
# 才展开「治理／战略／风险机遇／指标目标」四要素；否则收敛成单块摘要。
# 晟原是精密电子制造商（营收 6.5 亿、员工 860、自有电镀与表面处理产线），
# 23 个议题按其真实经营后果判定，全部构成双重重要性——制造业的环境、用工与
# 产品责任议题既有外部影响，也有可量化的财务后果（处罚、停产、索赔、客户流失）。
SCORES: dict[int, tuple[float, float]] = {
    10: (4.4, 4.3),   # 应对气候变化：客户碳足迹要求与能耗成本
    11: (4.3, 4.4),   # 环境合规管理：环保处罚直接触发停产整改
    12: (4.2, 4.5),   # 污染物排放管理：电镀废水与废气是核心合规风险
    13: (4.1, 4.3),   # 废弃物管理：危废处置成本与转移联单合规
    14: (4.0, 4.2),   # 促进循环经济：金属边角料回用直抵原材料成本
    15: (4.4, 4.2),   # 能源管理：外购电力 42,500 MWh 是主要可控成本
    16: (4.0, 4.2),   # 水资源管理：园区取水指标与超额水价
    17: (4.0, 4.1),   # 生态系统与生物多样性保护：厂区周边水体影响与用地合规
    19: (4.3, 4.4),   # 人力资本发展：技术工人留存决定良率与交付
    20: (4.5, 4.5),   # 职业健康与安全：冲压与电镀工序工伤直接中断生产
    21: (4.0, 4.1),   # 乡村振兴：本地用工与技能培训的属地经营基础
    22: (4.0, 4.1),   # 社会贡献：属地社区关系影响用地与用工
    23: (4.4, 4.3),   # 可持续供应链管理：单一来源中断即断线
    24: (4.2, 4.1),   # 平等对待中小企业：账期合规与中小供应商稳定性
    25: (4.0, 4.0),   # 科技伦理：智能检测算法的数据使用边界
    26: (4.4, 4.2),   # 创新驱动：模具与工艺研发决定议价能力
    27: (4.6, 4.5),   # 产品质量与安全：批量退货索赔与客户资格取消
    28: (4.2, 4.3),   # 客户服务质量管理：响应时效进入客户年度考评
    29: (4.3, 4.4),   # 数据安全与客户隐私保护：客户图纸与工艺参数保密义务
    31: (4.2, 4.4),   # 反商业贿赂与反贪污：采购环节合规与客户审核红线
    32: (4.0, 4.2),   # 反不正当竞争：同业技术秘密与宣传合规
    33: (4.1, 4.1),   # 风险管理：单一来源中断、客户资格与合规处罚的统一识别与应对
    34: (4.1, 4.2),   # 尽职调查：供应商准入与客户合规审核前置条件
}

# 指标 key -> (数值, 数据提供部门, 备注)。r07 由范围一与范围二自动求和，不写。
METRICS: dict[str, tuple[object, str, str]] = {
    "economic_environment_r03": (650, "财务部", "2025 年度营业收入"),
    "economic_environment_r04": (12800, "EHS", "苏州主厂区天然气与工艺直接排放"),
    "economic_environment_r05": (26400, "EHS", "外购电力产生的间接排放（市场法）"),
    "economic_environment_r11": (12100, "EHS", "外购电力 42,500 MWh 折算标准煤"),
    # 轻量版工作簿不含 r14 柴油 / r15 天然气（lightweight=False），直接能源合并进「传统能源：其他」。
    "economic_environment_r17": (6500, "EHS", "天然气 6,474 与柴油 26 折算标准煤合计"),
    "economic_environment_r19": (1200, "EHS", "屋顶光伏发电自用部分折算"),
    "economic_environment_r28": (1.6, "EHS", "主厂区废气治理设施处理后排放口"),
    "economic_environment_r30": (8600, "EHS", "预处理后纳入园区污水管网的工业废水"),
    "economic_environment_r38": (42.6, "EHS", "一般工业固废、危险废弃物与生活垃圾合计"),
    "economic_environment_r39": (39.4, "EHS", "金属边角料、一般包装废弃物与生活垃圾"),
    "economic_environment_r40": (3.2, "EHS", "含油抹布、废包装桶与废活性炭等"),
    "economic_environment_r41": (18.0, "EHS", "金属边角料与部分包装材料资源化利用"),
    "economic_environment_r42": (12500, "EHS", "主厂区生产与生活用水"),
    "social_r02": (860, "人力资源部", "报告期末在册员工"),
    "social_r03": (520, "人力资源部", "报告期末男性员工"),
    "social_r04": (340, "人力资源部", "报告期末女性员工"),
    "social_r10": (2, "人力资源部", "博士学历"),
    "social_r11": (48, "人力资源部", "硕士学历"),
    "social_r12": (310, "人力资源部", "本科学历"),
    "social_r13": (500, "人力资源部", "本科以下学历"),
    "social_r17": (8.5, "人力资源部", "离职 72 人 ÷ 平均在册 846 人"),
    "social_r25": (86, "人力资源部", "覆盖生产、质量与管理岗位"),
    "social_r26": (12, "人力资源部", "培训总学时 10,248 ÷ 员工总数 860"),
    "social_r53": (0, "质量部", "报告期无产品召回"),
    "social_r54": (0, "质量部", "报告期无产品召回"),
    "social_r55": (6, "客户服务部", "产品及服务相关投诉"),
    "social_r56": (100, "客户服务部", "已受理投诉处理完成率"),
    "social_r64": (186, "采购部", "纳入采购系统的活跃供应商"),
    "social_r65": (168, "采购部", "中国内地供应商"),
    "social_r66": (4, "采购部", "港澳台地区供应商"),
    "social_r67": (14, "采购部", "国外供应商"),
    "social_r69": (6, "采购部", "年度评价与退出记录"),
    "governance_r06": (4, "法务合规", "采购、财务、销售与供应商对接岗位廉洁培训"),
    "governance_r07": (128, "法务合规", "廉洁培训签到统计"),
}

QUANTITATIVE_SHEETS = ("经济+环境", "社会", "治理")

# 每项指标必须且只能填「数值」或「无值原因」之一（quantitative_parser.py:326 fail-loud）。
# 未纳入本 corpus 的指标统一给出未收集的理由，而非留空。
DEFAULT_NO_VALUE_REASON = "not_collected"

# 填写了温室气体数值时必须选择核算标准；下拉值须精确匹配。
GHG_STANDARD_CELL = ("定量填写说明", "B8")
GHG_STANDARD = "ISO 14064-1:2018及《工业企业温室气体排放核算和报告通则》（GB/T 32150-2015）"


def fill(out_path: Path) -> Path:
    """生成填写后的工作簿；只写可写单元格，隐藏元数据保持原样。"""

    blank = FIXTURE_DIR / ".blank-unified-workbook.xlsx"
    blank.write_bytes(blank_unified_workbook_template())
    workbook = openpyxl.load_workbook(blank)

    basics = workbook["基础资料"]
    for row, value in BASICS.items():
        basics.cell(row=row, column=6, value=value)

    scoring = workbook["重要性评分表"]
    for row, (financial, impact) in SCORES.items():
        scoring.cell(row=row, column=2, value=financial)
        scoring.cell(row=row, column=3, value=impact)

    written: set[str] = set()
    for name in QUANTITATIVE_SHEETS:
        sheet = workbook[name]
        for row in range(9, sheet.max_row + 1):
            key = sheet.cell(row=row, column=1).value
            if not isinstance(key, str) or not key.startswith(
                ("economic_environment_", "social_", "governance_")
            ):
                continue
            if key not in METRICS:
                # 自动求和行由公式给出，既不填数值也不填无值原因。
                if sheet.cell(row=row, column=5).value is None:
                    sheet.cell(row=row, column=6, value=DEFAULT_NO_VALUE_REASON)
                continue
            value, department, remark = METRICS[key]
            sheet.cell(row=row, column=5, value=value)
            sheet.cell(row=row, column=7, value=department)
            sheet.cell(row=row, column=8, value=remark)
            written.add(key)

    standard_sheet, standard_cell = GHG_STANDARD_CELL
    workbook[standard_sheet][standard_cell] = GHG_STANDARD

    missing = set(METRICS) - written
    if missing:
        raise SystemExit(f"以下指标 key 在工作簿中未找到，请核对官方目录：{sorted(missing)}")

    workbook.save(out_path)
    blank.unlink()
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    path = fill(args.out)
    print(f"已生成 {path}（基础资料 {len(BASICS)} 项 / 评分 {len(SCORES)} 议题 / 定量 {len(METRICS)} 指标）")


if __name__ == "__main__":
    main()
