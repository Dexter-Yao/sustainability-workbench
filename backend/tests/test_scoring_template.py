# ABOUTME: 动态重要性评分模板测试，确保模板仅投影当前适用议题与统一评分尺度。
# ABOUTME: 模板必须可被同源解析器回读，避免下载模板与上传合同产生两套规则。
from io import BytesIO
from uuid import UUID

import openpyxl

from sustainability_desk.assets.scoring import parse_scoring
from sustainability_desk.assets.scoring_template import SCORING_SHEET_TITLE, create_scoring_template
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.models import (
    AssessmentScoreScale,
    DisclosureProfile,
    Field,
    Report,
)
from sustainability_desk.contract.structured_inputs import StructuredInputContext
from knowledge_package_fixtures import SSE_PACKAGE

CONTEXT = StructuredInputContext(
    reportId=UUID("00000000-0000-0000-0000-000000000001"),
    contractVersion="cv-test",
    compiledSemanticsVersion="semantics-test",
)


def _report(*, technology_ethics: str = "否") -> Report:
    return Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="t",
        sections=[],
        fields={
            "has_technology_ethics_sensitive_activity": Field(
                key="has_technology_ethics_sensitive_activity",
                label="科技伦理适用性",
                type="enum",
                source="user_input",
                value=technology_ethics,
                options=["是", "否"],
            )
        },
        disclosureProfile=DisclosureProfile(),
    )


def _filled_template(report: Report) -> bytes:
    workbook = openpyxl.load_workbook(
        BytesIO(create_scoring_template(report, context=CONTEXT))
    )
    worksheet = workbook["重要性评分表"]
    for row in worksheet.iter_rows():
        if row[0].value and row[1].value is None and row[2].value is None and row[0].value not in {"环境", "社会", "治理"}:
            row[1].value = 4.0
            row[2].value = 4.0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_scoring_template_uses_current_applicable_topics():
    workbook = openpyxl.load_workbook(
        BytesIO(create_scoring_template(_report(), context=CONTEXT))
    )
    worksheet = workbook["重要性评分表"]

    assert worksheet["B2"].value == "重要性评分表"
    assert worksheet["A8"].value == "议题名称"
    assert worksheet["B8"].value == "财务重要性"
    assert worksheet["C8"].value == "影响重要性"
    # 表头不放品牌字标图片。
    assert len(worksheet._images) == 0
    assert worksheet.freeze_panes == "A9"
    assert "22 个适用议题" in worksheet["B4"].value
    assert "评分范围：(0, 5]；步长：0.1" in worksheet["B5"].value
    # 指纹已移入隐藏元数据页，不再在用户可见区域显示
    assert "_模板元数据" in workbook.sheetnames


def test_scoring_template_follows_technology_ethics_applicability():
    no_ethics = openpyxl.load_workbook(
        BytesIO(
            create_scoring_template(
                _report(technology_ethics="否"),
                context=CONTEXT,
            )
        )
    )["重要性评分表"]
    with_ethics = openpyxl.load_workbook(
        BytesIO(
            create_scoring_template(
                _report(technology_ethics="是"),
                context=CONTEXT,
            )
        )
    )["重要性评分表"]

    no_names = {row[0].value for row in no_ethics.iter_rows()}
    with_names = {row[0].value for row in with_ethics.iter_rows()}
    assert "科技伦理" not in no_names
    assert "科技伦理" in with_names
    assert "23 个适用议题" in with_ethics["B4"].value


def test_filled_scoring_template_round_trips_through_parser():
    report = _report()

    parsed = parse_scoring(
        _filled_template(report),
        report=report,
        context=CONTEXT,
    )

    assert len(parsed.scored) == 22
    assert parsed.scored[0].assessmentTopicId == "climate_change"


def _column_b_validation(workbook) -> object:
    worksheet = workbook["重要性评分表"]
    return next(
        dv
        for dv in worksheet.data_validations.dataValidation
        if str(dv.sqref).split()[0].startswith("B")
    )


def test_score_validation_formula_is_float_tolerant():
    """旧 MOD(x,0.1) 在双精度下误拒 0.3/4.3 等合法评分；整数空间校验必须全数放行。"""
    workbook = openpyxl.load_workbook(
        BytesIO(create_scoring_template(_report(), context=CONTEXT))
    )
    validation = _column_b_validation(workbook)
    assert "MOD(ROUND(B" in validation.formula1
    assert ",1)=0" in validation.formula1

    def excel_passes(x: float, factor: int, units: int) -> bool:
        scaled = x * factor
        return (
            x > 0
            and x <= 5
            and abs(scaled - round(scaled)) < 1e-9
            and round(scaled) % units == 0
        )

    assert all(excel_passes(i / 10, 10, 1) for i in range(1, 51))
    assert not excel_passes(4.35, 10, 1)


def test_score_validation_formula_respects_half_step_scale():
    report = _report()
    report.assessmentScoreScale = AssessmentScoreScale(
        minimumExclusive=0, maximum=5, multipleOf=0.5
    )
    workbook = openpyxl.load_workbook(
        BytesIO(create_scoring_template(report, context=CONTEXT))
    )
    validation = _column_b_validation(workbook)
    assert "MOD(ROUND(B" in validation.formula1
    assert ",5)=0" in validation.formula1

    def excel_passes(x: float) -> bool:
        scaled = x * 10
        return (
            x > 0
            and x <= 5
            and abs(scaled - round(scaled)) < 1e-9
            and round(scaled) % 5 == 0
        )

    assert excel_passes(4.5)
    assert not excel_passes(4.3)


def test_scoring_sheet_carries_materiality_definitions_from_the_contract() -> None:
    """双重重要性是最难的概念，离线填表的人也必须拿到判断依据。

    若释义只硬编码在 React 组件里，下载的 xlsx 仅有「财务重要性/影响重要性」两个
    裸列头。释义由合同评分尺度承载，网页与 Excel 共用同一份，不在任一端另写。
    """

    report = load_contract(
        SSE_PACKAGE.report_contract_path
    )
    scale = report.assessmentScoreScale
    workbook = openpyxl.load_workbook(
        BytesIO(create_scoring_template(report, context=CONTEXT))
    )
    text = workbook[SCORING_SHEET_TITLE]["B6"].value or ""

    for part in (
        scale.financialMaterialityDefinition,
        scale.financialMaterialityExplanation,
        scale.impactMaterialityDefinition,
        scale.impactMaterialityExplanation,
    ):
        assert part, "合同必须声明双重重要性两个维度的释义"
        assert part in text, f"评分表缺少合同释义：{part}"
