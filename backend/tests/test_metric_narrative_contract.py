# ABOUTME: 指标与目标 source contract、Planner 分支和 typed 无值正文授权的端到端合同测试。
# ABOUTME: 测试直接覆盖重要性、目录映射和用户数值三类输入，避免从运行态结构反向猜测规则。
from pathlib import Path

import pytest
import yaml
from docx import Document

from assessment_fixtures import complete_assessment
from sustainability_desk.contract.loader import load_contract
from sustainability_desk.contract.metric_narrative import MetricNarrativePolicy
from sustainability_desk.contract.models import Report, Section
from sustainability_desk.contract.topic_section_template import (
    MetricDisclosureSpec,
    TopicSectionTemplate,
    compile_topic_section_template,
    compiled_metric_disclosure_keys,
)
from sustainability_desk.contract.visibility import visible
from sustainability_desk.llm.prompts import build_model_context, render_prompt
from sustainability_desk.export.docx_renderer import render_docx
from sustainability_desk.planner import assemble_report
from sustainability_desk.quantitative_metrics import (
    quantitative_metric_display_label,
    quantitative_metrics_by_key,
)
from knowledge_package_fixtures import SSE_PACKAGE
from sustainability_desk.planner import load_topic_templates_from

BACKEND = Path(__file__).resolve().parents[1]
CONTRACT = SSE_PACKAGE.report_contract_path
TOPIC_DIR = SSE_PACKAGE.topic_sections_dir


def _assembled(
    report_section_id: str,
    materiality: str,
    *,
    metrics: dict[str, dict[str, str]] | None = None,
) -> tuple[Report, Section]:
    fixed = load_contract(CONTRACT)
    if metrics is not None:
        fixed = fixed.model_copy(
            update={"meta": {"quantitativeMetrics": {"metrics": metrics}}}
        )
    assessment = complete_assessment(
        fixed,
        materialities={report_section_id: materiality},
    )
    report = assemble_report(
        fixed,
        load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE),
        assessment=assessment,
    ).report
    section = next(
        section
        for module in report.sections
        for section in module.children or []
        if section.reportSectionId == report_section_id
    )
    return report, section


def _assembled_with_materialities(
    report_section_id: str,
    materialities: dict[str, str],
) -> tuple[Report, Section]:
    fixed = load_contract(CONTRACT)
    assessment = complete_assessment(fixed, materialities=materialities)
    report = assemble_report(
        fixed,
        load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE),
        assessment=assessment,
    ).report
    section = next(
        section
        for module in report.sections
        for section in module.children or []
        if section.reportSectionId == report_section_id
    )
    return report, section


def _metrics_h3(section: Section) -> Section | None:
    return next((child for child in section.children or [] if child.title == "指标与目标"), None)


def _metric_narrative_block(metrics: Section):
    return next(
        block
        for block in metrics.blocks
        if block.id.endswith(".metrics_narrative_body")
    )


def _word_h2_segment(document: Document, title: str) -> list[tuple[str, str]]:
    segment: list[tuple[str, str]] = []
    collecting = False
    for paragraph in document.paragraphs:
        style = paragraph.style.name
        if style == "Heading 2":
            if collecting:
                break
            collecting = paragraph.text.endswith(title)
        if collecting:
            segment.append((style, paragraph.text))
    assert segment, f"Word 未找到 H2：{title}"
    return segment


def test_impact_and_non_materiality_do_not_assemble_metric_h3() -> None:
    for materiality in ("impact", "non"):
        _report, section = _assembled("energy_management", materiality)
        assert _metrics_h3(section) is None


def test_mapped_metric_with_any_value_shows_chart_and_hides_ai_narrative() -> None:
    report, section = _assembled(
        "energy_management",
        "dual",
        metrics={"economic_environment_r11": {"value": "0"}},
    )
    metrics = _metrics_h3(section)
    assert metrics is not None
    assert [block.id for block in metrics.blocks] == [
        "energy_management.metrics_summary",
        "energy_management.metrics_narrative_body",
        "energy_management.layout_assets",
    ]
    assert not visible(_metric_narrative_block(metrics), report)


def test_mapped_metric_without_values_gets_catalog_policy() -> None:
    report, section = _assembled("energy_management", "financial")
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)
    assert visible(block, report)
    context = build_model_context(block, report, report)
    assert context.metric_narrative_policy is not None
    assert context.evidence_posture.level == "metric_narrative"
    assert context.metric_narrative_policy.indicator_source == "catalog"
    assert context.metric_narrative_policy.catalog_metric_labels[0] == "电力消耗量"
    assert context.display_title_task is None


def test_metric_narrative_catalog_deduplicates_display_labels() -> None:
    report, section = _assembled("pollutant_emissions_management", "dual")
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)

    context = build_model_context(block, report, report)

    assert context.metric_narrative_policy is not None
    labels = context.metric_narrative_policy.catalog_metric_labels
    assert labels.count("其它") == 1
    assert len(labels) == len(set(labels))


def test_metric_narrative_prompt_replaces_generic_pillar_and_duplicate_metric_evidence() -> None:
    report, section = _assembled("innovation_driven", "dual")
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)

    context = build_model_context(block, report, report)
    system, user = render_prompt(context, n=1)
    prompt = "\n".join((system, user))

    assert context.pillar_purpose is not None
    assert context.pillar_purpose.name == "指标与目标"
    assert "已显示章节标题下的局部正文" not in context.section_task
    assert "<pillar_purpose>" not in system
    assert "适合关注的指标、目标或披露维度" not in system
    assert "<metric_narrative_policy>" in system
    assert "以下句式仅用于校准信息顺序和承诺强度，不要求逐字照抄" in system
    assert "公司高度重视……管理" in system
    assert "篇幅约 60–160 字" in system
    assert "彼此在表达、结构与侧重上有明显差异" not in system
    assert "已授权专利项目数（境内）" in prompt
    assert prompt.count("已授权专利项目数（境内）") == 1
    assert "（件）" not in prompt
    assert "本节没有报告期指标数值" in user


def test_metric_narrative_note_is_visible_once_as_user_fact() -> None:
    report, section = _assembled(
        "energy_management",
        "dual",
        metrics={
            "economic_environment_r11": {
                "note": "按全年账单口径统计，暂未录入数值。",
            }
        },
    )
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)

    context = build_model_context(block, report, report)
    system, user = render_prompt(context, n=1)

    assert context.metric_narrative_policy is not None
    assert context.evidence.substantive_input_present
    assert "按全年账单口径统计，暂未录入数值。" not in system
    assert user.count("按全年账单口径统计，暂未录入数值。") == 1
    assert "指标名称：电力消耗量" in user
    assert "分类路径" not in user
    assert "单位：" not in user
    assert "本节没有报告期指标数值" not in user


def test_metric_narrative_multiple_candidates_only_vary_natural_wording() -> None:
    report, section = _assembled("innovation_driven", "dual")
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)

    context = build_model_context(block, report, report)
    system, _user = render_prompt(context, n=3)

    assert "<expression_guidance>" not in system
    assert "不先概括企业背景、议题定位或管理态度" not in system
    assert "返回恰好 3 个完整版本" in system
    assert "均覆盖上述相同信息层级" in system
    assert "仅在自然措辞上适度变化" in system
    assert "结构与侧重上有明显差异" not in system


def test_unmapped_metric_without_values_gets_general_indicator_policy() -> None:
    report, section = _assembled("sme_fair_treatment", "dual")
    metrics = _metrics_h3(section)
    assert metrics is not None
    block = _metric_narrative_block(metrics)
    context = build_model_context(block, report, report)
    assert context.metric_narrative_policy is not None
    assert context.metric_narrative_policy.indicator_source == "generated_general"
    assert context.metric_narrative_policy.catalog_metric_labels == ()


def test_impact_summary_prompt_receives_metric_catalog_without_metric_policy() -> None:
    report, section = _assembled("innovation_driven", "impact")
    block = section.blocks[0]
    template = Report(knowledgePackageId=SSE_PACKAGE.id, 
        title="测试模板",
        sections=[load_topic_templates_from(TOPIC_DIR, package=SSE_PACKAGE)["innovation_driven"]],
    )
    context = build_model_context(block, template, report)
    prompt = "\n".join(render_prompt(context, n=1))

    assert "已授权专利项目数（境内）" in prompt
    assert "在申请中专利项目数" in prompt
    assert context.evidence.substantive_input_present is False
    assert context.metric_narrative_policy is None
    assert "正在着手建立指标管理体系" not in prompt


@pytest.mark.parametrize("path", sorted(TOPIC_DIR.glob("*.yaml")), ids=lambda path: path.stem)
def test_concise_disclosure_inherits_compiled_metric_catalog(path: Path) -> None:
    """全部 H2 摘要从唯一 metricDisclosure 编译结果继承目录，不在摘要 YAML 重复 key。"""

    source = TopicSectionTemplate.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")), context={"package": SSE_PACKAGE})
    section = compile_topic_section_template(source, package=SSE_PACKAGE)
    assert compiled_metric_disclosure_keys(section) == source.metricDisclosure.catalogMetricKeys
    # 指标摘要图不带题注：所属议题与「指标与目标」小节标题已给全，题注只会复述上两级标题。
    def _walk(node):
        yield from node.blocks
        for child in node.children or []:
            yield from _walk(child)

    summaries = [
        block for block in _walk(section)
        if block.image is not None and block.image.derivedVisualization is not None
    ]
    if source.metricDisclosure.catalogMetricKeys:
        assert [block.image.caption for block in summaries] == [None]
    else:
        assert summaries == []
    report = Report(knowledgePackageId=SSE_PACKAGE.id, title="测试报告", sections=[section])
    block = section.conciseDisclosure
    assert block is not None

    context = build_model_context(block, report, report)
    catalog = quantitative_metrics_by_key(SSE_PACKAGE)
    assert [metric.metric_name for metric in context.evidence.metric_evidence] == [
        quantitative_metric_display_label(catalog[key])
        for key in source.metricDisclosure.catalogMetricKeys
    ]
    assert context.evidence.substantive_input_present is False
    assert context.metric_narrative_policy is None


@pytest.mark.parametrize(
    ("report_section_id", "title", "materiality", "expect_metrics"),
    [
        ("sme_fair_treatment", "平等对待中小企业", "dual", True),
        ("circular_economy_promotion", "促进循环经济", "impact", False),
        ("circular_economy_promotion", "促进循环经济", "financial", True),
        ("innovation_driven", "创新驱动", "impact", False),
        ("innovation_driven", "创新驱动", "dual", True),
    ],
)
def test_word_projects_metric_branch_for_representative_topics(
    report_section_id: str,
    title: str,
    materiality: str,
    expect_metrics: bool,
    base_template: Path,
    tmp_path: Path,
) -> None:
    report, _section = _assembled(report_section_id, materiality)
    output = render_docx(
        report,
        base_template,
        tmp_path / f"{report_section_id}-{materiality}.docx",
    )
    segment = _word_h2_segment(Document(str(output)), title)

    assert any(style == "Heading 3" and text.endswith("指标与目标") for style, text in segment) is expect_metrics
    assert not any(
        style == "Heading 4" and text.endswith("指标管理与披露进展")
        for style, text in segment
    )


def test_word_aggregates_merged_h2_materiality(
    base_template: Path,
    tmp_path: Path,
) -> None:
    report, _section = _assembled_with_materialities(
        "rural_revitalization_social_contribution",
        {"rural_revitalization": "impact", "social_contribution": "financial"},
    )
    output = render_docx(report, base_template, tmp_path / "rural-merged.docx")
    segment = _word_h2_segment(Document(str(output)), "乡村振兴与社会贡献")

    assert any(style == "Heading 3" and text.endswith("指标与目标") for style, text in segment)
    assert not any(
        style == "Heading 4" and text.endswith("指标管理与披露进展")
        for style, text in segment
    )


def test_metric_disclosure_rejects_unknown_and_duplicate_keys() -> None:
    """目录只有一份、无子集校验；未知与重复仍拒绝。"""
    with pytest.raises(ValueError, match="未知指标"):
        MetricDisclosureSpec.model_validate(
            {"catalogMetricKeys": ("unknown",)}, context={"package": SSE_PACKAGE}
        )
    with pytest.raises(ValueError, match="不得重复"):
        MetricDisclosureSpec.model_validate(
            {"catalogMetricKeys": ("economic_environment_r11",) * 2},
            context={"package": SSE_PACKAGE},
        )


def test_metric_narrative_policy_rejects_ambiguous_indicator_sources() -> None:
    with pytest.raises(ValueError, match="必须提供目录指标名称"):
        MetricNarrativePolicy(indicator_source="catalog")
    with pytest.raises(ValueError, match="不得携带目录指标名称"):
        MetricNarrativePolicy(
            indicator_source="generated_general",
            catalog_metric_labels=("循环材料使用率（%）",),
        )
    with pytest.raises(ValueError, match="catalog_metric_labels 不得重复"):
        MetricNarrativePolicy(
            indicator_source="catalog",
            catalog_metric_labels=("其它", "其它"),
        )
    with pytest.raises(ValueError, match="allowed_processes 不得重复"):
        MetricNarrativePolicy(
            indicator_source="generated_general",
            allowed_processes=("statistics", "statistics"),
        )


def test_source_template_rejects_hand_authored_metric_h3_and_ai_placeholder_content() -> None:
    metric_h3 = Section(key="x.metrics", title="指标与目标", headingLevel=3)
    with pytest.raises(ValueError, match="三个 H3"):
        TopicSectionTemplate.model_validate(
            {
                "section": {
                    "key": "x",
                    "title": "X",
                    "headingLevel": 2,
                    "reportSectionId": "x",
                    "children": [metric_h3.model_dump(mode="json")],
                },
                "metricDisclosure": {"catalogMetricKeys": []},
            }, context={"package": SSE_PACKAGE}
        )
    with pytest.raises(ValueError, match="不得预置模板正文"):
        TopicSectionTemplate.model_validate(
            {
                "section": {
                    "key": "x",
                    "title": "X",
                    "headingLevel": 2,
                    "reportSectionId": "x",
                    "children": [
                        {
                            "key": "x.gov",
                            "title": "治理",
                            "pillar": "governance",
                            "headingLevel": 3,
                            "blocks": [
                                {
                                    "id": "x.body",
                                    "type": "paragraph",
                                    "blockType": "constrained",
                                    "source": "ai",
                                    "generation": {"task": {"focus": "形成正文"}},
                                    "content": [{"kind": "text", "text": "占位正文"}],
                                }
                            ],
                        },
                        {"key": "x.strategy", "title": "战略", "pillar": "strategy", "headingLevel": 3},
                        {"key": "x.iro", "title": "影响、风险与机遇管理", "pillar": "iro_management", "headingLevel": 3},
                    ],
                    "conciseDisclosure": {
                        "id": "x.concise",
                        "type": "paragraph",
                        "blockType": "constrained",
                        "source": "ai",
                        "generation": {"task": {"focus": "形成摘要"}},
                    },
                },
                "metricDisclosure": {"catalogMetricKeys": []},
            }, context={"package": SSE_PACKAGE}
        )
